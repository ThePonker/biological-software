"""Tabella -- Refresh Records Cache (v2 -- May 2026)

Queries observatum.db for personal observations, commercial observations,
and specimen records. Scans active workbooks for pending (un-imported)
entries.

Writes a CSV to Tabella/output/.records_cache.csv that the VBA macro
reads into the workbook.

Called from VBA via Shell, or manually:
    python -m Tabella.refresh_records
    python -m Tabella.refresh_records --skip "C:\\path\\to\\my.xlsm"

v2 changes (May 2026):
  - Added --skip <path> argument (repeatable). Workbooks listed are
    excluded from the pending scan. VBA passes ThisWorkbook.FullName
    so the user's currently-open workbook is excluded from the CSV
    (CountOwnPending in VBA covers that workbook live).
  - Pending scan now reads the TVK column directly (header lookup)
    instead of resolving species names. Fixes silent miscount for
    species typed with " agg.", " s.l.", or " s.s." suffixes.

Output columns: TVK, PersonalCount, PersonalLastDate, CommercialCount,
                CommercialProjects, SpecimenCount, PendingPersonal,
                PendingCommercial
"""

import argparse
import csv
import os
import sqlite3
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import paths

OUTPUT = Path(__file__).resolve().parent / "output" / ".records_cache.csv"

# Active workbooks folder -- scanned for pending entries
ACTIVE_WORKBOOKS_DIR = Path.home() / "OneDrive" / "Active Record Books"


def _normalise(p) -> str:
    """Canonical path string for comparison (case-insensitive on Windows)."""
    try:
        return os.path.normcase(os.path.normpath(str(Path(p).resolve())))
    except Exception:
        return os.path.normcase(os.path.normpath(str(p)))


def _scan_workbooks(workbooks_dir: Path, skip_paths: set) -> tuple:
    """Scan active workbooks for pending personal and commercial entries.

    Pending entries are identified by reading the TVK column directly
    (header "TVK" in row 4) -- so " agg." / " s.l." / " s.s." species
    are counted correctly because they share the same TVK as the base
    species.

    Skips workbooks whose normalised path is in skip_paths.

    Returns:
        (pending_personal, pending_commercial) -- each {tvk: count}
    """
    pending_personal = {}
    pending_commercial = {}

    if not workbooks_dir.exists():
        return pending_personal, pending_commercial

    try:
        import openpyxl
    except ImportError:
        print("[Refresh] openpyxl not installed -- skipping workbook scan",
              file=sys.stderr)
        return pending_personal, pending_commercial

    xlsm_files = list(workbooks_dir.glob("*.xlsm"))
    if not xlsm_files:
        return pending_personal, pending_commercial

    sheet_map = {
        "Personal Observations": pending_personal,
        "Commercial Observations": pending_commercial,
    }

    for xlsm_path in xlsm_files:
        # Skip backups
        if ".bak-" in xlsm_path.stem:
            continue

        # Skip explicitly excluded workbooks (caller's own workbook)
        if _normalise(xlsm_path) in skip_paths:
            print(f"[Refresh] Skipped (own workbook): {xlsm_path.name}")
            continue

        try:
            wb = openpyxl.load_workbook(
                str(xlsm_path), read_only=True, data_only=True,
                keep_vba=False
            )
        except Exception as e:
            print(f"[Refresh] Could not open {xlsm_path.name}: {e}",
                  file=sys.stderr)
            continue

        for sheet_name, target_dict in sheet_map.items():
            if sheet_name not in wb.sheetnames:
                continue

            ws = wb[sheet_name]

            # Locate the TVK column by header in row 4
            tvk_col = None
            for cell in ws[4]:
                if cell.value and str(cell.value).strip() == "TVK":
                    tvk_col = cell.column
                    break
            if tvk_col is None:
                continue

            for row in ws.iter_rows(min_row=5, min_col=tvk_col,
                                    max_col=tvk_col, values_only=True):
                tvk = row[0]
                if not tvk or not isinstance(tvk, str):
                    continue
                tvk = tvk.strip()
                if not tvk or tvk == "?":
                    continue
                target_dict[tvk] = target_dict.get(tvk, 0) + 1

        wb.close()
        print(f"[Refresh] Scanned {xlsm_path.name}")

    return pending_personal, pending_commercial


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--skip", action="append", default=[],
        help="Path of a workbook to exclude from the pending scan. "
             "Pass once per workbook to skip. Typically this is "
             "ThisWorkbook.FullName from the VBA caller."
    )
    args = parser.parse_args()

    skip_paths = {_normalise(p) for p in args.skip}
    if skip_paths:
        for p in args.skip:
            print(f"[Refresh] Will skip: {p}")

    db_path = paths.OBSERVATUM_DB
    if not db_path.exists():
        print(f"Database not found: {db_path}", file=sys.stderr)
        sys.exit(1)

    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row

    # Personal observations: count + last date per TVK
    personal = {}
    for row in conn.execute("""
        SELECT species_tvk, COUNT(*) as cnt, MAX(date) as last_date
        FROM observations
        WHERE record_type = 'Personal' AND species_tvk IS NOT NULL
              AND species_tvk != ''
        GROUP BY species_tvk
    """):
        personal[row["species_tvk"]] = {
            "count": row["cnt"],
            "last_date": (row["last_date"] or "")[:10],
        }

    # Commercial observations: count + project names per TVK
    commercial = {}
    for row in conn.execute("""
        SELECT species_tvk, COUNT(*) as cnt,
               GROUP_CONCAT(DISTINCT project_name) as projects
        FROM observations
        WHERE record_type = 'Commercial' AND species_tvk IS NOT NULL
              AND species_tvk != ''
        GROUP BY species_tvk
    """):
        projects = row["projects"] or ""
        proj_list = [p.strip() for p in projects.split(",") if p.strip()]
        if len(proj_list) > 3:
            proj_display = ", ".join(proj_list[:3]) + "..."
        else:
            proj_display = ", ".join(proj_list)
        commercial[row["species_tvk"]] = {
            "count": row["cnt"],
            "projects": proj_display,
        }

    # Specimens: count per TVK
    specimens = {}
    for row in conn.execute("""
        SELECT species_tvk, COUNT(*) as cnt
        FROM specimens
        WHERE species_tvk IS NOT NULL AND species_tvk != ''
        GROUP BY species_tvk
    """):
        specimens[row["species_tvk"]] = row["cnt"]

    conn.close()

    # Scan active workbooks for pending entries (excluding the caller's own)
    pending_personal, pending_commercial = _scan_workbooks(
        ACTIVE_WORKBOOKS_DIR, skip_paths
    )

    if pending_personal or pending_commercial:
        print(f"[Refresh] Pending: {sum(pending_personal.values())} personal, "
              f"{sum(pending_commercial.values())} commercial entries "
              f"from other workbooks")

    # Merge all TVKs
    all_tvks = (set(personal) | set(commercial) | set(specimens)
                | set(pending_personal) | set(pending_commercial))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT, "w", newline="", encoding="ascii",
              errors="replace") as f:
        writer = csv.writer(f)
        writer.writerow([
            "TVK", "PersonalCount", "PersonalLastDate",
            "CommercialCount", "CommercialProjects", "SpecimenCount",
            "PendingPersonal", "PendingCommercial"
        ])
        for tvk in sorted(all_tvks):
            p = personal.get(tvk, {})
            c = commercial.get(tvk, {})
            s = specimens.get(tvk, 0)
            pp = pending_personal.get(tvk, 0)
            pc = pending_commercial.get(tvk, 0)
            writer.writerow([
                tvk,
                p.get("count", 0),
                p.get("last_date", ""),
                c.get("count", 0),
                c.get("projects", ""),
                s,
                pp,
                pc,
            ])

    total = len(all_tvks)
    print(f"Records cache: {total} species -> {OUTPUT}")


if __name__ == "__main__":
    main()
