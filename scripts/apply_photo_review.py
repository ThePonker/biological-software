"""apply_photo_review.py -- write the links Wil chose in the photo review sheet (14 Oct 2026).

    py -3.14 scripts\\apply_photo_review.py                  # dry run on the newest sheet
    py -3.14 scripts\\apply_photo_review.py --sheet X.xlsx   # a particular sheet
    py -3.14 scripts\\apply_photo_review.py --apply          # backup, then link

Reads the "Review" sheet of scripts\\photo_review_sheet.py's workbook. Every filled "Link to
specimen id" is checked: it must be one of that row's candidates, the photo file must exist
under the library, and the photo must not already be linked to a specimen. Rows with a problem
are listed and skipped. --apply backs observatum.db up first (C:\\BiologicalSoftware_Backups\\
reference\\observatum_pre_photo_review_<time>.db), then writes the links with method 'by hand'.
Close Observatum first.
"""
import argparse
import datetime
import glob
import os
import re
import sqlite3
import sys
from pathlib import Path

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for _p in (_ROOT, os.path.join(_ROOT, "Observatum")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import paths  # noqa: E402
from scripts.add_photo_links_table import BACKUP_DIR  # noqa: E402
from scripts.photo_review_sheet import SHEET  # noqa: E402
from shared import photo_links as pl  # noqa: E402


def newest_sheet():
    found = sorted(glob.glob(os.path.join(_ROOT, "scripts", "_oneoff", "photo_review_*.xlsx")),
                   key=lambda p: (os.path.getmtime(p), p))
    return found[-1] if found else None


def read_sheet(path):
    """[(sheet row number, photo path, candidate ids, link cell value)] for rows with a photo."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET]
        rows = ws.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(rows)]
        ci = {h: head.index(h) for h in ("Candidates", "Link to specimen id", "Photo path")}
        out = []
        for n, r in enumerate(rows, 2):
            if not r or not r[ci["Photo path"]]:
                continue
            out.append((n, pl.norm(r[ci["Photo path"]]),
                        [int(x) for x in re.findall(r"\d+", str(r[ci["Candidates"]] or ""))],
                        r[ci["Link to specimen id"]]))
        return out
    finally:
        wb.close()


def plan(sheet_rows, root, linked):
    """(links [(photo path, specimen id)], problems [text], blank count)."""
    links, problems, blank, seen = [], [], 0, set()
    for n, p, cands, cell in sheet_rows:
        if cell is None or str(cell).strip() == "":
            blank += 1
            continue
        txt = str(cell).strip()
        try:
            f = float(txt)
            sid = int(f) if f == int(f) else None
        except ValueError:
            sid = None
        if sid is None:
            problems.append(f"row {n}: '{txt}' is not a specimen id ({p})")
        elif sid not in cands:
            problems.append(f"row {n}: id {sid} is not one of this photo's candidates {cands} ({p})")
        elif not (Path(root) / p).is_file():
            problems.append(f"row {n}: photo file not found ({p})")
        elif p in linked:
            problems.append(f"row {n}: already linked to a specimen ({p})")
        elif p in seen:
            problems.append(f"row {n}: photo listed twice ({p})")
        else:
            seen.add(p)
            links.append((p, sid))
    return links, problems, blank


def backup(db_path, backup_dir) -> str:
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    dest = os.path.join(backup_dir, f"observatum_pre_photo_review_{stamp}.db")
    os.makedirs(backup_dir, exist_ok=True)
    if os.path.exists(dest):
        raise FileExistsError(f"backup already exists: {dest}")
    src, d = sqlite3.connect(str(db_path)), sqlite3.connect(dest)
    try:
        src.backup(d)
    finally:
        d.close()
        src.close()
    return dest


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--root", default=str(paths.MEDIA_ROOT))
    ap.add_argument("--db", default=str(paths.OBSERVATUM_DB), help=argparse.SUPPRESS)
    ap.add_argument("--backup-dir", default=BACKUP_DIR, help=argparse.SUPPRESS)
    a = ap.parse_args(argv)
    sheet = a.sheet or newest_sheet()
    print(f"\nPhotos -- apply review sheet\n  sheet: {sheet}\n  library: {a.root}\n  observatum.db: {a.db}")
    if not sheet or not os.path.isfile(sheet):
        print("  no review sheet found (give it with --sheet). Nothing has been changed.")
        return 1
    ro = sqlite3.connect(f"file:{a.db}?mode=ro", uri=True)
    try:
        if not pl.has_table(ro):
            print("  the photo_links table is not there yet. Nothing has been changed.")
            return 1
        linked = {pl.norm(r[0]) for r in ro.execute(
            f"SELECT photo_path FROM {pl.TABLE} WHERE specimen_id IS NOT NULL")}
        before = ro.execute(f"SELECT COUNT(*) FROM {pl.TABLE}").fetchone()[0]
    finally:
        ro.close()
    links, problems, blank = plan(read_sheet(sheet), a.root, linked)
    print(f"  links to write: {len(links):,}   rows left blank: {blank:,}   problems: {len(problems):,}")
    for p in problems:
        print(f"    PROBLEM {p}")
    if not a.apply:
        for p, sid in links[:50]:
            print(f"    {p}  ->  specimen {sid}")
        if len(links) > 50:
            print(f"    ... and {len(links) - 50:,} more")
        print("\nDRY RUN. Nothing has been changed. Run with --apply to write the links.")
        return 0
    if not links:
        print("  nothing to link. Nothing has been changed.")
        return 0 if not problems else 1
    dest = backup(a.db, a.backup_dir)
    print(f"  backed up to {dest}")
    conn = sqlite3.connect(a.db)
    try:
        with conn:
            made = sum(pl.add_link(conn, p, specimen_id=sid, method=pl.BY_HAND) for p, sid in links)
        after = conn.execute(f"SELECT COUNT(*) FROM {pl.TABLE}").fetchone()[0]
    finally:
        conn.close()
    ok = after == before + made == before + len(links)
    print(f"  links written: {made:,}; table rows {before:,} -> {after:,}")
    print("  done." if ok else "  NOT as expected -- check the database.")
    if problems:
        print(f"  {len(problems):,} rows had problems and were skipped (listed above).")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
