"""Can uksi.db be rebuilt from the NHM 'UKSI Simplified Copy' (July 2025)?   READ ONLY

Derives every taxa field from the spreadsheet (kingdom..genus and superfamily by walking
PARENT_TVK chains) and compares with the current uksi.db for taxa present in both.
Also shows sort_code examples, NAMES sheet structure (statuses, languages) and what
is in one copy but not the other.

  python scripts\\compare_uksi_release.py "<path to the .xlsx>"
"""
import os, sqlite3, sys, time
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import paths
from openpyxl import load_workbook

if len(sys.argv) < 2 or not os.path.isfile(sys.argv[1]):
    sys.exit('  usage: python scripts\\compare_uksi_release.py "<path to the .xlsx>"')
t0 = time.time()
wb = load_workbook(sys.argv[1], read_only=True, data_only=True)


def sheet(name):
    it = wb[name].iter_rows(values_only=True)
    h = [str(x or "").strip() for x in next(it)]
    for r in it:
        yield {k: ("" if v is None else str(v).strip()) for k, v in zip(h, r)}


T = {}
for r in sheet("TAXA"):                       # a TVK may have several rows: never let a redundant row win
    k = r.get("TAXON_VERSION_KEY")
    if k and (k not in T or (T[k].get("REDUNDANT_FLAG") == "Y" and r.get("REDUNDANT_FLAG") != "Y")):
        T[k] = r
print("UKSI release comparison   READ ONLY")
print("=" * 88)
print(f"  TAXA sheet: {len(T)} rows   ({time.time() - t0:.0f}s)")
print(f"  RANK values (top): {Counter(r.get('RANK', '') for r in T.values()).most_common(14)}")
print(f"  REDUNDANT_FLAG: {Counter(r.get('REDUNDANT_FLAG', '') for r in T.values()).most_common(4)}")
ex = next(iter(T.values()))
print(f"  example row: { {k: ex[k][:40] for k in ex} }")

WANT = {"Kingdom": "kingdom", "Phylum": "phylum", "Class": "class", "Order": "order",
        "Superfamily": "superfamily", "Family": "family", "Genus": "genus"}


def derive(tvk):
    out, seen, cur = {}, set(), tvk                      # include the taxon's own rank (a genus is its own genus)
    while cur and cur in T and cur not in seen:
        seen.add(cur); p = T[cur]
        k = WANT.get(p.get("RANK", ""))
        if k and k not in out:
            out[k] = p.get("TAXON_NAME", "")
        cur = p.get("PARENT_TVK")
    return out


u = sqlite3.connect(f"file:{paths.UKSI_DB}?mode=ro", uri=True)
cols = [r[1] for r in u.execute("PRAGMA table_info(taxa)")]
ours = {r[0]: dict(zip(cols, r)) for r in u.execute(f"SELECT {', '.join('[' + c + ']' for c in cols)} FROM taxa")}
both = [t for t in ours if t in T]
print(f"\n  our taxa {len(ours)}   in both {len(both)}   only ours {len(ours) - len(both)}   "
      f"only July (not redundant) {sum(1 for t, r in T.items() if t not in ours and r.get('REDUNDANT_FLAG', '') in ('', 'N', '0'))}")

agree, diff_ex = Counter(), {}
GROUPS = ("Coleoptera", "Diptera", "Hymenoptera", "Hemiptera", "Lepidoptera", "Araneae", "Odonata", "Orthoptera")
gagree, gn = Counter(), Counter()
for t in both:
    o, j = ours[t], T[t]
    d = derive(t)
    checks = {"scientific_name": (o["scientific_name"], j.get("TAXON_NAME")),
              "rank": ((o["rank"] or "").lower(), j.get("RANK", "").lower()),
              "parent_tvk": (o["parent_tvk"], j.get("PARENT_TVK"))}
    for k in ("kingdom", "phylum", "class", "order", "superfamily", "family", "genus"):
        if k in o:
            checks[k] = (o[k] or "", d.get(k, ""))
    grp = o.get("order") if o.get("order") in GROUPS else ("Animalia (other)" if o.get("kingdom") == "Animalia" else (o.get("kingdom") or "?"))
    for k, (a, b) in checks.items():
        ok = (a or "") == (b or "")
        if ok:
            agree[k] += 1
        else:
            diff_ex.setdefault(k, []).append((o["scientific_name"], a, b, t))
        if k in ("order", "family", "genus", "superfamily", "scientific_name"):
            gn[(grp, k)] += 1; gagree[(grp, k)] += ok
print("\n  FIELD AGREEMENT for taxa in both copies (ours vs derived from July):")
for k in ("scientific_name", "rank", "parent_tvk", "kingdom", "phylum", "class", "order", "superfamily", "family", "genus"):
    n = agree[k] + len(diff_ex.get(k, []))
    if n:
        print(f"    {k:<16} {100 * agree[k] / n:6.2f}%   differ {len(diff_ex.get(k, [])):>6}"
              + (f"   e.g. {diff_ex[k][0][0][:28]}: ours {str(diff_ex[k][0][1])[:22]!r} July {str(diff_ex[k][0][2])[:22]!r}" if diff_ex.get(k) else ""))

print("\n  AGREEMENT BY GROUP (name / superfamily / family / genus):")
for g in list(GROUPS) + ["Animalia (other)", "Plantae", "Fungi", "Chromista", "Protozoa", "?"]:
    if gn[(g, "family")]:
        f = lambda k: f"{100 * gagree[(g, k)] / gn[(g, k)]:5.1f}%" if gn[(g, k)] else "   -  "
        print(f"    {g:<18} n={gn[(g, 'family')]:>6}   name {f('scientific_name')}  superfamily {f('superfamily')}  family {f('family')}  genus {f('genus')}")
fam_beetle = [x for x in diff_ex.get("family", []) if ours[x[3]].get("order") == "Coleoptera"][:8]
if fam_beetle:
    print("    beetle family differences, e.g.:", [(a[:24], b, c) for a, b, c, _ in fam_beetle])

red = {t for t, r in T.items() if r.get("REDUNDANT_FLAG") == "Y"}
our_red = [t for t in ours if t in red]
print(f"\n  REDUNDANT in July but in our taxa: {len(our_red)}   by rank: {Counter((ours[t]['rank'] or '') for t in our_red).most_common(6)}")
o_db = sqlite3.connect(f"file:{paths.OBSERVATUM_DB}?mode=ro", uri=True)
for tbl in ("observations", "specimens", "recording_scheme"):
    try:
        n_r = o_db.execute(f"SELECT COUNT(1), COUNT(DISTINCT species_tvk) FROM {tbl} WHERE species_tvk IN ({','.join('?' * len(our_red))})", our_red).fetchone() if our_red else (0, 0)
        print(f"    {tbl}: {n_r[0]} records on {n_r[1]} redundant TVKs")
    except sqlite3.Error as e:
        print(f"    {tbl}: (not checked: {e})")
c_db = sqlite3.connect(f"file:{paths.CODEX_DB}?mode=ro", uri=True)
n_c = c_db.execute(f"SELECT COUNT(DISTINCT tvk) FROM status_summary WHERE tvk IN ({','.join('?' * len(our_red))})", our_red).fetchone()[0] if our_red else 0
print(f"    codex status_summary: {n_c} species on redundant TVKs")

print("\n  SORT fields, ours beside July (5 species):")
sp = [t for t in both if (ours[t]["rank"] or "").lower() == "species"][:5]
for t in sp:
    print(f"    {ours[t]['scientific_name'][:30]:30} sort_code={ours[t]['sort_code']!r:24} sort_order={ours[t]['sort_order']!r:12} "
          f"July SORT_ORDER={T[t].get('SORT_ORDER', '')!r} LINEAGE={T[t].get('LINEAGE', '')[:24]!r}")

st, lang, form, ntype = Counter(), Counter(), Counter(), Counter()
en_pref = Counter()
n = 0
for r in sheet("NAMES"):
    n += 1
    st[r.get("NAME_STATUS", "")] += 1; lang[r.get("LANGUAGE", "")] += 1
    form[r.get("NAME_FORM", "")] += 1; ntype[r.get("NAME_TYPE", "")] += 1
    if r.get("LANGUAGE", "") == "en":
        en_pref[r.get("NAME_STATUS", "")] += 1
print(f"\n  NAMES sheet: {n} rows")
print(f"    NAME_STATUS: {st.most_common(6)}\n    LANGUAGE: {lang.most_common(6)}")
print(f"    NAME_FORM: {form.most_common(6)}\n    NAME_TYPE: {ntype.most_common(6)}\n    English names by status: {en_pref.most_common(5)}")
c_syn = u.execute("SELECT COUNT(1) FROM synonyms").fetchone()[0]
c_cn = u.execute("SELECT COUNT(1) FROM common_names").fetchone()[0]
print(f"    ours: synonyms {c_syn}, common_names {c_cn}")
print(f"\n  ({time.time() - t0:.0f}s)\nREAD ONLY -- nothing has been changed.")
