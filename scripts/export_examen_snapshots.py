"""Export the frozen assessments held in examen.db to Excel workbooks (backlog E13).
READ ONLY on examen.db.

Examen's Freeze / Assessment Archive feature was retired on 9 Oct 2026: the frozen
record of an assessment is now the report that was issued. This script gives anything
that was ever frozen a permanent home as plain files, before the code that read it
goes to _archive.

One workbook per stored assessment (a row of site_snapshots), with sheets:
  About      site, project, client, year, date frozen, mode, source, archive note
  Assessment every stored field of the site_snapshots row
  Species    its rows in snapshot_species
  Log        its rows in snapshot_log
  (and one sheet for any other table carrying a snapshot_id column)
plus index.csv listing every workbook. Rows that point at no assessment, and tables
that belong to no assessment, are reported and written to examen_other_tables.xlsx.

Nothing is ever overwritten: an existing file gets a _2, _3 ... suffix instead.

  python scripts\\export_examen_snapshots.py                 dry run: list what it would write
  python scripts\\export_examen_snapshots.py --write         write to _archive\\examen_snapshots_<YYYYMMDD>\\
  python scripts\\export_examen_snapshots.py --out <folder>  another folder (still needs --write)
  python scripts\\export_examen_snapshots.py --db <path>     read another copy of examen.db
"""
import csv
import os
import re
import sys
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [ROOT, os.path.join(ROOT, "Observatum")]

import paths  # noqa: E402
from shared.db_open import connect_ro  # noqa: E402

MAIN = "site_snapshots"
CHILD_SHEETS = {"snapshot_species": "Species", "snapshot_log": "Log"}
ARCHIVE_NOTE = ("This is an archived frozen assessment, exported from examen.db when "
                "Examen's Freeze / Assessment Archive feature was retired (backlog E13). "
                "The figures are as they were when frozen, against the Codex and Pantheon "
                "of that date; they are not recalculated.")


# ---- reading ---------------------------------------------------------------

def read_db(db_path):
    """Everything in the database: {table: (columns, rows)} for every table."""
    conn = connect_ro(db_path)
    try:
        names = [r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        tables = {}
        for t in names:
            cols = [r[1] for r in conn.execute(f'PRAGMA table_info("{t}")')]
            rows = conn.execute(f'SELECT * FROM "{t}" ORDER BY rowid').fetchall()
            tables[t] = (cols, [tuple(r) for r in rows])
        return tables
    finally:
        conn.close()


def _as_dicts(table):
    cols, rows = table
    return [dict(zip(cols, r)) for r in rows]


def collect(tables):
    """Group the tables into assessments.

    Returns (assessments, leftovers). Each assessment is
    {"row": dict, "children": {table: (cols, rows)}}; leftovers is
    {label: (cols, rows)} for orphan child rows and tables tied to no assessment.
    """
    main = _as_dicts(tables[MAIN]) if MAIN in tables else []
    ids = {r["id"] for r in main}
    order = list(CHILD_SHEETS)     # Species, then Log, then anything else by name
    child_tables = {t: tables[t] for t in sorted(
        (t for t, v in tables.items() if t != MAIN and "snapshot_id" in v[0]),
        key=lambda t: (order.index(t) if t in order else len(order), t))}
    assessments = []
    for row in main:
        children = {}
        for t, (cols, rows) in child_tables.items():
            i = cols.index("snapshot_id")
            children[t] = (cols, [r for r in rows if r[i] == row["id"]])
        assessments.append({"row": row, "children": children})

    leftovers = {}
    for t, (cols, rows) in child_tables.items():
        i = cols.index("snapshot_id")
        orphans = [r for r in rows if r[i] not in ids]
        if orphans:
            leftovers[f"{t} (no assessment)"] = (cols, orphans)
    for t, v in tables.items():
        if t == MAIN or t in child_tables or t.startswith("sqlite_"):
            continue
        if v[1]:
            leftovers[t] = v
    return assessments, leftovers


# ---- naming ----------------------------------------------------------------

def _slug(text, limit=40):
    s = re.sub(r"[^A-Za-z0-9]+", "_", str(text or "")).strip("_")
    return s[:limit].rstrip("_")


def workbook_name(row):
    parts = [f"{int(row.get('id') or 0):03d}", _slug(row.get("site_name")),
             _slug(row.get("project_name")), str(row.get("survey_year") or ""),
             _slug(row.get("analysis_mode"), 20)]
    return "_".join(p for p in parts if p) + ".xlsx"


def free_path(folder, name, taken=()):
    """`folder/name`, or name_2, name_3 ... if that exists on disk or is already planned."""
    base, ext = os.path.splitext(name)
    n, candidate = 1, name
    while os.path.exists(os.path.join(folder, candidate)) or candidate in taken:
        n += 1
        candidate = f"{base}_{n}{ext}"
    return os.path.join(folder, candidate)


def plan(assessments, leftovers, out_dir):
    """[(path, kind, payload)] -- every file that would be written, none overwriting."""
    out, taken = [], set()
    for a in assessments:
        p = free_path(out_dir, workbook_name(a["row"]), taken)
        taken.add(os.path.basename(p))
        out.append((p, "assessment", a))
    if leftovers:
        p = free_path(out_dir, "examen_other_tables.xlsx", taken)
        taken.add(os.path.basename(p))
        out.append((p, "other", leftovers))
    p = free_path(out_dir, "index.csv", taken)
    out.append((p, "index", None))
    return out


# ---- writing ---------------------------------------------------------------

def _sheet(wb, title, cols, rows):
    from openpyxl.styles import Font
    ws = wb.create_sheet(title[:31])
    ws.append(list(cols))
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    for i, col in enumerate(cols, 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width + 2, 8), 60)
    return ws


def _about(wb, pairs):
    from openpyxl.styles import Alignment, Font
    ws = wb.active
    ws.title = "About"
    for k, v in pairs:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 90


def write_assessment(path, a, db_path, exported):
    from openpyxl import Workbook
    r = a["row"]
    wb = Workbook()
    _about(wb, [
        ("Site", r.get("site_name") or ""),
        ("Project", r.get("project_name") or ""),
        ("Client", r.get("client") or ""),
        ("Survey year", r.get("survey_year") or ""),
        ("Date frozen", r.get("frozen_date") or ""),
        ("Analysis mode", r.get("analysis_mode") or ""),
        ("Codex version", r.get("codex_version") or ""),
        ("Pantheon version", r.get("pantheon_version") or ""),
        ("Source", f"{db_path}  (table {MAIN}, id {r.get('id')})"),
        ("Exported", exported),
        ("Note", ARCHIVE_NOTE),
    ])
    _sheet(wb, "Assessment", ["Field", "Value"], list(r.items()))
    for t, (cols, rows) in a["children"].items():
        _sheet(wb, CHILD_SHEETS.get(t, t), cols, rows)
    wb.save(path)


def write_other(path, leftovers, db_path, exported):
    from openpyxl import Workbook
    wb = Workbook()
    _about(wb, [("Source", str(db_path)), ("Exported", exported),
                ("Note", "Rows in examen.db that belong to no stored assessment. " + ARCHIVE_NOTE)])
    for label, (cols, rows) in leftovers.items():
        _sheet(wb, _slug(label, 31) or "table", cols, rows)
    wb.save(path)


def write_index(path, planned):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["file", "id", "site", "project", "client", "survey_year",
                    "analysis_mode", "frozen_date", "species_rows", "sqi"])
        for p, kind, payload in planned:
            if kind == "assessment":
                r = payload["row"]
                sp = payload["children"].get("snapshot_species", ((), []))[1]
                w.writerow([os.path.basename(p), r.get("id"), r.get("site_name"),
                            r.get("project_name"), r.get("client"), r.get("survey_year"),
                            r.get("analysis_mode"), r.get("frozen_date"), len(sp), r.get("sqi")])
            elif kind == "other":
                w.writerow([os.path.basename(p), "", "(rows tied to no assessment)",
                            "", "", "", "", "", "", ""])


def export(db_path, out_dir, write=False):
    """Read, plan, print, and (with write=True) write. Returns the plan."""
    tables = read_db(db_path)
    print(f"examen.db: {db_path}  (opened READ ONLY)")
    for t, (cols, rows) in tables.items():
        print(f"  {t:22} {len(rows):>6} rows   {', '.join(cols)}")
    assessments, leftovers = collect(tables)
    print(f"\n{len(assessments)} stored assessment(s).")
    for a in assessments:
        r = a["row"]
        sp = a["children"].get("snapshot_species", ((), []))[1]
        print(f"  id {r.get('id')}: {r.get('site_name')} | {r.get('project_name')} | "
              f"{r.get('survey_year')} | {r.get('analysis_mode')} | frozen {r.get('frozen_date')} | "
              f"{len(sp)} species | SQI {r.get('sqi')}")
    for label, (_, rows) in leftovers.items():
        print(f"  NOT IN ANY ASSESSMENT: {label}: {len(rows)} row(s)")

    if not assessments and not leftovers:
        print("\nexamen.db holds no assessments -- there is nothing to export.")
        print("Nothing has been changed.")
        return []

    planned = plan(assessments, leftovers, out_dir)
    print(f"\n{'Writing' if write else 'Would write'} to {out_dir}:")
    for p, _, _ in planned:
        print(f"  {os.path.basename(p)}")
    if not write:
        print("\nDRY RUN -- nothing written. Re-run with --write.  Nothing has been changed.")
        return planned

    os.makedirs(out_dir, exist_ok=True)
    exported = datetime.now().strftime("%d/%m/%Y %H:%M")
    for p, kind, payload in planned:
        if kind == "assessment":
            write_assessment(p, payload, db_path, exported)
        elif kind == "other":
            write_other(p, payload, db_path, exported)
        else:
            write_index(p, planned)
    print(f"\n{len(planned)} file(s) written. examen.db was opened read-only and has not been changed.")
    return planned


def main(argv):
    def arg(flag):
        return argv[argv.index(flag) + 1] if flag in argv else None
    db_path = arg("--db") or str(paths.EXAMEN_DB)
    if not os.path.exists(db_path):
        print(f"{db_path} does not exist -- nothing to export. Nothing has been changed.")
        return 0
    out_dir = arg("--out") or os.path.join(
        ROOT, "_archive", f"examen_snapshots_{datetime.now():%Y%m%d}")
    export(db_path, out_dir, write="--write" in argv)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
