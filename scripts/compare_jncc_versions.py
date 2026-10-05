"""Compare the 2023 and 2026 JNCC designation spreadsheets before a Codex rebuild.  READ ONLY.

  1. sheet names, header rows and row counts -- will build_codex_db.py read the new file?
  2. designation codes in 2026 that the build does not route (DESIG_TO_TRACK)
  3. sources new in 2026 (the reviews JNCC has added) and sources dropped
  4. how paths.py and the build find the spreadsheet (folder and file name)

Run:  python scripts\\compare_jncc_versions.py
"""
import glob, os, re, sys
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [ROOT, os.path.join(ROOT, "scripts")]
import openpyxl

OLD = glob.glob(os.path.join(ROOT, "data", "conservation-designations-20231206", "*esignations*20231206*.xlsx"))
NEW = glob.glob(os.path.join(ROOT, "data", "conservation-designations-20260609", "*esignations*20260609*.xlsx"))
if not OLD or not NEW:
    sys.exit(f"  x spreadsheet not found -- old: {OLD}  new: {NEW}")


def read(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        sheets = {ws.title: ws for ws in wb.worksheets}
        info = {n: ws.max_row for n, ws in sheets.items()}
        main = "Master List" if "Master List" in sheets else max(info, key=info.get)
        rows = list(sheets[main].iter_rows(values_only=True))
    finally:
        wb.close()
    hi = next(i for i, r in enumerate(rows[:10]) if sum(1 for v in r if isinstance(v, str)) >= 5)
    hdr = [str(h).strip() if h is not None else "" for h in rows[hi]]
    data = [dict(zip(hdr, r)) for r in rows[hi + 1:] if any(v is not None for v in r)]
    return info, main, hdr, data


print("JNCC spreadsheet: 2023 vs 2026   READ ONLY")
print("=" * 90)
o_info, o_main, o_hdr, o = read(OLD[0])
n_info, n_main, n_hdr, n = read(NEW[0])
print(f"  2023: {os.path.basename(OLD[0])}  sheets {list(o_info)}  main '{o_main}' {len(o)} rows")
print(f"  2026: {os.path.basename(NEW[0])}  sheets {list(n_info)}  main '{n_main}' {len(n)} rows")
print(f"  columns only in 2023: {[h for h in o_hdr if h not in n_hdr]}")
print(f"  columns only in 2026: {[h for h in n_hdr if h not in o_hdr]}")

ab = next((h for h in n_hdr if "abbreviation" in h.lower()), None)
src = next((h for h in n_hdr if h.strip().lower() == "source"), None)
if ab:
    try:
        import build_codex_db as b
        routed = set(getattr(b, "DESIG_TO_TRACK", {}))
    except Exception as e:
        routed = set(); print(f"  (could not import DESIG_TO_TRACK: {e})")
    oc = Counter(str(r.get(ab) or "") for r in o); nc = Counter(str(r.get(ab) or "") for r in n)
    new_codes = {c: k for c, k in nc.items() if c not in oc}
    print(f"\n  designation codes new in 2026: {len(new_codes)}")
    for c, k in sorted(new_codes.items(), key=lambda x: -x[1])[:25]:
        print(f"    {k:>6}  {c:<45} {'routed' if c in routed else 'NOT ROUTED' if routed else ''}")
    if routed:
        un = {c: k for c, k in nc.items() if c and c not in routed}
        print(f"  all unrouted codes in 2026: {len(un)} codes, {sum(un.values())} rows (2023 build: 14 codes, 121 rows)")
if src:
    os_ = Counter(str(r.get(src) or "") for r in o); ns = Counter(str(r.get(src) or "") for r in n)
    print(f"\n  sources new in 2026 ({sum(1 for s in ns if s not in os_)}):")
    for s, k in sorted(((s, k) for s, k in ns.items() if s not in os_), key=lambda x: -x[1])[:30]:
        print(f"    {k:>6}  {s[:95]}")
    print(f"  sources dropped since 2023 ({sum(1 for s in os_ if s not in ns)}):")
    for s, k in sorted(((s, k) for s, k in os_.items() if s not in ns), key=lambda x: -x[1])[:15]:
        print(f"    {k:>6}  {s[:95]}")

print("\n  how the build finds the spreadsheet:")
for f in ("paths.py", os.path.join("scripts", "build_codex_db.py")):
    for i, line in enumerate(open(os.path.join(ROOT, f), encoding="utf-8-sig"), 1):
        if re.search(r"JNCC|20231206|designations.*xlsx|Master List", line):
            print(f"    {f}:{i}: {line.rstrip()[:100]}")
print("\nREAD ONLY -- nothing has been changed.")
