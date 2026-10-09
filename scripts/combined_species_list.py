"""Combined species list -- Birmingham - Wheels Park, Wil + J. Moore.

Reads Wil's records from staging job 6 and J. Moore's from his spreadsheet,
resolves every name against UKSI (TVK, then name, then synonym -- the same
resolve() the review importer uses), and writes one spreadsheet:

  Combined    one row per species, taxonomic order: name, common name, order,
              family, TVK, records by Wil, records by J. Moore, recorded by
  Unmatched   names that did not resolve, and whose list they came from

READ ONLY on the databases. Writes only the output spreadsheet.

Run:
  python scripts\\combined_species_list.py "data\\contributed\\Birmingham_Wheels.xlsx"
"""
import os, sqlite3, sys
from collections import defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import paths
from import_status_review import resolve      # import the rule, don't restate it

import openpyxl
from openpyxl.styles import Font, PatternFill

JOB_ID = 6
SITE = "Birmingham - Wheels Park"
OUT = os.path.join(ROOT, "Birmingham_Wheels_Park_combined_species.xlsx")

if len(sys.argv) < 2 or not os.path.exists(sys.argv[1]):
    sys.exit("Give the path to J. Moore's spreadsheet.")
moore_path = sys.argv[1]

u = sqlite3.connect(f"file:{paths.UKSI_DB}?mode=ro", uri=True)
o = sqlite3.connect(f"file:{paths.OBSERVATUM_DB}?mode=ro", uri=True)

species = {}                         # tvk -> info
counts = defaultdict(lambda: {"wil": 0, "moore": 0})
unmatched = []


def add(name, tvk, who):
    r = resolve(u, (name or "").strip(), (tvk or "").strip() or None)
    if not r:
        if (name, who) not in unmatched:
            unmatched.append((name, who))
        return
    t, sci, family, order, how = r
    if t not in species:
        srt = u.execute("SELECT sort_order, sort_code FROM taxa WHERE tvk=?", (t,)).fetchone()
        cn = u.execute("SELECT common_name FROM common_names WHERE tvk=? AND preferred=1",
                       (t,)).fetchone()
        species[t] = {"name": sci, "family": family or "", "order": order or "",
                      "common": cn[0] if cn else "",
                      "sort": (srt[0] or "") if srt else "", "code": (srt[1] or 0) if srt else 0}
    counts[t][who] += 1


# ---------------------------------------------------------------- Wil
n_w = 0
for name, tvk in o.execute("""SELECT species_name, species_tvk FROM entry_staging
                              WHERE job_id=? AND species_name IS NOT NULL
                                AND trim(species_name)!=''""", (JOB_ID,)):
    add(name, tvk, "wil")
    n_w += 1

# ---------------------------------------------------------------- J. Moore
wb = openpyxl.load_workbook(moore_path, read_only=True, data_only=True)
ws = wb["Sample Data"] if "Sample Data" in wb.sheetnames else wb.worksheets[0]
rows = ws.iter_rows(values_only=True)
header = [str(h or "").strip() for h in next(rows)]
col = next((i for i, h in enumerate(header) if h.lower().startswith("species")), 0)
n_m = 0
for r in rows:
    name = r[col] if col < len(r) else None
    if name and str(name).strip():
        add(str(name).strip(), None, "moore")
        n_m += 1

# ---------------------------------------------------------------- write
out = openpyxl.Workbook()
sh = out.active
sh.title = "Combined"
head = ["Species", "Common name", "Order", "Family", "TVK",
        "Records (W. Heeney)", "Records (J. Moore)", "Recorded by"]
sh.append([f"{SITE} -- combined species list"])
sh["A1"].font = Font(bold=True, size=13)
sh.append([f"W.J. Heeney: {n_w} records.  J. Moore: {n_m} records.  "
           f"{len(species)} species.  Taxonomy: UKSI (NHM), CC BY 4.0."])
sh.append([])
sh.append(head)
for c in sh[4]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4A7C59")

both = only_w = only_m = 0
for t, s in sorted(species.items(), key=lambda kv: (kv[1]["sort"] or "~", kv[1]["code"], kv[1]["name"])):
    w, m = counts[t]["wil"], counts[t]["moore"]
    by = "Both" if w and m else ("W. Heeney" if w else "J. Moore")
    both += bool(w and m); only_w += bool(w and not m); only_m += bool(m and not w)
    sh.append([s["name"], s["common"], s["order"], s["family"], t, w or None, m or None, by])
    sh.cell(sh.max_row, 1).font = Font(italic=True)

for letter, width in zip("ABCDEFGH", (34, 28, 16, 20, 20, 12, 12, 12)):
    sh.column_dimensions[letter].width = width
sh.freeze_panes = "A5"

un = out.create_sheet("Unmatched")
un.append(["Name as written", "From"])
for c in un[1]:
    c.font = Font(bold=True)
for name, who in unmatched:
    un.append([name, "W. Heeney" if who == "wil" else "J. Moore"])
un.column_dimensions["A"].width = 40

out.save(OUT)

print("Combined species list -- Birmingham - Wheels Park")
print("=" * 60)
print(f"  W.J. Heeney: {n_w} records    J. Moore: {n_m} records")
print(f"  species: {len(species)}   both: {both}   W. Heeney only: {only_w}   J. Moore only: {only_m}")
print(f"  unmatched names: {len(unmatched)}")
for name, who in unmatched:
    print(f"    {name!r}  ({'W. Heeney' if who == 'wil' else 'J. Moore'})")
print(f"\n  written: {OUT}")
