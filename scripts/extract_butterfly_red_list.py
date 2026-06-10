"""
Extract butterfly Red List 2022 from Butterfly Conservation Excel file.

Prerequisites:
    pip install openpyxl --break-system-packages

Usage:
    1. Go to: https://butterfly-conservation.org/red-list-of-butterflies-in-great-britain
    2. Find and download the Excel file with qualifying criteria
       (the page says "More detailed information on the qualifying criteria
        for each butterfly species can be found here (downloads as an Excel document)")
    3. Save it, e.g.: data/reviews/butterfly-red-list-2022.xlsx
    
    4. Run:
       python scripts/extract_butterfly_red_list.py data/reviews/butterfly-red-list-2022.xlsx
    
    5. Output: data/reviews/butterfly_red_list_2022.csv
    
    6. Import into Codex (supersedes the 2010 butterfly Red List already in JNCC):
       python scripts/import_codex_review.py data/reviews/butterfly_red_list_2022.csv ^
           --name "GB Butterfly Red List 2022" ^
           --author "Fox et al., 2022" ^
           --group "Lepidoptera" ^
           --track gb_red_list ^
           --date 2022-05-25 ^
           --dry-run

       Then remove --dry-run to import for real.
       Then run: python scripts/seed_codex.py   (for SQS gap-fill)
"""

import sys
import csv
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Run:")
    print("  pip install openpyxl --break-system-packages")
    sys.exit(1)


# Valid IUCN Red List categories
VALID_CATEGORIES = {
    "RE", "CR", "CR(PE)", "EN", "VU", "NT", "DD", "LC", "NA", "NE",
    "Regionally Extinct", "Critically Endangered", "Endangered",
    "Vulnerable", "Near Threatened", "Data Deficient", "Least Concern",
    "Not Applicable", "Not Evaluated",
}

# Map full names to abbreviations
CATEGORY_MAP = {
    "Regionally Extinct": "RE",
    "Critically Endangered": "CR",
    "Endangered": "EN",
    "Vulnerable": "VU",
    "Near Threatened": "NT",
    "Data Deficient": "DD",
    "Least Concern": "LC",
    "Not Applicable": "NA",
    "Not Evaluated": "NE",
}


def extract_butterfly_red_list(xlsx_path: str, output_dir: str = None):
    """Extract butterfly Red List data from the Excel file."""
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}")
        return None

    if output_dir:
        out_dir = Path(output_dir)
    else:
        out_dir = xlsx_path.parent

    print(f"Reading: {xlsx_path}")
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)

    # Try to find the right sheet
    print(f"  Sheets: {wb.sheetnames}")
    ws = None
    for name in wb.sheetnames:
        if "red" in name.lower() or "species" in name.lower() or "list" in name.lower():
            ws = wb[name]
            print(f"  Using sheet: {name}")
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
        print(f"  Using first sheet: {wb.sheetnames[0]}")

    # Find header row and column indices
    headers = []
    header_row = None
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        cells = [str(c).strip() if c else "" for c in row]
        # Look for a row with "species" and some category-like header
        lower = [c.lower() for c in cells]
        if any("species" in c or "scientific" in c or "taxon" in c for c in lower):
            if any("red" in c or "category" in c or "status" in c or "iucn" in c for c in lower):
                headers = cells
                header_row = row_idx
                break
        # Also check if first row has headers
        if row_idx <= 3 and any("species" in c or "scientific" in c for c in lower):
            headers = cells
            header_row = row_idx

    if not headers:
        # Fallback: treat row 1 as headers
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(c).strip() if c else f"col_{i}" for i, c in enumerate(row)]
            header_row = 1
            break

    print(f"  Header row: {header_row}")
    print(f"  Headers: {headers[:8]}...")

    # Find key columns
    name_col = None
    cat_col = None
    for i, h in enumerate(headers):
        hl = h.lower()
        if name_col is None and ("scientific" in hl or "species" in hl or "taxon" in hl):
            if "common" not in hl:
                name_col = i
        if cat_col is None and ("red list" in hl or "category" in hl or "2022" in hl or "iucn" in hl or "status" in hl):
            if "2010" not in hl and "previous" not in hl and "old" not in hl:
                cat_col = i

    if name_col is None:
        print(f"  ERROR: Could not find species name column")
        print(f"  Headers found: {headers}")
        return None

    if cat_col is None:
        # Try to find any column with valid categories
        print(f"  WARNING: Could not identify Red List column by header")
        print(f"  Scanning data for IUCN categories...")
        for i in range(len(headers)):
            if i == name_col:
                continue
            for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + 10, values_only=True):
                val = str(row[i]).strip() if row[i] else ""
                if val in VALID_CATEGORIES or val in CATEGORY_MAP:
                    cat_col = i
                    print(f"  Found categories in column {i}: '{headers[i]}'")
                    break
            if cat_col is not None:
                break

    if cat_col is None:
        print(f"  ERROR: Could not find Red List category column")
        return None

    print(f"  Species column: {name_col} ('{headers[name_col]}')")
    print(f"  Category column: {cat_col} ('{headers[cat_col]}')")

    # Extract data
    species_data = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        name = str(row[name_col]).strip() if row[name_col] else ""
        cat = str(row[cat_col]).strip() if row[cat_col] else ""

        if not name or not cat:
            continue
        if name.lower() in ("none", "nan", ""):
            continue

        # Normalise category
        if cat in CATEGORY_MAP:
            cat = CATEGORY_MAP[cat]
        cat = cat.upper().strip()

        if cat not in {"RE", "CR", "EN", "VU", "NT", "DD", "LC", "NA", "NE"}:
            print(f"    Skipping unknown category: {name} → '{cat}'")
            continue

        # Clean species name (remove any leading/trailing quotes or whitespace)
        name = name.strip("'\"").strip()

        species_data.append({"species_name": name, "status_value": cat})

    wb.close()

    if not species_data:
        print(f"\n  WARNING: No species data extracted.")
        print(f"  Check the Excel structure manually.")
        return None

    # Write CSV
    csv_path = out_dir / "butterfly_red_list_2022.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["species_name", "status_value"])
        for entry in species_data:
            writer.writerow([entry["species_name"], entry["status_value"]])

    # Summary
    categories = {}
    for e in species_data:
        c = e["status_value"]
        categories[c] = categories.get(c, 0) + 1

    print(f"\n  Output: {csv_path}")
    print(f"  Total species: {len(species_data)}")
    print(f"  Breakdown:")
    for cat in ["RE", "CR", "EN", "VU", "NT", "DD", "LC", "NA", "NE"]:
        if cat in categories:
            print(f"    {cat:4s}  {categories[cat]:>3}")

    return csv_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python scripts/extract_butterfly_red_list.py <path-to-xlsx>")
        print("")
        print("Download the Excel file from:")
        print("  https://butterfly-conservation.org/red-list-of-butterflies-in-great-britain")
        print("  (link: 'More detailed information on the qualifying criteria')")
        sys.exit(1)

    result = extract_butterfly_red_list(sys.argv[1])
    if result:
        print(f"\nDone. Import with:")
        print(f'  python scripts/import_codex_review.py "{result}" ^')
        print(f'      --name "GB Butterfly Red List 2022" ^')
        print(f'      --author "Fox et al., 2022" ^')
        print(f'      --group "Lepidoptera" ^')
        print(f'      --track gb_red_list ^')
        print(f'      --date 2022-05-25 ^')
        print(f'      --dry-run')
