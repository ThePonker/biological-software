"""
Extract sawfly Red List + Rarity data from all 3 phase Excel data tables.

Reads:
  data/reviews/Sawfly-Review-Phase-1-DATA-TABLE-FINAL.xlsx
  data/reviews/Sawfly-Review-Phase-2-Data-Table-FINAL.xlsx
  data/reviews/Sawfly-Review-Phase-3-Data-Table-FINAL.xlsx

Outputs:
  data/reviews/sawfly_all_red_list.csv    (IUCN Red List: gb_red_list track)
  data/reviews/sawfly_all_rarity.csv      (NR/NS: gb_rarity track)

Run: python scripts/extract_sawfly_all.py
"""

import csv
import openpyxl
from pathlib import Path

REVIEWS_DIR = Path("data/reviews")

VALID_RED_LIST = {"RE", "CR", "CR(PE)", "EN", "VU", "NT", "DD", "LC", "NA", "NE"}
VALID_RARITY = {"NR", "NS", "Nationally Rare", "Nationally Scarce"}
RARITY_MAP = {"Nationally Rare": "NR", "Nationally Scarce": "NS", "NR": "NR", "NS": "NS"}

FILES = [
    "Sawfly-Review-Phase-1-DATA-TABLE-FINAL.xlsx",
    "Sawfly-Review-Phase-2-Data-Table-FINAL.xlsx",
    "Sawfly-Review-Phase-3-Data-Table-FINAL.xlsx",
]


def find_col(headers, candidates):
    """Find column index matching any candidate (case-insensitive partial match)."""
    for i, h in enumerate(headers):
        hl = (h or "").lower().strip()
        for c in candidates:
            if c.lower() in hl:
                return i
    return None


def extract_phase(filepath):
    """Extract species + IUCN status + rarity from one Excel file."""
    wb = openpyxl.load_workbook(str(filepath), read_only=True)
    ws = wb[wb.sheetnames[0]]

    # Read all rows
    all_rows = []
    for row in ws.iter_rows(values_only=True):
        all_rows.append([str(c).strip() if c else "" for c in row])
    wb.close()

    if len(all_rows) < 3:
        print(f"  WARNING: Too few rows in {filepath.name}")
        return [], []

    # Find header row — look for "Taxon" or "Binomial" in first few rows
    header_row_idx = None
    headers = []
    for i in range(min(5, len(all_rows))):
        row = all_rows[i]
        lower = [c.lower() for c in row]
        if "taxon" in lower or "binomial" in lower:
            headers = row
            header_row_idx = i
            break

    if header_row_idx is None:
        # Phase 1 has headers in row 0 with "Binomial"
        # Try row 0
        headers = all_rows[0]
        header_row_idx = 0

    # Find columns
    name_col = find_col(headers, ["taxon", "binomial", "species"])
    iucn_col = find_col(headers, ["gb iucn", "iucn status", "red list status"])
    rarity_col = find_col(headers, ["gb rarity", "rarity status"])
    tvk_col = find_col(headers, ["taxon version key"])

    # Phase 1 has different structure — IUCN/rarity might be unnamed
    # Check if we found the key columns
    if name_col is None:
        # Try column 0
        name_col = 0

    print(f"  {filepath.name}:")
    print(f"    Headers at row {header_row_idx}: name_col={name_col}, iucn_col={iucn_col}, rarity_col={rarity_col}, tvk_col={tvk_col}")

    # If IUCN column not found by header, scan data for IUCN values
    if iucn_col is None:
        for col_idx in range(len(headers)):
            if col_idx == name_col:
                continue
            for row in all_rows[header_row_idx + 1:header_row_idx + 10]:
                if col_idx < len(row):
                    val = row[col_idx].upper().strip()
                    if val in VALID_RED_LIST:
                        iucn_col = col_idx
                        print(f"    Found IUCN values in column {col_idx}")
                        break
            if iucn_col is not None:
                break

    if rarity_col is None:
        for col_idx in range(len(headers)):
            if col_idx in (name_col, iucn_col):
                continue
            for row in all_rows[header_row_idx + 1:header_row_idx + 10]:
                if col_idx < len(row):
                    val = row[col_idx].strip()
                    if val in VALID_RARITY:
                        rarity_col = col_idx
                        print(f"    Found rarity values in column {col_idx}")
                        break
            if rarity_col is not None:
                break

    # Extract data
    red_list = []
    rarity = []
    data_start = header_row_idx + 1
    # Phase 1 might have a second header row
    if data_start < len(all_rows):
        first_data = all_rows[data_start]
        # Check if this row looks like another header
        if first_data[name_col] and first_data[name_col].lower() in ("", "none", "taxon", "binomial", "species"):
            data_start += 1

    for row in all_rows[data_start:]:
        if name_col >= len(row):
            continue
        name = row[name_col].strip()
        if not name or name.lower() in ("none", ""):
            continue
        # Must look like a binomial
        words = name.split()
        if len(words) < 2:
            continue
        if not words[0][0].isupper():
            continue

        species = f"{words[0]} {words[1]}"

        tvk = ""
        if tvk_col is not None and tvk_col < len(row):
            tvk = row[tvk_col].strip()

        # IUCN status
        if iucn_col is not None and iucn_col < len(row):
            val = row[iucn_col].strip().upper()
            # Handle variations
            val = val.replace("REGIONALLY EXTINCT", "RE")
            val = val.replace("CRITICALLY ENDANGERED", "CR")
            val = val.replace("ENDANGERED", "EN")
            val = val.replace("VULNERABLE", "VU")
            val = val.replace("NEAR THREATENED", "NT")
            val = val.replace("DATA DEFICIENT", "DD")
            val = val.replace("LEAST CONCERN", "LC")
            val = val.replace("NOT APPLICABLE", "NA")
            if val in VALID_RED_LIST:
                red_list.append((species, val, tvk))

        # Rarity
        if rarity_col is not None and rarity_col < len(row):
            val = row[rarity_col].strip()
            mapped = RARITY_MAP.get(val)
            if mapped:
                rarity.append((species, mapped, tvk))

    print(f"    Red List entries: {len(red_list)}")
    print(f"    Rarity entries:  {len(rarity)}")
    return red_list, rarity


def main():
    all_red_list = []
    all_rarity = []

    for fname in FILES:
        fpath = REVIEWS_DIR / fname
        if not fpath.exists():
            print(f"  SKIP: {fname} not found")
            continue
        rl, rar = extract_phase(fpath)
        all_red_list.extend(rl)
        all_rarity.extend(rar)

    # Deduplicate by species name (keep first)
    seen = set()
    unique_rl = []
    for name, cat, tvk in all_red_list:
        if name not in seen:
            seen.add(name)
            unique_rl.append((name, cat, tvk))

    seen = set()
    unique_rar = []
    for name, cat, tvk in all_rarity:
        if name not in seen:
            seen.add(name)
            unique_rar.append((name, cat, tvk))

    # Write Red List CSV
    rl_path = REVIEWS_DIR / "sawfly_all_red_list.csv"
    with open(rl_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["species_name", "status_value", "tvk"])
        for name, cat, tvk in unique_rl:
            writer.writerow([name, cat, tvk])

    # Write Rarity CSV
    rar_path = REVIEWS_DIR / "sawfly_all_rarity.csv"
    with open(rar_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["species_name", "status_value", "tvk"])
        for name, cat, tvk in unique_rar:
            writer.writerow([name, cat, tvk])

    # Summary
    rl_cats = {}
    for _, c, _ in unique_rl:
        rl_cats[c] = rl_cats.get(c, 0) + 1

    rar_cats = {}
    for _, c, _ in unique_rar:
        rar_cats[c] = rar_cats.get(c, 0) + 1

    print(f"\n{'='*50}")
    print(f"SAWFLY RED LIST — ALL PHASES")
    print(f"{'='*50}")
    print(f"  Total species (Red List): {len(unique_rl)}")
    for c in ["RE", "CR", "EN", "VU", "NT", "DD", "LC", "NA"]:
        if c in rl_cats:
            print(f"    {c:4s}  {rl_cats[c]:>4}")
    print(f"\n  Total species (Rarity): {len(unique_rar)}")
    for c in ["NR", "NS"]:
        if c in rar_cats:
            print(f"    {c:4s}  {rar_cats[c]:>4}")

    print(f"\n  Red List CSV: {rl_path}")
    print(f"  Rarity CSV:   {rar_path}")

    # Import commands
    print(f"\nImport commands:")
    print(f'  python scripts/import_codex_review.py "{rl_path}" ^')
    print(f'      --name "GB Sawfly Red List (All Phases)" ^')
    print(f'      --author "Musgrove, 2022-2024" ^')
    print(f'      --group "Hymenoptera: Symphyta" ^')
    print(f'      --track gb_red_list ^')
    print(f'      --date 2022-10-01 ^')
    print(f'      --dry-run')
    print()
    print(f'  python scripts/import_codex_review.py "{rar_path}" ^')
    print(f'      --name "GB Sawfly Rarity (All Phases)" ^')
    print(f'      --author "Musgrove, 2022-2024" ^')
    print(f'      --group "Hymenoptera: Symphyta" ^')
    print(f'      --track gb_rarity ^')
    print(f'      --date 2022-10-01 ^')
    print(f'      --dry-run')


if __name__ == "__main__":
    main()
