"""UKSI name mapping: every name's TVK -> its current recommended TVK.

Source: the NAMES sheet of the NHM 'UKSI Simplified Copy' spreadsheet (July 2025),
column TAXON_VERSION_KEY -> RECOMMENDED_TAXON_VERSION_KEY. This is the mapping the
UKSI Nameserver provides, and the one our December 2023 extraction lacks.

Writes table  name_map  in uksi.db:
    tvk TEXT PRIMARY KEY, recommended_tvk, name, recommended_name, name_status, deprecated, source

DRY RUN (default) reads the sheet and reports, for TVKs we cannot currently resolve
(JNCC designations in codex.db, your observations/specimens in observatum.db), how many
map onto a species that IS in our uksi.db taxa -- i.e. what the mapping alone fixes --
and how many point to a taxon newer than our copy (needs a UKSI rebuild).

  python scripts\\build_uksi_name_map.py "<path to the .xlsx>"
  python scripts\\build_uksi_name_map.py "<path to the .xlsx>" --apply     (backs up uksi.db first)
"""
import os, sqlite3, sys, time
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import paths

args = [a for a in sys.argv[1:] if not a.startswith("--")]
if not args or not os.path.isfile(args[0]):
    sys.exit('  usage: python scripts\\build_uksi_name_map.py "<path to the .xlsx>" [--apply]')
XLSX, APPLY = args[0], "--apply" in sys.argv

from openpyxl import load_workbook

print("UKSI name mapping -- " + ("APPLY" if APPLY else "DRY RUN"))
print("=" * 80)
t0 = time.time()
wb = load_workbook(XLSX, read_only=True, data_only=True)
if "NAMES" not in wb.sheetnames:
    sys.exit(f"  x no NAMES sheet (sheets: {wb.sheetnames})")
it = wb["NAMES"].iter_rows(values_only=True)
hdr = [str(h or "").strip() for h in next(it)]
need = ["TAXON_VERSION_KEY", "RECOMMENDED_TAXON_VERSION_KEY", "TAXON_NAME", "RECOMMENDED_SCIENTIFIC_NAME"]
miss = [n for n in need if n not in hdr]
if miss:
    sys.exit(f"  x NAMES sheet lacks {miss}")
ix = {h: i for i, h in enumerate(hdr)}
get = lambda r, k: (str(r[ix[k]]).strip() if k in ix and ix[k] < len(r) and r[ix[k]] is not None else "")
rows, seen = [], set()
for r in it:
    tvk, rec = get(r, "TAXON_VERSION_KEY"), get(r, "RECOMMENDED_TAXON_VERSION_KEY")
    if not tvk or not rec or tvk in seen:
        continue
    seen.add(tvk)
    rows.append((tvk, rec, get(r, "TAXON_NAME"), get(r, "RECOMMENDED_SCIENTIFIC_NAME"),
                 get(r, "NAME_STATUS"), get(r, "DEPRECATED_DATE")))
wb.close()
print(f"  NAMES sheet: {len(rows)} name TVKs read in {time.time() - t0:.0f}s")
m = {r[0]: r for r in rows}

u = sqlite3.connect(str(paths.UKSI_DB))
taxa = {t for (t,) in u.execute("SELECT tvk FROM taxa")}
print(f"  our uksi.db taxa: {len(taxa)}")


def report(label, tvks):
    tvks = {t for t in tvks if t and t not in taxa}
    fixed, newer, absent = [], [], []
    for t in tvks:
        r = m.get(t)
        if not r:
            absent.append(t)
        elif r[1] in taxa:
            fixed.append(r)
        else:
            newer.append(r)
    print(f"\n  {label}: {len(tvks)} TVKs not in our taxa")
    print(f"     map to a species we HAVE (the mapping fixes these): {len(fixed)}")
    print(f"     map to a taxon newer than our copy (needs UKSI rebuild): {len(newer)}")
    print(f"     not in the July 2025 release at all (newer still): {len(absent)}")
    for r in sorted(fixed, key=lambda r: r[2])[:6]:
        print(f"        e.g. {r[2][:34]:34} {r[0]} -> {r[3][:30]} {r[1]}")
    for r in sorted(newer, key=lambda r: r[3])[:4]:
        print(f"        newer: {r[2][:34]:34} -> {r[3][:30]} {r[1]}")
    return fixed, newer, absent


c = sqlite3.connect(f"file:{paths.CODEX_DB}?mode=ro", uri=True)
report("JNCC designations (codex.db)", {t for (t,) in c.execute("SELECT DISTINCT tvk FROM designations")})
o = sqlite3.connect(f"file:{paths.OBSERVATUM_DB}?mode=ro", uri=True)
rec = set()
for tbl in ("observations", "specimens"):
    try:
        rec |= {t for (t,) in o.execute(f"SELECT DISTINCT species_tvk FROM {tbl} WHERE COALESCE(species_tvk,'')!=''")}
    except sqlite3.Error:
        pass
report("your records (observations + specimens)", rec)

if not APPLY:
    sys.exit("\n  DRY RUN -- nothing changed. Re-run with --apply to write table name_map into uksi.db.\n")

bk = os.path.join(os.path.dirname(str(paths.UKSI_DB)), f"uksi_pre_namemap_{datetime.now():%Y%m%d_%H%M%S}.db")
src = sqlite3.connect(str(paths.UKSI_DB)); dst = sqlite3.connect(bk); src.backup(dst); dst.close(); src.close()
print(f"\n  backup: {bk}")
try:
    u.execute("DROP TABLE IF EXISTS name_map")
    u.execute("""CREATE TABLE name_map (tvk TEXT PRIMARY KEY, recommended_tvk TEXT NOT NULL, name TEXT,
                 recommended_name TEXT, name_status TEXT, deprecated TEXT, source TEXT)""")
    srcname = "UKSI Simplified Copy 20250703a (NHM Data Portal), NAMES sheet"
    u.executemany("INSERT INTO name_map VALUES (?,?,?,?,?,?,?)", [r + (srcname,) for r in rows])
    u.execute("CREATE INDEX IF NOT EXISTS idx_name_map_rec ON name_map(recommended_tvk)")
    u.commit()
except Exception as e:
    u.rollback(); sys.exit(f"  x FAILED, rolled back: {type(e).__name__}: {e}")
n = u.execute("SELECT COUNT(1) FROM name_map").fetchone()[0]
print(f"  name_map written: {n} rows   (to taxa we hold: "
      f"{u.execute('SELECT COUNT(1) FROM name_map WHERE recommended_tvk IN (SELECT tvk FROM taxa)').fetchone()[0]})")
print("  Nothing else changed: Codex and your records are untouched until the remap step.")
