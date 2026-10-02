"""Import a collaborator's species records into contributed_observations.

Never touches `observations`: contributed records stay out of your stats,
mapping and iRecord export by construction. Examen sees them through the
assessment_records view.

Columns are found by heading (species, date, method, site, identifier /
determiner, recorder, quantity / count, sex, stage, grid ref, notes). Anything
the file lacks comes from the command line.

  * Names resolve against UKSI by the same resolve() the review importer uses
    (TVK, then name, then synonym). Apply REFUSES if any name is unmatched --
    correct the file and re-run.
  * Vice-county is derived from the grid reference, never copied.
  * Every row carries the import batch, source file, contributor, date received
    and permission. --replace removes an earlier import of the same file first.

DRY RUN by default; --apply to write.

Example (Birmingham - Wheels Park, J. Moore):
  python scripts\\import_contributed.py "data\\contributed\\Birmingham_Wheels.xlsx" ^
      --contributor "Moore, J." --project "Birmingham - Wheels Park" ^
      --site "Birmingham - Wheels Park" --grid SP094868 --stage Adult
"""
import argparse, datetime, os, sqlite3, sys
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import paths
import ast


def _guarded(mod):
    """True if importing scripts/<mod>.py will not run its main()."""
    src = open(os.path.join(ROOT, "scripts", mod + ".py"), encoding="utf-8-sig").read()
    return any(isinstance(n, ast.If) and "__main__" in ast.dump(n.test)
               for n in ast.parse(src).body)


for _m in ("import_status_review", "load_workbook_to_staging"):
    if not _guarded(_m):
        sys.exit(f"  x scripts/{_m}.py has no __main__ guard -- importing it would run it. "
                 "Nothing written.")
from import_status_review import resolve                      # one matcher
from load_workbook_to_staging import normalise_date, resolve_vc  # one date rule, one VC rule

import openpyxl

BACKUP_DIR = r"C:\BiologicalSoftware_Backups\reference"
DEFAULT_METHODS = {"swept": "Net", "sweeping": "Net", "beaten": "Beating tray",
                   "beating": "Beating tray", "incidental": "Field observation"}
HEADINGS = {
    "species":    ("species", "taxon", "scientific"),
    "date":       ("date",),
    "method":     ("method", "sweeping", "beating"),
    "site":       ("site", "location"),
    "determiner": ("identifier", "determiner", "det"),
    "recorder":   ("recorder", "collector"),
    "quantity":   ("quantity", "count", "number", "abundance"),
    "sex":        ("sex",),
    "stage":      ("stage",),
    "grid":       ("grid",),
    "comment":    ("notes", "comment", "status"),
}

ap = argparse.ArgumentParser()
ap.add_argument("xlsx")
ap.add_argument("--contributor", required=True, help='e.g. "Moore, J."')
ap.add_argument("--project", required=True)
ap.add_argument("--client", default=None)
ap.add_argument("--site", default=None, help="overrides the file's site column")
ap.add_argument("--grid", default=None, help="used where a row has no grid reference")
ap.add_argument("--stage", default=None, help="used where a row has no stage")
ap.add_argument("--determiner", default=None, help="overrides the file's identifier column")
ap.add_argument("--sheet", default=None)
ap.add_argument("--method", action="append", default=[], help='"Swept=Net"; adds to the defaults')
ap.add_argument("--received", default=datetime.date.today().isoformat())
ap.add_argument("--permission", default="Permission given to store and to use in assessments")
ap.add_argument("--replace", action="store_true", help="remove an earlier import of this file first")
ap.add_argument("--apply", action="store_true")
a = ap.parse_args()

methods = dict(DEFAULT_METHODS)
for m in a.method:
    k, _, v = m.partition("=")
    methods[k.strip().lower()] = v.strip()

src_file = os.path.basename(a.xlsx)
print(f"Contributed import  --  {'APPLY' if a.apply else 'DRY RUN'}")
print("=" * 76)
print(f"  file: {a.xlsx}")
print(f"  contributor: {a.contributor}   project: {a.project}   client: {a.client}")

# ---------------------------------------------------------------- read
wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
ws = wb[a.sheet] if a.sheet else (wb["Sample Data"] if "Sample Data" in wb.sheetnames
                                  else wb.worksheets[0])
it = ws.iter_rows(values_only=True)
header = [str(h or "").strip() for h in next(it)]
col = {}
for i, h in enumerate(header):
    hl = h.lower()
    for key, words in HEADINGS.items():
        if key not in col and any(hl.startswith(w) or w in hl.split() for w in words):
            col[key] = i
            break
print(f"  sheet: {ws.title}")
print("  columns found: " + ", ".join(f"{k}<-{header[i]!r}" for k, i in col.items()))
if "species" not in col or "date" not in col:
    sys.exit("  x a species column and a date column are required -- nothing written")


def get(row, key):
    i = col.get(key)
    v = row[i] if i is not None and i < len(row) else None
    return v.strip() if isinstance(v, str) else v


u = sqlite3.connect(f"file:{paths.UKSI_DB}?mode=ro", uri=True)
try:
    from DataEntry import bootstrap
    vc_service = bootstrap.get_vc_service_safe()
except Exception as e:
    vc_service = None
    print(f"  (VC service unavailable: {e} -- vice-county will be left empty)")

records, unmatched, bad_dates, how = [], Counter(), [], Counter()
for n, row in enumerate(it, start=2):
    name = get(row, "species")
    if not name:
        continue
    name = str(name)
    r = resolve(u, name, None)
    if not r:
        unmatched[name] += 1
        continue
    tvk, sci, family, order, via = r
    how[via] += 1
    d, ok = normalise_date(get(row, "date"))
    if not d or not ok:
        bad_dates.append((n, get(row, "date")))
        continue
    grid = str(get(row, "grid") or a.grid or "").strip().upper().replace(" ", "") or None
    vcn, vcname = resolve_vc(vc_service, grid) if grid else (None, None)
    rank = u.execute("SELECT rank FROM taxa WHERE tvk=?", (tvk,)).fetchone()
    cn = u.execute("SELECT common_name FROM common_names WHERE tvk=? AND preferred=1",
                   (tvk,)).fetchone()
    raw_method = str(get(row, "method") or "").strip()
    qty = get(row, "quantity")
    try:
        qty = int(qty) if qty not in (None, "") else 1
    except (TypeError, ValueError):
        qty = 1
    det = a.determiner or get(row, "determiner") or a.contributor
    records.append({
        "species_name": sci, "species_tvk": tvk, "common_name": cn[0] if cn else None,
        "order_name": order, "family": family, "taxon_rank": rank[0] if rank else None,
        "date": d, "grid_ref": grid,
        "grid_precision": {2: 10000, 4: 1000, 6: 100, 8: 10, 10: 1}.get(len(grid) - 2) if grid else None,
        "vice_county": vcname, "vc_number": vcn,
        "site_name": a.site or get(row, "site"),
        "recorder": get(row, "recorder") or a.contributor, "determiner": det,
        "sex": get(row, "sex"), "stage": get(row, "stage") or a.stage, "quantity": qty,
        "method": methods.get(raw_method.lower(), raw_method or None),
        "comment": get(row, "comment"),
        "record_type": "Commercial", "project_name": a.project, "client": a.client,
        "contributor": a.contributor, "source_file": src_file,
        "received_date": a.received, "permission": a.permission,
    })

# ---------------------------------------------------------------- report
print(f"\n  rows read: {len(records) + sum(unmatched.values()) + len(bad_dates)}   "
      f"importable: {len(records)}   species: {len({r['species_tvk'] for r in records})}")
print(f"  matched by: {dict(how)}")
print(f"  dates: {dict(Counter(r['date'] for r in records))}")
print(f"  methods: {dict(Counter(r['method'] for r in records))}   (from {sorted(methods)})")
print(f"  grid refs: {dict(Counter(r['grid_ref'] for r in records))}")
print(f"  vice-county: {dict(Counter((r['vc_number'], r['vice_county']) for r in records))}")
print(f"  site: {dict(Counter(r['site_name'] for r in records))}   "
      f"stage: {dict(Counter(r['stage'] for r in records))}")
print(f"  recorder / determiner: "
      f"{dict(Counter((r['recorder'], r['determiner']) for r in records))}")
if unmatched:
    print(f"\n  x {len(unmatched)} unmatched names (correct them in the file and re-run):")
    for nm, c in unmatched.most_common():
        print(f"      {nm!r} x{c}")
if bad_dates:
    print(f"\n  x {len(bad_dates)} rows with an unreadable date: {bad_dates[:5]}")

o = sqlite3.connect(str(paths.OBSERVATUM_DB))
if not o.execute("SELECT 1 FROM sqlite_master WHERE name='contributed_observations'").fetchone():
    sys.exit("\n  x contributed_observations does not exist -- run patch_contributed_schema.py first")
prior = o.execute("SELECT import_batch, COUNT(1) FROM contributed_observations "
                  "WHERE source_file=? AND contributor=? GROUP BY 1",
                  (src_file, a.contributor)).fetchall()
if prior:
    print(f"\n  this file was imported before: {prior}")
    if not a.replace:
        print("  x re-run with --replace to remove that import first")

blocked = bool(unmatched or bad_dates or (prior and not a.replace) or not records)
if not a.apply:
    print("\n" + ("Fix the items marked x before applying.\n" if blocked else "All checks pass.\n")
          + "DRY RUN -- nothing has been changed.")
    sys.exit(0)
if blocked:
    sys.exit("\nABORTED -- fix the items marked x. Nothing has been changed.")

# ---------------------------------------------------------------- write
os.makedirs(BACKUP_DIR, exist_ok=True)
bp = os.path.join(BACKUP_DIR, "observatum_pre_contributed_import_" +
                  datetime.datetime.now().strftime("%Y%m%d_%H%M%S") + ".db")
t = sqlite3.connect(bp); o.backup(t); t.close()
print(f"\n  backup: {bp}")

batch = f"Contributed {a.contributor} {datetime.datetime.now().isoformat(timespec='seconds')}"
cols = list(records[0].keys()) + ["import_batch"]
with o:
    if prior:
        o.execute("DELETE FROM contributed_observations WHERE source_file=? AND contributor=?",
                  (src_file, a.contributor))
        print(f"  removed earlier import: {sum(n for _, n in prior)} rows")
    o.executemany(f"INSERT INTO contributed_observations ({', '.join(cols)}) "
                  f"VALUES ({', '.join('?' * len(cols))})",
                  [[r[c] for c in cols[:-1]] + [batch] for r in records])

# ---------------------------------------------------------------- read back
n = o.execute("SELECT COUNT(1), COUNT(DISTINCT species_tvk) FROM contributed_observations "
              "WHERE import_batch=?", (batch,)).fetchone()
print(f"  written: {n[0]} records, {n[1]} species   batch: {batch!r}")
print("\n  In assessment_records for this project:")
for r in o.execute("""SELECT origin, substr(date,1,4), COUNT(1), COUNT(DISTINCT species_name)
                      FROM assessment_records WHERE project_name=? GROUP BY 1,2""", (a.project,)):
    print(f"    {r[0]:<12} {r[1]}  {r[2]} records  {r[3]} species")
own = o.execute("SELECT COUNT(1) FROM observations WHERE project_name=?", (a.project,)).fetchone()[0]
print(f"  observations under this project (your own, committed): {own}  -- untouched")
print("\nDone.")
