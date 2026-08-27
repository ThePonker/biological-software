"""Load a Tabella workbook's observation rows into a DataEntry staging job.

One-off migration tool (Session 28, Aug 2026). Reads a .xlsm from the Active
Record Books folder and creates ONE entry_jobs row plus its entry_staging rows.

    python scripts/load_workbook_to_staging.py "Jukes Toddington.xlsm" --dry-run
    python scripts/load_workbook_to_staging.py "Jukes Toddington.xlsm"
    python scripts/load_workbook_to_staging.py "RWE Elmley.xlsm" --sheet personal

Reads row 4 as headers and data from row 5, matching columns BY NAME so a
reordered workbook cannot misfile data. Blank rows *within* the data are kept
(they carry Wil's date/trap spacing); trailing blanks after the last species
are dropped.

Never writes to the workbook. Never writes to `observations` — staging only.
Undo is deleting the job's staging rows.
"""
from __future__ import annotations

import argparse
import os
import sqlite3
import sys
from datetime import datetime

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "Observatum"))
import paths  # noqa: E402

HEADER_ROW = 4
FIRST_DATA_ROW = 5

SHEETS = {
    "commercial": "Commercial Observations",
    "personal": "Personal Observations",
}

# workbook header -> entry_staging column
COLUMN_MAP = {
    "Species": "species_name",
    "TVK": "species_tvk",
    "Common Name": "common_name",
    "Order": "order_name",
    "Family": "family",
    "Date": "date",
    "Grid Ref": "grid_ref",
    "Site Name": "site_name",
    "VC": "vice_county",
    "Recorder": "recorder",
    "Determiner": "determiner",
    "Method": "method",
    "Stage": "stage",
    "Sex": "sex",
    "Qty": "quantity",
    "Certainty": "certainty",
    "Comment": "comment",
}

# Deliberately ignored: Conservation, Red List, Rarity, S41, Legal Protection,
# BAP, Class, Sort Code, Personal, Commercial, Specimens. Display aids, not data.


def normalise_date(value):
    """ISO-normalise a workbook date cell. Returns (value_or_None, ok)."""
    if value is None:
        return None, True
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d"), True
    text = str(value).strip()
    if not text:
        return None, True
    try:
        from DataEntry import date_utils
        out = date_utils.normalise(text)
        if out:
            return out, True
    except Exception:
        pass
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d"), True
        except ValueError:
            continue
    return text, False  # keep it, but flag it


def as_int(value):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def read_rows(path, sheet_name):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"Sheet {sheet_name!r} not in {os.path.basename(path)}. "
                         f"Found: {wb.sheetnames}")
    ws = wb[sheet_name]

    grid = list(ws.iter_rows(min_row=HEADER_ROW, values_only=True))
    wb.close()
    if not grid:
        return [], {}

    headers = [(h or "").strip() if isinstance(h, str) else h for h in grid[0]]
    index = {}
    for pos, head in enumerate(headers):
        if isinstance(head, str) and head in COLUMN_MAP:
            index[COLUMN_MAP[head]] = pos

    missing = sorted(set(COLUMN_MAP.values()) - set(index))
    if "species_name" not in index:
        raise SystemExit("No 'Species' column found on row 4 — wrong sheet or layout?")

    body = grid[1:]  # from FIRST_DATA_ROW

    # last row that actually has a species: everything after it is trailing blank
    last = -1
    for i, row in enumerate(body):
        cell = row[index["species_name"]] if index["species_name"] < len(row) else None
        if cell is not None and str(cell).strip():
            last = i
    if last < 0:
        return [], {"missing": missing, "read": 0}

    return body[: last + 1], {"missing": missing, "read": last + 1, "index": index}


_VC_CACHE = {}


def resolve_vc(service, grid_ref):
    """(number, name) for a grid ref, cached per distinct ref. (None, None) on failure."""
    if not service or not grid_ref:
        return None, None
    key = str(grid_ref).strip().upper().replace(" ", "")
    if key in _VC_CACHE:
        return _VC_CACHE[key]
    try:
        got = service.get_vc_from_grid_ref(key)
        out = (got[0], got[1]) if got else (None, None)
    except Exception:
        out = (None, None)
    _VC_CACHE[key] = out
    return out


def build_records(body, index, vc_service=None):
    records, blanks, no_tvk, bad_dates, vc_ok = [], 0, 0, [], 0
    for order, row in enumerate(body, start=1):
        def cell(col):
            pos = index.get(col)
            if pos is None or pos >= len(row):
                return None
            v = row[pos]
            if isinstance(v, str):
                v = v.strip()
            return v or None

        species = cell("species_name")
        rec = {"row_order": order}

        if not species:
            blanks += 1
            records.append(rec)  # spacer row, preserved
            continue

        for col in COLUMN_MAP.values():
            rec[col] = cell(col)

        iso, ok = normalise_date(rec.get("date"))
        rec["date"] = iso
        if not ok:
            bad_dates.append((order, iso))

        rec["quantity"] = as_int(rec.get("quantity"))
        rec["vc_number"] = as_int(rec.get("vice_county"))
        if rec["vc_number"] is not None:
            rec["vice_county"] = str(rec["vice_county"]).strip()

        num, name = resolve_vc(vc_service, rec.get("grid_ref"))
        if num is not None:
            rec["vc_number"] = num
            rec["vice_county"] = name or rec.get("vice_county")
            vc_ok += 1

        if not rec.get("species_tvk"):
            no_tvk += 1

        records.append(rec)

    return records, {"blanks": blanks, "no_tvk": no_tvk, "bad_dates": bad_dates, "vc_ok": vc_ok}


def insert(conn, job_name, mode, records, source_file):
    now = datetime.now().isoformat(timespec="seconds")
    cur = conn.execute(
        "INSERT INTO entry_jobs (name, mode, status, notes, created_at, updated_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (job_name, mode, "active",
         f"Loaded from {source_file} on {now}", now, now),
    )
    job_id = cur.lastrowid

    cols = ["job_id", "row_order"] + list(COLUMN_MAP.values()) + \
           ["vc_number", "created_at", "updated_at"]
    placeholders = ",".join("?" for _ in cols)
    sql = f"INSERT INTO entry_staging ({','.join(cols)}) VALUES ({placeholders})"

    for rec in records:
        values = [job_id, rec.get("row_order")]
        values += [rec.get(c) for c in COLUMN_MAP.values()]
        values += [rec.get("vc_number"), now, now]
        conn.execute(sql, values)

    conn.commit()
    return job_id


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", help="Filename inside Active Record Books, or a full path.")
    ap.add_argument("--sheet", choices=sorted(SHEETS), default="commercial")
    ap.add_argument("--mode", default="Commercial", help="Job mode (default: Commercial).")
    ap.add_argument("--name", default=None, help="Job name (default: workbook filename).")
    ap.add_argument("--dry-run", action="store_true", help="Report only; insert nothing.")
    args = ap.parse_args()

    path = args.workbook
    if not os.path.isabs(path):
        path = os.path.join(os.environ["USERPROFILE"], "OneDrive",
                            "Active Record Books", args.workbook)
    if not os.path.exists(path):
        raise SystemExit(f"Not found: {path}")

    sheet = SHEETS[args.sheet]
    body, meta = read_rows(path, sheet)
    if not body:
        raise SystemExit(f"No data rows on {sheet!r}.")

    from DataEntry import bootstrap
    vc_service = bootstrap.get_vc_service_safe()
    records, stats = build_records(body, meta["index"], vc_service)
    real = len(records) - stats["blanks"]

    print(f"\n  workbook   {os.path.basename(path)}")
    print(f"  sheet      {sheet}")
    print(f"  rows read  {meta['read']}  (to last species row)")
    print(f"  records    {real}")
    print(f"  spacers    {stats['blanks']}  (blank rows kept)")
    print(f"  no TVK     {stats['no_tvk']}")
    print(f"  VC derived {stats['vc_ok']} of {real}")
    if meta["missing"]:
        print(f"  columns not found: {', '.join(meta['missing'])}")
    if stats["bad_dates"]:
        print(f"  unparsed dates: {len(stats['bad_dates'])} "
              f"(first few: {stats['bad_dates'][:5]})")

    if args.dry_run:
        print("\n  DRY RUN — nothing written.\n")
        return

    job_name = args.name or os.path.splitext(os.path.basename(path))[0]
    conn = sqlite3.connect(str(paths.OBSERVATUM_DB))
    job_id = insert(conn, job_name, args.mode, records, os.path.basename(path))
    conn.close()
    print(f"\n  created job {job_id}: {job_name!r} ({args.mode})")
    print(f"  undo: DELETE FROM entry_staging WHERE job_id={job_id}; "
          f"DELETE FROM entry_jobs WHERE id={job_id};\n")


if __name__ == "__main__":
    main()
