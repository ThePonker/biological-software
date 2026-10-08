"""
Examen -- the assessment report as plain data, read from the workbook.

The PDF (pdf_export) and Word (word_export) reports are both laid out from the
assessment workbook, so every figure is the workbook's own. This module is the one
place that reads the workbook's sheets into sections; each renderer only does layout.

    from Examen.report_model import read_report, build_report
    rep = read_report("Glory Park assessment.xlsx")        # from a saved workbook
    rep = build_report(result, detail, project, jurisdiction="England")   # from an analysis

A cell value is either plain (str / number / None) or, where the workbook greys a
designation that does not apply in the jurisdiction, a list of (text, applies) parts.
"""
from __future__ import annotations

import os
import tempfile

# Columns shown in the report's key species table (the workbook has more)
KEY_TABLE = ["Tier", "Order", "Family", "Species", "Common name", "Conservation status", "SQS",
             "Broad biotope", "Habitat"]
ACCOUNTS_NOTE = ("Accounts are the author's where written; otherwise the current published "
                 "review account, quoted and cited where its licence allows. \u201cThis survey\u201d "
                 "gives the number of individuals, places and months recorded.")
SECTIONS = (("Habitats", "Habitats"),
            ("Assemblages", "Specific assemblage types"),
            ("Guilds", "Feeding guilds"))


def value(v):
    """Workbook cell -> plain value, or [(text, applies), ...] for rich text."""
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock
    except ImportError:  # pragma: no cover
        return v
    if isinstance(v, CellRichText):
        parts = []
        for p in v:
            if isinstance(p, TextBlock):
                grey = p.font is not None and bool(p.font.i or p.font.color is not None)
                parts.append((p.text, not grey))
            else:
                parts.append((str(p), True))
        return parts
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def plain(v) -> str:
    """Any cell value as plain text."""
    if v is None:
        return ""
    if isinstance(v, list):
        return "".join(t for t, _ in v)
    return str(v)


def _rows(ws):
    out = []
    for r in ws.iter_rows(values_only=True):
        r = [value(c) for c in r]
        while r and r[-1] in (None, ""):
            r.pop()
        out.append(r)
    return out


def _header_index(rows, start=2):
    for i in range(start, len(rows)):
        if sum(1 for c in rows[i] if c not in (None, "")) >= 3:
            return i
    return None


def _sub(rows):
    return plain(rows[1][0]) if len(rows) > 1 and rows[1] else ""


def _summary(ws):
    rows = _rows(ws)
    figures, basis, attribution, mode = [], [], [], "figures"
    for r in rows[2:]:
        if not r:
            continue
        if r[0] == "Basis of assessment":
            mode = "basis"
            continue
        if mode == "figures":
            figures.append((r[0], r[1] if len(r) > 1 else "", r[2] if len(r) > 2 else ""))
        elif mode == "basis" and len(r) >= 2:
            basis.append((plain(r[0]), r[1]))
        else:
            attribution.append(plain(r[0]))
    subtitle = _sub(rows).replace(" — figures and their basis", "")
    return {"title": plain(rows[0][0]) if rows and rows[0] else "Assessment",
            "subtitle": subtitle, "figures": figures, "basis": basis,
            "attribution": attribution}


def _key_species(ws):
    rows = _rows(ws)
    h = _header_index(rows)
    out = {"subtitle": _sub(rows), "rows": [], "accounts": []}
    if h is None:
        return out
    heads = [plain(c) for c in rows[h]]
    col = {name: i for i, name in enumerate(heads)}
    occ_col = len(heads) - 1 if heads and heads[-1].startswith("Occurrence") else None
    for r in rows[h + 1:]:
        if not (r and r[0] and plain(r[0]).endswith("Key")):
            continue
        g = lambda c: r[col[c]] if col.get(c) is not None and col[c] < len(r) else None  # noqa: E731
        out["rows"].append([g(c) for c in KEY_TABLE])
        occ = plain(r[occ_col]) if occ_col is not None and occ_col < len(r) else ""
        src = plain(g("Account source"))
        out["accounts"].append({
            "species": plain(g("Species")), "common": plain(g("Common name")),
            "status": g("Conservation status"), "account": plain(g("Species account")),
            # the author's own account needs no credit; a review's does
            "source": "" if src == "Your account" else src,
            # the workbook's "[write occurrence]" marker is for the author, not the reader
            "evidence": occ.replace("[write occurrence]", "").strip(" ·"),
        })
    return out


def _table(ws):
    rows = _rows(ws)
    h = _header_index(rows)
    out = {"subtitle": _sub(rows), "heads": [], "rows": [], "notes": []}
    if h is None:
        return out
    out["heads"] = [plain(c) for c in rows[h]]
    for r in rows[h + 1:]:
        if not r:
            continue
        filled = sum(1 for c in r if c not in (None, ""))
        if filled == 1 and len(plain(r[0])) > 60:
            out["notes"].append(plain(r[0]))
        else:
            out["rows"].append(r)
    return out


def _appendix(ws):
    rows = _rows(ws)
    h = _header_index(rows)
    out = {"subtitle": _sub(rows), "heads": [plain(c) for c in rows[h]] if h is not None else [],
           "rows": [], "footer": []}
    if h is None:
        return out
    for r in rows[h + 1:]:
        if not r:
            continue
        if r[0] in (None, "") and len(r) > 2:
            out["footer"].append((plain(r[1]), r[2]))
        else:
            out["rows"].append(r)
    return out


def _definitions(ws):
    out = []
    for r in _rows(ws)[1:]:
        if not r:
            continue
        out.append(("section", plain(r[0])) if len(r) == 1 else ("row", r[0], r[1]))
    return out


def read_report(xlsx_path) -> dict:
    """The workbook's sheets as report sections."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, rich_text=True)
    rep = _summary(wb["Summary"])
    stamp = {k: plain(v) for k, v in rep["basis"]}
    rep["run"] = stamp.get("Assessment run", "")
    rep["codex"] = stamp.get("Codex version", "").split(",")[0]
    rep["key"] = _key_species(wb["Key species"]) if "Key species" in wb.sheetnames else None
    rep["tables"] = [(label, _table(wb[sheet])) for sheet, label in SECTIONS
                     if sheet in wb.sheetnames]
    rep["appendix"] = _appendix(wb["Species appendix"]) if "Species appendix" in wb.sheetnames else None
    rep["definitions"] = (_definitions(wb["Status definitions"])
                          if "Status definitions" in wb.sheetnames else [])
    return rep


def build_report(result, detail, project, jurisdiction="England",
                 sqs_basis="Pantheon published", pooled_years=False) -> dict:
    """Build the workbook to a temporary file and read it back as report sections."""
    try:
        from Examen.workbook_export import export_workbook
    except ImportError:  # pragma: no cover
        from workbook_export import export_workbook
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        export_workbook(result, detail, project, tmp, jurisdiction=jurisdiction,
                        sqs_basis=sqs_basis, pooled_years=pooled_years)
        return read_report(tmp)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
