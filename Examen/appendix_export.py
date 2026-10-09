"""
Examen - Species Appendix Export

Generates the standard invertebrate survey appendix as Excel (.xlsx).
Columns: Species Name, Common Name, Conservation Status, SQS, Tier,
Broad Biotope, Habitat, Family, Order. Key species grouped at top, each group
in taxonomic order (examen_data.in_taxonomic_order, the workbook's rule).
Summary row at bottom.

Matches the workbook (backlog E17): statuses named for the jurisdiction, other
jurisdictions' designations greyed; "(derived)" on a score derived from current
status; both SQIs in the totals when any score was derived.

Requires openpyxl.
"""

from PySide6.QtWidgets import QFileDialog, QMessageBox

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Colours matching Naturalist theme
MOSS = "4A7C59"
ACCENT = "7C6C9F"
RED = "A63D40"
AMBER = "C2956E"
LIGHT_BG = "F0EEEA"
HEADER_BG = "5A4D78"
KEY_BG = "F0EDF5"


def _get_taxonomy(tvks: list[str]) -> dict[str, dict]:
    """Look up common_name, family, order for each TVK -- via examen_data.load_taxonomy."""
    from .examen_data import load_taxonomy
    return load_taxonomy(list(tvks or []))


def export_appendix(parent_widget, site_name: str, result, detail=None,
                    jurisdiction="England"):
    """Export species appendix to Excel.

    Args:
        parent_widget: Qt parent for file dialog
        site_name: Site name for filename default
        result: AnalysisResult from PantheonAnalysisService
        detail: Optional SiteDetail with full species list
        jurisdiction: whose designations apply; others are greyed, as in the workbook
    """
    from .examen_data import in_taxonomic_order
    from .workbook_export import (_parts_from_string, _sqi_cell, appendix_name, status_cell,
                                  status_parts)
    if not HAS_OPENPYXL:
        QMessageBox.warning(parent_widget, "Examen",
                            "openpyxl required for Excel export.\npip install openpyxl")
        return

    safe_name = "".join(c if c.isalnum() or c in " _-" else "_" for c in site_name)
    path, _ = QFileDialog.getSaveFileName(
        parent_widget, "Export Species Appendix",
        f"{safe_name}_species_appendix.xlsx", "Excel (*.xlsx)")
    if not path:
        return

    # Gather all species TVKs
    key_tvks = {k.tvk for k in result.key_species}
    all_tvks = list(key_tvks)
    non_key_species = []
    if detail:
        for sp in detail.species_list:
            if sp.tvk and sp.tvk not in key_tvks:
                all_tvks.append(sp.tvk)
                non_key_species.append(sp)

    taxonomy = _get_taxonomy(all_tvks)
    derived = getattr(result, "derived_sqs_tvks", set()) or set()

    def sqs_cell(tvk, sqs):
        if not sqs:
            return ""
        return f"{sqs} (derived)" if tvk in derived else sqs

    # Taxonomic order within each group -- one rule, shared with the workbook
    key_sorted = in_taxonomic_order(result.key_species, name=lambda k: k.species_name)
    non_key_sorted = in_taxonomic_order(non_key_species)

    # Build workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Species Appendix"

    # Styles
    hdr_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    hdr_fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    key_fill = PatternFill(start_color=KEY_BG, end_color=KEY_BG, fill_type="solid")
    body_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    border = Border(
        left=Side(style="thin", color="D4D0C8"), right=Side(style="thin", color="D4D0C8"),
        top=Side(style="thin", color="D4D0C8"), bottom=Side(style="thin", color="D4D0C8"))
    center = Alignment(horizontal="center")

    # Title row
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"Species Appendix -- {site_name}"
    title_cell.font = Font(name="Georgia", bold=True, size=14, color=ACCENT)

    # Headers (row 3)
    headers = ["Species Name", "Common Name", "Conservation Status", "SQS", "Tier",
               "Broad Biotope", "Habitat", "Family", "Order"]
    widths = [32, 24, 20, 8, 10, 18, 18, 18, 16]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        cell.border = border
        ws.column_dimensions[chr(64 + col)].width = w

    # Key species section
    row = 4
    if key_sorted:
        for k in key_sorted:
            tax = taxonomy.get(k.tvk, {})
            name = (f"{k.species_name} ({k.recorded_note})"
                    if getattr(k, "recorded_note", "") else k.species_name)
            values = [name, tax.get("common", ""),
                      status_cell(status_parts(k, jurisdiction), jurisdiction),
                      sqs_cell(k.tvk, k.sqs), k.tier, k.broad_biotope, k.habitat,
                      k.family or tax.get("family", ""), tax.get("order", "")]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font = bold_font if col == 1 else body_font
                cell.fill = key_fill
                cell.border = border
                if col in (4, 5):
                    cell.alignment = center
            row += 1

    # Separator row
    if key_sorted and non_key_sorted:
        for col in range(1, 10):
            cell = ws.cell(row=row, column=col)
            cell.border = border
        row += 1

    # Non-key species
    for sp in non_key_sorted:
        tax = taxonomy.get(sp.tvk, {})
        status = status_cell(_parts_from_string(
            getattr(sp, "status_full", "") or sp.status or "", jurisdiction), jurisdiction)
        values = [appendix_name(sp), tax.get("common", ""), status, sqs_cell(sp.tvk, sp.sqs),
                  "", sp.broad_biotope or "", sp.habitat or "",
                  tax.get("family", ""), tax.get("order", "")]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = body_font
            cell.border = border
            if col in (4, 5):
                cell.alignment = center
        row += 1

    # Summary row
    row += 1
    summary_font = Font(name="Arial", size=10, bold=True, color=ACCENT)
    ws.cell(row=row, column=1, value="TOTALS").font = summary_font
    total_species = len(key_sorted) + len(non_key_sorted)
    ws.cell(row=row, column=2, value=f"{total_species} species").font = summary_font
    ws.cell(row=row, column=5, value=f"{len(key_sorted)} key").font = summary_font
    sqi = result.overall_sqi
    if sqi and sqi.sqi:
        ws.cell(row=row, column=4, value=f"SQI: {_sqi_cell(sqi)}").font = summary_font
    ws.cell(row=row, column=6, value=f"R:{result.rare_count} S:{result.scarce_count} P:{result.priority_count}").font = summary_font
    pub = getattr(result, "overall_sqi_published", None)
    if derived and pub is not None:
        row += 1
        ws.cell(row=row, column=4, value=f"SQI on Pantheon scores only: {_sqi_cell(pub)}").font = summary_font
        ws.cell(row=row, column=6, value=f"{len(derived)} score(s) derived from current status "
                                         "by Pantheon's published rule").font = Font(
            name="Arial", size=9, italic=True)
    if jurisdiction:
        row += 1
        ws.cell(row=row, column=1, value=f"Statuses as they apply in {jurisdiction}; "
                                         "other jurisdictions' designations in grey.").font = Font(
            name="Arial", size=9, italic=True)

    ws.freeze_panes = "A4"

    try:
        wb.save(path)
        QMessageBox.information(parent_widget, "Examen",
                                f"Species appendix exported.\n{total_species} species, "
                                f"{len(key_sorted)} key.\n{path}")
    except Exception as e:
        QMessageBox.warning(parent_widget, "Examen", f"Export failed:\n{e}")
