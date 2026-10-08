"""
Examen -- Word assessment report (backlog E4)

The same sections as the PDF (pdf_export), laid out as an editable .docx for copying
into a report template. Both come from the assessment workbook via report_model, so
every figure is the workbook's own. Built for copy and paste: real Word headings
(Heading 1 / 2, so they take the template's styles when pasted), real Word tables,
plain body text; scientific names in italic; designations that do not apply in the
jurisdiction grey italic, as in the workbook.

    from Examen.word_export import export_word, workbook_to_word
    export_word(result, detail, project, "Glory Park assessment.docx", jurisdiction="England")
    workbook_to_word("Glory Park assessment.xlsx", "Glory Park assessment.docx")

Requires python-docx (py -3.14 -m pip install python-docx).
"""
from __future__ import annotations

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False

try:
    from Examen.report_model import ACCOUNTS_NOTE, KEY_TABLE, build_report, read_report
except ImportError:  # pragma: no cover
    from report_model import ACCOUNTS_NOTE, KEY_TABLE, build_report, read_report

MOSS = "4A7C59"
MUTED = RGBColor(0x6B, 0x72, 0x80) if HAS_DOCX else None
NA_GREY = RGBColor(0x9C, 0xA3, 0xAF) if HAS_DOCX else None
FONT = "Arial"
TABLE_CM = {   # column widths per section, in cm (17.4 cm of text width)
    "Habitats": [3.6, 4.6, 1.8, 1.8, 2.0, 2.4],
    "Specific assemblage types": [3.8, 1.4, 1.4, 1.3, 1.7, 1.5, 1.9, 4.4],
    "Feeding guilds": [3.0, 8.0, 2.4],
}


# ------------------------------------------------------------------ helpers
def _runs(par, v, italic=False, bold=False, size=None):
    """Write a cell value into a paragraph; rich parts that do not apply grey italic."""
    parts = v if isinstance(v, list) else [("" if v is None else str(v), True)]
    for text, applies in parts:
        r = par.add_run(text)
        r.italic = italic or not applies
        r.bold = bold
        if not applies:
            r.font.color.rgb = NA_GREY
        if size:
            r.font.size = Pt(size)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _table(doc, heads, rows, widths_cm=None, italic_cols=(), size=8):
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = t.rows[0]
    _repeat_header(hdr)
    for j, h in enumerate(heads):
        c = hdr.cells[j]
        c.text = ""
        _runs(c.paragraphs[0], h, bold=True, size=size)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, MOSS)
    for r in rows:
        cells = t.add_row().cells
        for j in range(len(heads)):
            v = r[j] if j < len(r) else None
            _runs(cells[j].paragraphs[0], v, italic=(j in italic_cols), size=size)
    if widths_cm:
        # Word reads the cell widths, LibreOffice the grid: set both, autofit off
        t.autofit = False
        widths = [Cm(w) for w in widths_cm[:len(heads)]]
        for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(w.twips)))
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = w
    return t


def _note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    return p


def _field(par, instr):
    """A Word field (e.g. PAGE) so page numbers update in Word."""
    r = par.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        r._r.append(el)


# ------------------------------------------------------------------ document
def report_to_word(rep, docx_path):
    """Lay out report sections (report_model) as a Word document."""
    if not HAS_DOCX:
        raise ImportError("python-docx is required for the Word report: "
                          "py -3.14 -m pip install python-docx")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Title"):
        doc.styles[name].font.name = FONT
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    head = sec.header.paragraphs[0]
    _runs(head, rep["title"], size=7)
    foot = sec.footer.paragraphs[0]
    _runs(foot, f"Examen — run {rep['run']}" + (f" — Codex {rep['codex']}" if rep["codex"] else "")
          + "    Page ", size=7)
    _field(foot, "PAGE")

    doc.add_heading(rep["title"], level=0)
    sub = doc.add_paragraph()
    _runs(sub, rep["subtitle"], size=11)
    sub.runs[0].font.color.rgb = MUTED

    doc.add_heading("Basis of assessment", level=2)
    _table(doc, ["Item", "Value"], [[k, v] for k, v in rep["basis"]], [5.0, 12.4], size=8)

    doc.add_heading("1  Summary of figures", level=1)
    _table(doc, ["Measure", "Value", "Basis / note"], rep["figures"], [4.8, 1.8, 10.8], size=8.5)
    for a in rep["attribution"]:
        _note(doc, a)

    doc.add_heading("2  Key species", level=1)
    key = rep.get("key")
    if key and key["subtitle"]:
        _note(doc, key["subtitle"])
    if not key or not key["rows"]:
        doc.add_paragraph("No key species recorded.")
    else:
        _table(doc, KEY_TABLE, key["rows"], [1.4, 1.6, 1.9, 2.8, 1.9, 2.0, 0.8, 2.3, 2.7],
               italic_cols={KEY_TABLE.index("Species")}, size=7.5)
        doc.add_heading("Species accounts", level=2)
        for a in key["accounts"]:
            p = doc.add_paragraph()
            _runs(p, a["species"], italic=True, bold=True)
            if a["common"]:
                _runs(p, f" ({a['common']})")
            _runs(p, " — ")
            _runs(p, a["status"])
            if a["account"]:
                doc.add_paragraph(a["account"])
            bits = ([f"Account: {a['source']}"] if a["source"] else []) + \
                   ([f"This survey: {a['evidence']}"] if a["evidence"] else [])
            if bits:
                _note(doc, "    ".join(bits))
        _note(doc, ACCOUNTS_NOTE)

    for n, (label, t) in enumerate(rep["tables"], start=3):
        doc.add_heading(f"{n}  {label}", level=1)
        if t["subtitle"]:
            _note(doc, t["subtitle"])
        if t["heads"]:
            _table(doc, t["heads"], t["rows"], TABLE_CM.get(label), size=8)
        for x in t["notes"]:
            _note(doc, x)

    ap = rep.get("appendix")
    if ap:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Appendix A  Species list", level=1)
        if ap["subtitle"]:
            _note(doc, ap["subtitle"])
        heads = ap["heads"]
        _table(doc, heads, ap["rows"], [1.9, 2.1, 3.3, 2.2, 1.9, 0.9, 2.0, 2.2, 1.0],
               italic_cols={heads.index("Species")} if "Species" in heads else (), size=7.5)
        if ap["footer"]:
            doc.add_paragraph()
            _table(doc, ["Total", ""], [[k, v] for k, v in ap["footer"]], [6.0, 3.0], size=8)

    if rep["definitions"]:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Appendix B  Status definitions", level=1)
        for d in rep["definitions"]:
            if d[0] == "section":
                doc.add_heading(d[1], level=2)
            else:
                p = doc.add_paragraph()
                _runs(p, d[1], bold=True, size=9)
                _runs(p, "  —  ", size=9)
                _runs(p, d[2], size=9)
    doc.core_properties.title = rep["title"]
    doc.core_properties.subject = "Invertebrate assemblage assessment"
    doc.save(docx_path)
    return docx_path


def workbook_to_word(xlsx_path, docx_path):
    return report_to_word(read_report(xlsx_path), docx_path)


def export_word(result, detail, project, path, jurisdiction="England",
                sqs_basis="Pantheon published", pooled_years=False):
    return report_to_word(build_report(result, detail, project, jurisdiction=jurisdiction,
                                       sqs_basis=sqs_basis, pooled_years=pooled_years), path)
