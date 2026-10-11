"""scripts/combined_species_list.py for any joint survey (Sundon, 11 Oct 2026).

Wil's side now comes from his committed records under the project name (the
Birmingham staging job it used to read has been committed and cleared). Runs the
script against small temporary databases and a temporary spreadsheet; the real
data copies are not opened.
"""
import os
import runpy
import sqlite3
import sys

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(ROOT, "scripts", "combined_species_list.py")


def _uksi(path):
    c = sqlite3.connect(path)
    c.executescript("""
        CREATE TABLE taxa (tvk TEXT PRIMARY KEY, scientific_name TEXT, rank TEXT,
                           family TEXT, "order" TEXT, sort_code INTEGER, sort_order TEXT);
        CREATE TABLE common_names (common_name TEXT, tvk TEXT, preferred INTEGER);
        CREATE TABLE synonyms (synonym TEXT, tvk TEXT);
        INSERT INTO taxa VALUES ('T1','Bombus pascuorum','Species','Apidae','Hymenoptera',1,'A');
        INSERT INTO taxa VALUES ('T2','Diaperis boleti','Species','Tenebrionidae','Coleoptera',2,'B');
        INSERT INTO taxa VALUES ('T3','Larinus carlinae','Species','Curculionidae','Coleoptera',3,'C');
        INSERT INTO synonyms VALUES ('Larinus planus','T3');
    """)
    c.commit()
    c.close()


def _obs(path):
    c = sqlite3.connect(path)
    c.execute("CREATE TABLE observations (project_name TEXT, species_name TEXT, species_tvk TEXT)")
    c.executemany("INSERT INTO observations VALUES (?,?,?)", [
        ("Sundon", "Diaperis boleti", "T2"),
        ("Sundon", "Bombus pascuorum", "T1"),
        ("Other", "Larinus carlinae", "T3"),          # another project: not read
    ])
    c.commit()
    c.close()


def _moore(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Survey Data"
    ws.append(["Species Name", "Common Name", "Sweeping/Beating", "Site", "Date Recorded"])
    ws.append(["Diaperis boleti", "a darkling beetle", "Incidental", "Sundon", "2026-05-28"])
    ws.append(["Larinus planus", "a weevil", "Swept", "Sundon", "2026-06-22"])
    ws.append([None, None, None, None, None])                 # trailing blank rows
    wb.save(path)


@pytest.fixture
def run(tmp_path, monkeypatch):
    sys.path[:0] = [ROOT, os.path.join(ROOT, "scripts")]
    import paths
    _uksi(tmp_path / "uksi.db")
    _obs(tmp_path / "observatum.db")
    _moore(tmp_path / "moore.xlsx")
    monkeypatch.setattr(paths, "UKSI_DB", tmp_path / "uksi.db")
    monkeypatch.setattr(paths, "OBSERVATUM_DB", tmp_path / "observatum.db")

    def go(*extra):
        out = tmp_path / "out.xlsx"
        monkeypatch.setattr(sys, "argv", [SCRIPT, str(tmp_path / "moore.xlsx"),
                                          "--out", str(out), *extra])
        runpy.run_path(SCRIPT, run_name="__main__")
        rows = list(openpyxl.load_workbook(out).active.iter_rows(values_only=True))
        return rows
    return go


def test_project_reads_wils_committed_records(run, capsys):
    rows = run("--project", "Sundon")
    assert rows[0][0] == "Sundon -- combined species list"
    table = {r[0]: r[5:8] for r in rows[4:]}
    assert table == {
        "Bombus pascuorum": (1, None, "W. Heeney"),
        "Diaperis boleti": (1, 1, "Both"),
        "Larinus carlinae": (None, 1, "J. Moore"),       # matched through the synonym
    }
    out = capsys.readouterr().out
    assert "W.J. Heeney: 2 records    J. Moore: 2 records" in out
    assert "species: 3   both: 1   W. Heeney only: 1   J. Moore only: 1" in out


def test_default_project_is_birmingham(run):
    rows = run()
    assert rows[0][0] == "Birmingham - Wheels Park -- combined species list"
    assert all(r[5] is None for r in rows[4:])           # no Birmingham records here
