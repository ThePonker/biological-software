"""Round 9 agent B (12 Oct): data fixes.

B1 butterflies labelled "insect - moth" (F52, Q63): one order/family -> group map, and
   backfill_taxon_groups.py --relabel-only.
B2 "Date before 1700" in the cleaning report (Q65), for observations, the scheme and
   contributed records.
B3 scheme iRecord / NBN copies (F53, Q59): shared/scheme_duplicates.py, the marking script
   (dry run, --apply, --undo), and the one rule through shared/species_counting and the
   Scheme tab ("Show duplicates").
B4 Sundon reference figures (Q66) and the local Pantheon bridge correction for
   Sympetrum nigrescens (F54, Q67).

Offscreen. The real data copies are only read; writes go to small temp databases or a
backup-API copy in tmp_path.
"""
import json
import os
import sqlite3

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import paths  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HAVE_DATA = all(p.exists() for p in (paths.OBSERVATUM_DB, paths.UKSI_DB))
HAVE_CODEX = HAVE_DATA and paths.CODEX_DB.exists() and paths.PANTHEON_DB.exists()
needs_data = pytest.mark.skipif(not HAVE_DATA, reason="real data copies not present")
needs_codex = pytest.mark.skipif(not HAVE_CODEX, reason="codex / pantheon copies not present")

BUTTERFLY_FAMILIES = ("Hesperiidae", "Papilionidae", "Pieridae", "Nymphalidae", "Lycaenidae",
                      "Riodinidae")


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


# ================================================================== B1 taxon groups

def test_butterfly_families_get_butterfly_moths_get_moth():
    from shared.taxon_groups import BUTTERFLY, MOTH, taxon_group
    for fam in BUTTERFLY_FAMILIES:
        assert taxon_group("Lepidoptera", fam) == BUTTERFLY, fam
    for fam in ("Noctuidae", "Geometridae", "Crambidae", "Zygaenidae", ""):
        assert taxon_group("Lepidoptera", fam) == MOTH, fam


def test_no_second_order_to_group_map_in_app_code():
    """Every importer, Data Entry commit and Quick Entry use shared/taxon_groups.py: the
    literal 'insect - moth' appears nowhere else in application code."""
    hits = []
    for top in ("Observatum", "DataEntry", "shared", "Examen", "Codex"):
        for dirpath, dirs, files in os.walk(os.path.join(ROOT, top)):
            dirs[:] = [d for d in dirs if d not in ("__pycache__", "_archive")]
            for f in files:
                if f.endswith(".py"):
                    p = os.path.join(dirpath, f)
                    with open(p, encoding="utf-8", errors="replace") as fh:
                        if "insect - moth" in fh.read():
                            hits.append(os.path.relpath(p, ROOT))
    assert hits == [os.path.join("shared", "taxon_groups.py")]


def _obs_db(path):
    c = sqlite3.connect(path)
    c.execute("CREATE TABLE observations (id INTEGER PRIMARY KEY, species_tvk TEXT, source TEXT,"
              " order_name TEXT, family TEXT, species_name TEXT, taxon_group TEXT, kingdom TEXT)")
    return c


@needs_data
def test_backfill_relabel_only_fixes_our_butterflies_and_fills_nothing(tmp_path):
    from scripts import backfill_taxon_groups as b
    db = str(tmp_path / "obs.db")
    c = _obs_db(db)
    u = sqlite3.connect(f"file:{paths.UKSI_DB}?mode=ro", uri=True)
    pieris = u.execute("SELECT tvk FROM taxa WHERE scientific_name = 'Pieris brassicae' "
                       "AND rank = 'Species'").fetchone()[0]
    noctua = u.execute("SELECT tvk FROM taxa WHERE scientific_name = 'Noctua pronuba' "
                       "AND rank = 'Species'").fetchone()[0]
    u.close()
    c.executemany("INSERT INTO observations VALUES (?,?,?,?,?,?,?,?)", [
        (1, pieris, None, "Lepidoptera", "Pieridae", "Pieris brassicae", "insect - moth", "Animalia"),
        (2, pieris, "iRecord | x", "Lepidoptera", "Pieridae", "Pieris brassicae", "insect - moth", None),
        (3, noctua, None, "Lepidoptera", "Noctuidae", "Noctua pronuba", "insect - moth", None),
        (4, pieris, None, "Lepidoptera", "Pieridae", "Pieris brassicae", None, None),
    ])
    c.commit()
    c.close()
    bk = str(tmp_path / "bk")
    assert b.main(["--relabel-only", "--tables", "observations", "--db", db, "--backup-dir", bk]) == 0
    c = sqlite3.connect(db)
    assert c.execute("SELECT taxon_group FROM observations WHERE id = 1").fetchone()[0] == "insect - moth"
    c.close()
    b.main(["--relabel-only", "--tables", "observations", "--db", db, "--backup-dir", bk, "--apply"])
    c = sqlite3.connect(db)
    got = dict(c.execute("SELECT id, taxon_group FROM observations").fetchall())
    kingdom4 = c.execute("SELECT kingdom FROM observations WHERE id = 4").fetchone()[0]
    c.close()
    assert got == {1: "insect - butterfly", 2: "insect - moth", 3: "insect - moth", 4: None}
    assert kingdom4 is None                                # nothing filled
    assert os.listdir(bk)                                  # a backup was taken


@needs_data
def test_backfill_relabel_dry_run_on_copy_is_the_42_butterflies():
    from scripts import backfill_taxon_groups as b
    ro = sqlite3.connect(f"file:{paths.OBSERVATUM_DB}?mode=ro", uri=True)
    try:
        found = b.plan(ro, paths.UKSI_DB, ("observations",), ("taxon_group",))
        fam = dict(ro.execute("SELECT id, family FROM observations").fetchall())
    finally:
        ro.close()
    rel = found["observations"]["relabel"]
    # Before the relabel is applied: 42 on the 10 Oct copy. After it: none left. Either way
    # every relabel is moth -> butterfly on a butterfly family, and nothing else.
    assert all(old == "insect - moth" and new == "insect - butterfly"
               for _i, old, new in rel)
    assert all(fam[i] in BUTTERFLY_FAMILIES for i, _o, _n in rel)
    assert len(rel) <= 60


# ================================================================== B2 date before 1700

def test_early_date_problem():
    from Observatum.src.services.cleaning_report import early_date_problem
    assert early_date_problem("1500-01-01") == "Dated 1500 (before 1700)"
    assert early_date_problem("1675") == "Dated 1675 (before 1700)"
    assert early_date_problem("1700-01-01") is None
    assert early_date_problem("2020-05-01") is None
    assert early_date_problem("") is None and early_date_problem(None) is None


def _cleaning_db(path):
    c = sqlite3.connect(path)
    for t in ("observations", "recording_scheme"):
        c.execute(f"CREATE TABLE {t} (id INTEGER PRIMARY KEY, irecord_id INTEGER, species_name TEXT,"
                  " species_tvk TEXT, taxon_rank TEXT, date TEXT, date_type TEXT, grid_ref TEXT,"
                  " site_name TEXT, recorder TEXT, latitude REAL, longitude REAL, source TEXT)")
    c.execute("CREATE TABLE contributed_observations (id INTEGER PRIMARY KEY, species_name TEXT,"
              " species_tvk TEXT, taxon_rank TEXT, date TEXT, date_type TEXT, grid_ref TEXT,"
              " site_name TEXT, recorder TEXT, contributor TEXT)")
    row = (None, "Alosterna tabacicolor", "", "Species", "1500-01-01", "D", "SP4512", "Wood", "X")
    c.execute("INSERT INTO recording_scheme VALUES (1, 77, ?,?,?,?,?,?,?,?, NULL, NULL, 'ds')", row[1:])
    c.execute("INSERT INTO recording_scheme VALUES (2, NULL, 'Rutpela maculata', '', 'Species',"
              " '2020-06-01', 'D', 'SP4512', 'Wood', 'X', NULL, NULL, 'ds')")
    c.execute("INSERT INTO observations VALUES (1, NULL, 'Clytus arietis', '', 'Species',"
              " '1675-01-01', 'D', 'SP4512', 'Wood', 'X', NULL, NULL, NULL)")
    c.execute("INSERT INTO contributed_observations VALUES (1, 'Clytus arietis', '', 'Species',"
              " '1690-05-01', 'D', 'SP4512', 'Wood', 'X', 'Moore, J.')")
    c.commit()
    c.close()
    return path


@needs_data
def test_cleaning_report_lists_early_dates_in_all_three_sources(tmp_path):
    from Observatum.src.services.cleaning_report import PROBLEM_LABELS, SOURCES, build_report
    assert PROBLEM_LABELS["early"] == "Date before 1700"
    assert "contributed_observations" in dict(SOURCES)
    db = _cleaning_db(str(tmp_path / "c.db"))
    got = {}
    for src, _label in SOURCES:
        rep = build_report(db, paths.UKSI_DB, src)
        got[src] = [(r.record_id, r.detail, r.irecord_id, r.dataset) for r in rep.rows["early"]]
    assert got["recording_scheme"] == [(1, "Dated 1500 (before 1700)", 77, "ds")]
    assert got["observations"] == [(1, "Dated 1675 (before 1700)", None, "")]
    assert got["contributed_observations"] == [(1, "Dated 1690 (before 1700)", None, "Moore, J.")]


@needs_data
def test_cleaning_export_has_the_early_sheet(tmp_path):
    from Observatum.src.services.cleaning_export import SHEET_NAMES, export_xlsx, have_openpyxl
    from Observatum.src.services.cleaning_report import build_report
    assert SHEET_NAMES["early"] == "Date before 1700"
    if not have_openpyxl():
        pytest.skip("openpyxl not installed")
    import openpyxl
    db = _cleaning_db(str(tmp_path / "c.db"))
    out = export_xlsx(build_report(db, paths.UKSI_DB, "recording_scheme"), str(tmp_path / "r.xlsx"))
    wb = openpyxl.load_workbook(out)
    rows = list(wb["Date before 1700"].iter_rows(values_only=True))
    assert rows[1][0] == "Alosterna tabacicolor" and rows[1][8].endswith("occurrence_id=77")
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]}
    assert summary["Date before 1700"] == 1


@needs_data
def test_cleaning_dashboard_early_tile_and_contributed_source(qapp, tmp_path):
    from Observatum.src.views.stats.cleaning_dashboard import CleaningDashboard
    db = _cleaning_db(str(tmp_path / "c.db"))
    dash = CleaningDashboard(main_db=db, uksi_db=str(paths.UKSI_DB))
    opened, shown = [], []
    dash.open_url = lambda url: opened.append(url.toString())
    dash.open_in_observatum = lambda row: shown.append(row.record_id)
    dash.cmb_source.setCurrentIndex(dash.cmb_source.findData("recording_scheme"))
    dash.refresh()
    assert dash._tiles["early"].text() == "Date before 1700\n1"
    view = dash._views["early"]
    dash._on_double_click(view.model().index(0, 0))
    assert opened == [] and len(shown) == 1                # 13 Oct 2026: detail, not iRecord
    dash.cmb_source.setCurrentIndex(dash.cmb_source.findData("contributed_observations"))
    dash.refresh()
    assert dash._tiles["early"].text() == "Date before 1700\n1"
    dash._on_double_click(dash._views["early"].model().index(0, 0))
    assert len(shown) == 2 and shown[1] == 1               # opens in Observatum


@needs_data
def test_cleaning_report_real_copy_known_early_records():
    from Observatum.src.services.cleaning_report import build_report
    rep = build_report(paths.OBSERVATUM_DB, paths.UKSI_DB, "recording_scheme")
    early = {(r.species, r.date[:4]) for r in rep.rows["early"]}
    assert ("Alosterna tabacicolor", "1500") in early
    assert ("Acanthocinus aedilis", "1675") in early
    assert ("Agapanthia villosoviridescens", "1675") in early


# ================================================================== B3 scheme copies

SCHEME_COLS = ("id INTEGER PRIMARY KEY, irecord_id INTEGER, nbn_atlas_id TEXT, occurrence_id TEXT,"
               " source TEXT, dataset_name TEXT, species_name TEXT, species_tvk TEXT,"
               " taxon_rank TEXT, date TEXT, date_type TEXT, grid_ref TEXT, site_name TEXT,"
               " recorder TEXT, family TEXT, vc_number INTEGER, verification_status TEXT,"
               " common_name TEXT")


def _scheme_db(path):
    c = sqlite3.connect(path)
    c.execute(f"CREATE TABLE recording_scheme ({SCHEME_COLS})")
    rows = [
        # iRecord rows
        (1, 1001, None, None, "iRecord", None, "Rutpela maculata", "T1", "Species", "2018-06-22", "D", "TQ443222"),
        (2, 1002, None, None, "iRecord", None, "Clytus arietis", "T2", "Species", "2019-05-01", "D", "SP1"),
        # NBN copies of 1 and 2 (grid ref blurred on 4)
        (3, None, "nbn-a", "1001", None, "Longhorn beetle (Cerambycidae) data from iRecord",
         "Rutpela maculata", "T1", "Species", "2018-06-22", "D", "TQ443222"),
        (4, None, "nbn-b", "1002", None, "Longhorn beetle (Cerambycidae) data from iRecord",
         "Clytus arietis", "T2", "Species", "2019-05-01", "D", "SP12"),
        # NBN row with the occurrence id but another date: not a copy
        (5, None, "nbn-c", "1001", None, "Other", "Rutpela maculata", "T1", "Species",
         "2018-06-23", "D", "TQ443222"),
        # a plain NBN record
        (6, None, "nbn-d", "zzz", None, "Other", "Rutpela maculata", "T1", "Species",
         "2001-07-01", "D", "TQ4422"),
    ]
    c.executemany("INSERT INTO recording_scheme (id, irecord_id, nbn_atlas_id, occurrence_id,"
                  " source, dataset_name, species_name, species_tvk, taxon_rank, date, date_type,"
                  " grid_ref) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    c.commit()
    c.close()
    return path


def test_copies_rule(tmp_path):
    from shared.scheme_duplicates import copies, load_rows
    c = sqlite3.connect(_scheme_db(str(tmp_path / "s.db")))
    try:
        assert copies(load_rows(c)) == {3: 1, 4: 2}
    finally:
        c.close()


def test_report_script_uses_the_shared_rule(tmp_path):
    from scripts import report_scheme_duplicates as rep
    c = sqlite3.connect(_scheme_db(str(tmp_path / "s.db")))
    rows = rep.load(c)
    c.close()
    pairs = rep.find_pairs(rows)
    same = sorted(k for k, v in pairs.items() if v[0] == "same iRecord id")
    assert same == [(1, 3), (2, 4)]


def test_mark_script_dry_run_apply_undo(tmp_path, capsys):
    from scripts import mark_scheme_irecord_duplicates as m
    from shared import scheme_duplicates as sd
    db = _scheme_db(str(tmp_path / "s.db"))
    bk = str(tmp_path / "bk")
    assert m.main(["--db", db, "--backup-dir", bk]) == 0
    out = capsys.readouterr().out
    assert "Nothing has been changed" in out and "to mark now (not marked yet): 2" in out
    c = sqlite3.connect(db)
    assert not m.has_column(c)                             # dry run: no column added
    c.close()
    assert not os.path.exists(bk)

    m.main(["--db", db, "--backup-dir", bk, "--apply"])
    c = sqlite3.connect(db)
    marks = dict(c.execute("SELECT id, duplicate_of FROM recording_scheme").fetchall())
    assert marks == {1: None, 2: None, 3: 1, 4: 2, 5: None, 6: None}
    assert c.execute("SELECT COUNT(*) FROM recording_scheme").fetchone()[0] == 6   # nothing deleted
    c.close()
    assert len(os.listdir(bk)) == 1
    sd.forget()
    assert sd.live_sql(db) == " AND duplicate_of IS NULL"
    assert sd.live_sql(db, "rs.") == " AND rs.duplicate_of IS NULL"

    m.main(["--db", db, "--backup-dir", bk, "--undo"])
    assert "Nothing has been changed" in capsys.readouterr().out
    m.main(["--db", db, "--backup-dir", bk, "--undo", "--apply"])
    c = sqlite3.connect(db)
    assert c.execute("SELECT COUNT(*) FROM recording_scheme WHERE duplicate_of IS NOT NULL"
                     ).fetchone()[0] == 0
    c.close()
    assert len(os.listdir(bk)) == 2


def test_live_sql_empty_without_the_column(tmp_path):
    from shared import scheme_duplicates as sd
    sd.forget()
    db = _scheme_db(str(tmp_path / "s.db"))
    assert sd.live_sql(db) == ""
    c = sqlite3.connect(db)
    assert sd.live_sql(c) == ""
    c.close()


def _marked_db(tmp_path):
    from scripts import mark_scheme_irecord_duplicates as m
    from shared import scheme_duplicates as sd
    db = _scheme_db(str(tmp_path / "s.db"))
    m.main(["--db", db, "--backup-dir", str(tmp_path / "bk"), "--apply"])
    sd.forget()
    return db


def test_counting_rule_leaves_copies_out(tmp_path):
    from shared import species_counting as sc
    db = _marked_db(tmp_path)
    assert sc.not_excluded_sql(db, "recording_scheme", "id", kinds=()).endswith(
        " AND duplicate_of IS NULL")
    assert sc.not_excluded_sql(db, "recording_scheme", "rs.id", kinds=()).endswith(
        " AND rs.duplicate_of IS NULL")
    assert "duplicate_of" not in sc.not_excluded_sql(db, "observations", "id", kinds=())
    assert sc.export_sql(db, "recording_scheme", sc.EXPORT_INCLUDE) == " AND duplicate_of IS NULL"
    assert sc.export_sql(db, "recording_scheme", sc.EXPORT_INCLUDE, include_copies=True) == ""
    c = sqlite3.connect(db)
    n = c.execute("SELECT COUNT(*) FROM recording_scheme WHERE 1"
                  + sc.not_excluded_sql(db, "recording_scheme", "id", kinds=())).fetchone()[0]
    c.close()
    assert n == 4


def test_map_counting_exclusions_include_copies_rule(tmp_path):
    from Observatum.src.services.map_selection import counting_exclusions
    db = _marked_db(tmp_path)
    c = sqlite3.connect(db)
    c.execute("CREATE TABLE observations (id INTEGER PRIMARY KEY, species_name TEXT,"
              " taxon_rank TEXT, species_tvk TEXT)")
    c.commit()
    c.close()
    assert "duplicate_of IS NULL" in counting_exclusions(db)["recording_scheme"]


def test_scheme_worker_hides_copies_unless_asked(qapp, tmp_path):
    from Observatum.src.views.scheme.scheme_data_worker import SchemeDataWorker
    db = _marked_db(tmp_path)
    w = SchemeDataWorker(db, {})
    w.run()
    assert sorted(r["id"] for r in w.results[0]) == [1, 2, 5, 6] and w.counted == 4
    w = SchemeDataWorker(db, {"show_duplicates": True})
    w.run()
    assert len(w.results[0]) == 6 and w.counted == 4
    assert w.dimmed == {3: "Duplicate of #1", 4: "Duplicate of #2"}


def test_scheme_model_duplicate_column(qapp):
    from PySide6.QtCore import Qt
    from Observatum.src.views.scheme.scheme_record_model import SchemeRecordModel
    model = SchemeRecordModel()
    model.set_records([{"id": 1, "species_name": "A b"}], {})
    assert "duplicate_of" not in [c[0] for c in model.get_columns()]
    model.set_records([{"id": 1, "species_name": "A b", "duplicate_of": None},
                       {"id": 3, "species_name": "A b", "duplicate_of": 1}], {3: "Duplicate of #1"})
    keys = [c[0] for c in model.get_columns()]
    col = keys.index("duplicate_of")
    assert model.headerData(col, Qt.Orientation.Horizontal) == "Duplicate of"
    assert model.data(model.index(1, col)) == "#1"
    assert model.data(model.index(0, col)) == ""
    assert "iRecord record #1" in model.data(model.index(1, col), Qt.ItemDataRole.ToolTipRole)


def test_scheme_filter_bar_show_duplicates_toggle(qapp):
    from Observatum.src.views.scheme.scheme_filter_bar import SchemeFilterBar
    bar = SchemeFilterBar()
    got = []
    bar.filters_changed.connect(got.append)
    assert bar.show_dupes_chk.text() == "Show duplicates"
    assert bar.get_filters()["show_duplicates"] is False
    bar.show_dupes_chk.setChecked(True)
    assert got and got[-1]["show_duplicates"] is True
    bar.clear_filters()
    assert bar.get_filters()["show_duplicates"] is False


@needs_data
def test_real_copy_dry_run_finds_the_irecord_copies():
    from scripts import mark_scheme_irecord_duplicates as m
    ro = sqlite3.connect(f"file:{paths.OBSERVATUM_DB}?mode=ro", uri=True)
    try:
        p = m.plan(ro)
    finally:
        ro.close()
    assert 30000 <= len(p["copies"]) <= 40000              # 34,121 on the 10 Oct copy
    assert p["datasets"].most_common(1)[0][0] == "Longhorn beetle (Cerambycidae) data from iRecord"


# ================================================================== B4 Sundon / Sympetrum

def test_sundon_reference_figures():
    with open(os.path.join(ROOT, "scripts", "reference_figures.json"), encoding="utf-8") as f:
        ref = json.load(f)
    assert ref["surveys"]["Sundon 2026"] == {"species": 231, "key": 19, "sqi": 125,
                                             "sqi_pantheon": 114, "sqi_strict": 114,
                                             "key_strict": 19}


def test_bridge_corrections_file():
    from shared.repositories.pantheon_repository import LOCAL_CORRECTION, load_corrections
    rows = load_corrections()
    pairs = {(p, u) for _a, p, u, _n in rows}
    assert ("NBNSYS0000005646", "NBNSYS0000005645") in pairs
    assert ("NBNSYS0000167746", "NBNSYS0000005645") in pairs
    assert all(n.startswith(LOCAL_CORRECTION) for *_x, n in rows)


@needs_codex
def test_sympetrum_striolatum_pantheon_only_not_key_and_flagged():
    from shared.repositories.codex_repository import AnalysisMode, CodexRepository
    from shared.repositories.pantheon_repository import PantheonRepository
    striolatum = "NBNSYS0000005645"
    cands = PantheonRepository().candidates(striolatum)
    assert "NBNSYS0000005646" not in cands and "NBNSYS0000167746" not in cands
    assert striolatum in cands
    st = CodexRepository().get_statuses_batch([striolatum], AnalysisMode.PANTHEON_ONLY)[striolatum]
    assert not st.is_key
    assert "local correction" in st.status_note
    assert st.sqs == 1
    # Codex Full is untouched by the correction
    full = CodexRepository().get_statuses_batch([striolatum], AnalysisMode.CODEX_FULL)[striolatum]
    assert "local correction" not in full.status_note
