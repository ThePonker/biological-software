"""Photo review sheet and its apply script (14 Oct 2026): small purpose-made data only."""
import sqlite3

import pytest
from openpyxl import load_workbook

import paths
from scripts import apply_photo_review as ap
from scripts import photo_review_sheet as prs
from shared import photo_links as pl

D = "Photos/Microscope Photos/Coleoptera/Carabidae"


def _spec(c, sid, sex, site="Bicester", grid="SP5822", coll="WJH", drawer="4"):
    c.execute("INSERT INTO specimens VALUES (?,?,?,?,?,?,?,?)",
              (sid, f"S{sid}", sex, site, grid, coll, drawer, "2026-06-05"))


@pytest.fixture
def world(tmp_path, monkeypatch):
    root = tmp_path / "Media"
    names = ["Amara aenea 1.jpg", "Amara aenea 2.jpg",                  # identical group (ids 1, 2)
             "Bembidion male gen.jpg", "Bembidion plain.jpg",           # differing (ids 3, 4)
             "Carabus linked.jpg"]                                       # already linked
    for n in names:
        p = root / D / n
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"x")
    db = tmp_path / "o.db"
    c = sqlite3.connect(db)
    c.execute("CREATE TABLE specimens (id INTEGER PRIMARY KEY, specimen_code TEXT, sex TEXT, "
              "site_name TEXT, grid_ref TEXT, collector TEXT, drawer_number TEXT, date_collected TEXT)")
    for sid in (1, 2):
        _spec(c, sid, "Male")
    _spec(c, 3, "Male")
    _spec(c, 4, "Female", site="Otmoor")
    _spec(c, 5, "Male")
    _spec(c, 6, "Female")
    pl.ensure_schema(c)
    pl.add_link(c, f"{D}/Carabus linked.jpg", specimen_id=5)
    c.commit()
    c.close()
    monkeypatch.setattr(paths, "MEDIA_ROOT", root)

    def row(name, ids, sp):
        return {"rel_path": f"{D}/{name}", "specimen_count": len(ids), "date": "2026-06-05",
                "uksi_name": sp, "binomial": sp,
                "specimen_ids": ";".join(f"{i}:S{i}" for i in ids)}
    scan = [row("Amara aenea 2.jpg", [1, 2], "Amara aenea"), row("Amara aenea 1.jpg", [1, 2], "Amara aenea"),
            row("Bembidion male gen.jpg", [3, 4], "Bembidion x"), row("Bembidion plain.jpg", [3, 4], "Bembidion x"),
            row("Carabus linked.jpg", [5, 6], "Carabus x"),
            {"rel_path": f"{D}/single.jpg", "specimen_count": 1, "specimen_ids": "9:S9"}]
    return root, db, scan


def _build(world):
    root, db, scan = world
    ro = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    linked = {r[0] for r in ro.execute("SELECT photo_path FROM photo_links WHERE specimen_id IS NOT NULL")}
    details = prs.load_details(ro, range(1, 10))
    ro.close()
    return prs.build_rows(scan, details, linked, root)


def test_pairing_sex_and_exclusion(world):
    rows, groups = _build(world)
    by = {r["Photo"]: r for r in rows}
    assert "Carabus linked.jpg" not in by and len(rows) == 4 and groups == 1
    assert (by["Amara aenea 1.jpg"]["Link to specimen id"], by["Amara aenea 2.jpg"]["Link to specimen id"]) == (1, 2)
    assert by["Amara aenea 1.jpg"]["Why"] == prs.IDENTICAL
    m = by["Bembidion male gen.jpg"]
    assert (m["Link to specimen id"], m["Why"], m["Photo sex"], m["Photo view"]) == (3, prs.SEX_MATCH, "male", "gen")
    assert by["Bembidion plain.jpg"]["Link to specimen id"] == ""
    assert dict(prs.counts(rows, groups))["Left for Wil"] == 1


def test_sheet_written_and_apply(world, tmp_path):
    root, db, _ = world
    rows, groups = _build(world)
    x = tmp_path / "photo_review.xlsx"
    prs.write_xlsx(rows, prs.counts(rows, groups), x)
    wb = load_workbook(x)
    assert wb.sheetnames == ["Review", "Summary"] and wb["Review"].freeze_panes == "A2"
    with pytest.raises(FileExistsError):                    # never over a filled-in sheet
        prs.write_xlsx(rows, prs.counts(rows, groups), x)
    # Wil fills the blank row with a non-candidate (bad) and edits nothing else
    ws = wb["Review"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value == "Bembidion plain.jpg":
            ws.cell(r, 8).value = 99
    wb.save(x)
    args = ["--sheet", str(x), "--root", str(root), "--db", str(db), "--backup-dir", str(tmp_path / "bk")]

    def n():
        c = sqlite3.connect(db)
        try:
            return c.execute("SELECT COUNT(*) FROM photo_links").fetchone()[0]
        finally:
            c.close()
    assert ap.main(args) == 0 and n() == 1                 # dry run writes nothing
    assert not (tmp_path / "bk").exists()
    ap.main(args + ["--apply"])
    assert n() == 4                                         # 1 + (1, 2, 3); id 99 refused
    c = sqlite3.connect(db)
    got = dict(c.execute("SELECT photo_path, specimen_id FROM photo_links WHERE specimen_id != 5"))
    assert got == {f"{D}/Amara aenea 1.jpg": 1, f"{D}/Amara aenea 2.jpg": 2, f"{D}/Bembidion male gen.jpg": 3}
    assert len(list((tmp_path / "bk").glob("observatum_pre_photo_review_*.db"))) == 1


def test_apply_problems(world):
    root, _db, _ = world
    rows = [(2, f"{D}/Amara aenea 1.jpg", [1, 2], 2), (3, f"{D}/missing.jpg", [1], 1),
            (4, f"{D}/Carabus linked.jpg", [5], 5), (5, f"{D}/Amara aenea 2.jpg", [1, 2], "x"),
            (6, f"{D}/Amara aenea 2.jpg", [1, 2], None)]
    links, problems, blank = ap.plan(rows, root, {f"{D}/Carabus linked.jpg"})
    assert links == [(f"{D}/Amara aenea 1.jpg", 2)] and len(problems) == 3 and blank == 1
