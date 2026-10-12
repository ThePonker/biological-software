"""Lector stage A (12 Oct 2026): fix2 restored, review findings LEC1-6, the
saved-search data layer, migration of lector.db and the service API.

No network: BHL is a FakeBHL test double. Databases are temp files; the real
data/lector.db and data/uksi.db are only read (copied with conn.backup()).
"""
import json
import sqlite3
import threading
from pathlib import Path

import pytest

from Lector import config, engine, saved, schema, upgrade
from Lector.bhl_client import BHLError
from Lector.excerpts import (classify, find_excerpts, group_duplicates, looks_like_index,
                             name_pattern, split_marked)
from Lector.harvest import HarvestOptions, apply_filters, extract_pages, first_year, title_has
from Lector.names import find_names
from Lector.service import LectorError, LectorService, NeedsUpdate
from Lector.species_source import UksiLookup
from Lector.store import LectorStore
from Lector.synonyms import SynonymRules, candidate_names

ROOT = Path(__file__).resolve().parents[1]
TEXTOR, SUTOR, SARTOR = "NBNSYS0000011050", "NHMSYS0020704741", "NHMSYS0020704740"

RECORD = ("THE ENTOMOLOGIST'S MONTHLY MAGAZINE.\nNotes on Coleoptera taken during the past season.\n"
          "By the kindness of a friend I was able to work the\ndistrict thoroughly. Lamia\n"
          "textor. I took a specimen\nof this species in the Black Wood, on a willow stump,\n"
          "in June last; it is, I believe, a new locality for it.\n"
          "Saperda carcharias was also common on aspen.\nE. A. Waterhouse, Ripon.")
ANNALS = ("On some British Coleoptera. The longhorn beetles of the northern counties were\n"
          "examined with care during the summer months and several species were taken\n"
          "in abundance under the bark of old willows near the river at Rannoch where\n"
          "Lamia textor occurred sparingly in company with Saperda populnea and the\n"
          "common Rhagium inquisitor, all of them on felled timber left by the woodmen\n"
          "through the winter, which is the best place to search for them in June.")
INDEX = "\n".join(f"{w}, {n}" for n, w in enumerate(
    "tabida textor tenebrioides testacea testaceus thoracica tibialis tomentosa "
    "trifasciatus tristis typhae umbellatarum undata varius".split(), 12))
AD = ("Dealer in natural history specimens. Coleoptera for sale: Lamia textor 2 fr.,\n"
      "Carabus auratus 1 fr., Cetonia aurata 50 c. Write for lists to the office.")


def _page(pid, num, ptype="Text"):
    return {"PageID": pid, "PageNumbers": [{"Prefix": "Page", "Number": str(num)}] if num else [],
            "PageTypes": [{"PageTypeName": ptype}]}


def bhl_data():
    """name -> GetNameMetadata Result; page_id -> OCR text."""
    textor = [{"NameConfirmed": "Lamia textor", "Titles": [
        {"TitleID": 1, "FullTitle": "The Entomologist's monthly magazine", "Items": [
            {"ItemID": 10, "Volume": "v.7 (1870-1871)", "Year": "1870", "Parts": [
                {"PartID": 5, "Title": "Notes on Coleoptera taken during the past season",
                 "Authors": [{"Name": "Waterhouse, E. A."}], "Pages": [_page(101, 81)]}]},
            {"ItemID": 11, "Volume": "v.8 (1872)", "Pages": [_page(105, 301, "Index")]}]},
        {"TitleID": 2, "FullTitle": "The Annals and magazine of natural history", "Items": [
            {"ItemID": 20, "Volume": "7", "Year": "1841", "Pages": [_page(102, 145)]},
            {"ItemID": 22, "Volume": "v.7=no.41-47 (1841)", "Pages": [_page(104, 145)]}]},
        {"TitleID": 3, "FullTitle": "Annals of natural history", "Items": [
            {"ItemID": 21, "Volume": "v.7(1841)", "Pages": [_page(103, 145)]}]},
        {"TitleID": 4, "FullTitle": "Le Naturaliste", "Items": [
            {"ItemID": 30, "Volume": "14", "Year": "1900", "Pages": [_page(106, 38), _page(107, 50)]}]},
        {"TitleID": 5, "FullTitle": "The Canadian entomologist", "Items": [
            {"ItemID": 40, "Volume": "v. 52 (1920)", "Pages": [_page(108, 49)]}]},
        {"TitleID": 6, "FullTitle": "The Entomologist's record and journal of variation", "Items": [
            {"ItemID": 50, "Volume": "no.1545", "Date": "1935", "Pages": [_page(109, 12)]}]},
    ]}]
    cerambyx = [{"NameConfirmed": "Cerambyx textor", "Titles": [
        {"TitleID": 7, "FullTitle": "Transactions of the Entomological Society of London",
         "Items": [{"ItemID": 60, "Volume": "v.2 (1837-1840)", "Pages": [_page(201, 12)]}]}]}]
    texts = {101: RECORD, 102: ANNALS, 103: ANNALS.replace("Rannoch", "Rannoeh"),
             104: ANNALS.replace("Lamia textor", "Lamia tcxtor").replace("summer", "sumnier"),
             105: INDEX, 106: AD, 107: AD, 108: "Lamia textor in Ontario.",
             109: "Lamia textor taken at Aviemore in 1934.",
             201: "Cerambyx textor Linn. was formerly taken near London on willows."}
    return {"Lamia textor": textor, "Cerambyx textor": cerambyx}, texts


class FakeBHL:
    """BHLClient test double: counts requests, can fail pages, never sleeps."""
    def __init__(self, names=None, texts=None, fail_pages=()):
        d_names, d_texts = bhl_data()
        self.names, self.texts = names or d_names, texts or d_texts
        self.fail_pages = set(fail_pages)
        self.request_count, self._interval = 0, 0
        self.asked = []

    def name_metadata(self, name):
        self.request_count += 1
        self.asked.append(("name", name))
        return self.names.get(name, [])

    def page_metadata(self, pid):
        self.request_count += 1
        self.asked.append(("page", pid))
        if pid in self.fail_pages:
            raise BHLError(f"GetPageMetadata failed after 4 attempts (HTTP 500) [{pid}]")
        return {"PageID": pid, "OcrText": self.texts.get(pid, ""), "TextSource": "OCR",
                "Names": [{"NameCanonical": "Lamia textor"}]}


def make_uksi(path):
    c = sqlite3.connect(path)
    c.executescript("""
        CREATE TABLE taxa (tvk TEXT PRIMARY KEY, scientific_name TEXT, rank TEXT, family TEXT,
                           genus TEXT, sort_code INTEGER);
        CREATE TABLE synonyms (synonym TEXT, tvk TEXT);
        CREATE TABLE name_map (tvk TEXT PRIMARY KEY, recommended_tvk TEXT, name TEXT,
                               recommended_name TEXT, name_status TEXT);""")
    c.executemany("INSERT INTO taxa VALUES (?,?,?,?,?,?)", [
        (TEXTOR, "Lamia textor", "Species", "Cerambycidae", "Lamia", 1),
        (SUTOR, "Monochamus sutor", "Species", "Cerambycidae", "Monochamus", 2),
        (SARTOR, "Monochamus sartor", "Species", "Cerambycidae", "Monochamus", 3),
        ("NHMSYS0099000001", "Cerambycidae", "Family", "Cerambycidae", "", 0),
        ("NHMSYS0099000002", "Rhagium mordax", "Species", "Cerambycidae", "Rhagium", 4)])
    c.executemany("INSERT INTO synonyms VALUES (?,?)",
                  [("Cerambyx textor", TEXTOR), ("Lamia sutor", TEXTOR)])
    c.executemany("INSERT INTO name_map VALUES (?,?,?,?,?)", [
        ("NBNSYS0000149844", TEXTOR, "Lamia sartor", "Lamia textor", "S"),
        ("NBNSYS0000149846", TEXTOR, "Lamia titillator", "Lamia textor", "S"),
        ("NBNSYS0000169525", TEXTOR, "Pine Sawyer", "Lamia textor", "S")])
    c.commit()
    c.close()


@pytest.fixture
def env(tmp_path):
    uksi = tmp_path / "uksi.db"
    make_uksi(uksi)
    fake = FakeBHL()
    svc = LectorService(db_path=tmp_path / "lector.db", uksi_db=uksi, client_factory=lambda: fake,
                        export_dir=tmp_path / "exports", synonym_file=tmp_path / "syn.txt")
    return svc, fake, tmp_path


BRITISH = {"include_words": ["Entomologist", "Annals", "Transactions"],
           "exclude_words": ["Canadian"], "max_pages": 0}


# --- fix2 restored --------------------------------------------------------------
def test_title_filter_whole_words_with_plurals():
    assert title_has(["Naturalist"], "Proceedings of the Bristol Naturalists' Society")
    assert title_has(["Entomologist"], "The Entomologist's monthly magazine")
    assert not title_has(["Naturalist"], "Le Naturaliste")
    assert not title_has(["Naturalist"], "Il Naturalista siciliano")


def test_extract_pages_keeps_page_type_authors_and_article_title():
    pages = {p["page_id"]: p for p in extract_pages(bhl_data()[0]["Lamia textor"])}
    assert pages[105]["page_type"] == "Index" and pages[101]["page_type"] == "Text"
    assert pages[101]["authors"] == "Waterhouse, E. A."
    assert pages[101]["article_title"].startswith("Notes on Coleoptera")
    assert pages[101]["title"] == "The Entomologist's monthly magazine" and pages[101]["year"] == 1870


def test_excerpt_bolds_name_split_over_two_lines_and_index_detection():
    ex = find_excerpts(RECORD, ["Lamia textor"])
    assert len(ex) == 1 and any("Lamia" in ln for ln in ex[0]) and any("textor" in ln for ln in ex[0])
    runs = split_marked("\n".join(ex[0]), name_pattern(["Lamia textor"]))
    assert [t for t, bold in runs if bold] == ["Lamia\ntextor"]
    assert looks_like_index(INDEX) and not looks_like_index(RECORD * 3)
    assert classify({"page_type": "Index", "ocr_text": "x"}, []) == "index"
    assert classify({"page_type": "Text", "ocr_text": "plate"}, []) == "no_match"


# --- LEC4 years ----------------------------------------------------------------
@pytest.mark.parametrize("text,year", [
    ("1841-1842", 1841), ("no.1545", None), ("t.32:no.1-2 (1859)", 1859),
    ("v.18-20 (1887-1890), Index v.11-20", 1887), ("ser.4:v.3-4=[v.20-21] (1910-1916)", 1910),
    ("new ser.:v.8 (1905-1908)", 1905), ("Nr. 1712 (1902)", 1902), ("pl. 1650", None),
    ("p.1800-1802 (1801)", 1801), ("1546", None), ("2099", None), ("anno 3 (1883-1884)", 1883)])
def test_year_parser(text, year):
    assert first_year(text) == year


def test_year_from_date_when_volume_is_an_issue_number():
    pages = {p["page_id"]: p for p in extract_pages(bhl_data()[0]["Lamia textor"])}
    assert pages[109]["year"] == 1935                 # 'no.1545' is not a year


# --- LEC6 names ------------------------------------------------------------------
def test_name_matching_case_spacing_accents_hyphenation():
    text = ("LAMIA TEXTOR; Lamia\n   textor; tex-\ntor; L. textor; Lami-\n a textor; "
            "Lamia rosenmüller; textorius; pretextor")
    found = [text[a:b] for a, b in find_names(text, ["Lamia textor", "Lamia rosenmuller"])]
    assert found == ["LAMIA TEXTOR", "Lamia\n   textor", "tex-\ntor", "L. textor",
                     "Lami-\n a textor", "Lamia rosenmüller"]
    ex = find_excerpts("a\nb\nnotes on tex-\ntor here\nc\nd\ne\nf", ["Lamia textor"])
    assert ex and "tor here" in ex[0]


# --- LEC2 duplicates --------------------------------------------------------------
class _E:
    def __init__(self, row):
        self.row, self.copies = row, []
        self.excerpts = find_excerpts(row["ocr_text"], ["Lamia textor"])
        self.section = classify(row, self.excerpts)


def _row(pid, item, title, vol, label, text, year=1841):
    return {"page_id": pid, "item_id": item, "title": title, "volume": vol, "page_label": label,
            "ocr_text": text, "year": year, "page_type": "Text"}


def test_duplicate_scans_grouped_across_titles_and_missed_names():
    data = bhl_data()[1]
    rows = [_row(102, 20, "The Annals and magazine of natural history", "7", "Page 145", data[102]),
            _row(103, 21, "Annals of natural history", "v.7(1841)", "Page 145", data[103]),
            _row(104, 22, "The Annals and magazine of natural history", "v.7=no.41-47 (1841)",
                 "Page 145", data[104])]
    kept = group_duplicates([_E(r) for r in rows])
    assert len(kept) == 1 and len(kept[0].copies) == 2 and kept[0].section == "main"


def test_repeated_text_in_one_volume_is_not_a_duplicate():
    rows = [_row(106, 30, "Le Naturaliste", "14", "Page 38", AD, 1900),
            _row(107, 30, "Le Naturaliste", "14", "Page 50", AD, 1900)]
    assert len(group_duplicates([_E(r) for r in rows])) == 2


def test_prefilter_dedupe_compares_volume_numbers():
    pages = [{"page_id": 1, "title": "Annals", "volume": "7", "page_label": "Page 145", "year": 1841},
             {"page_id": 2, "title": "Annals", "volume": "v.7 (1841)", "page_label": "p. 145",
              "year": 1841},
             {"page_id": 3, "title": "Annals", "volume": "", "page_label": "Page 145", "year": 1842}]
    kept = apply_filters(pages, HarvestOptions(max_pages=0))
    assert [p["page_id"] for p in kept] == [1, 3]


# --- LEC3 synonyms ---------------------------------------------------------------
def test_synonym_checks_hold_back_misapplied_names(tmp_path):
    make_uksi(tmp_path / "u.db")
    u = UksiLookup(tmp_path / "u.db")
    c = {x["name"]: x for x in candidate_names(u, "Lamia textor", TEXTOR, SynonymRules())}
    assert c["Lamia textor"]["origin"] == "accepted" and c["Cerambyx textor"]["included"]
    assert not c["Lamia sutor"]["included"] and c["Lamia sutor"]["conflict_tvk"] == SUTOR
    assert "Monochamus sartor" in c["Lamia sartor"]["held_back_reason"]
    assert c["Lamia titillator"]["included"]            # no current titillator: only a list stops it
    assert "Pine Sawyer" not in c                        # vernacular names never searched
    rules = SynonymRules.load(tmp_path / "syn.txt", suite_csv=ROOT / "scripts" / "uksi_synonym_corrections.csv")
    rules.remember("-", "Lamia titillator", "Lamia textor", "Nearctic Monochamus")
    rules.remember("+", "Lamia sutor", TEXTOR)          # '+' confirms a held-back name
    rules.remember("+", "Lamia rosenmulleri", "Lamia textor")
    rules = SynonymRules.load(tmp_path / "syn.txt", suite_csv=None)
    c = {x["name"]: x for x in candidate_names(u, "Lamia textor", TEXTOR, rules)}
    assert not c["Lamia titillator"]["included"] and "Nearctic" in c["Lamia titillator"]["held_back_reason"]
    assert c["Lamia sutor"]["included"] and c["Lamia rosenmulleri"]["origin"] == "user_added"
    assert (tmp_path / "syn.txt").read_text().startswith("# Lector synonym list")
    entry = u.resolve("Lamia textor", with_synonyms=True)
    assert "Lamia sartor" in dict(entry.held_back)
    u.close()


# --- saved searches, runs, LEC5 ------------------------------------------------------
def test_create_search_with_checked_synonyms(env):
    svc, _fake, _tmp = env
    sid = svc.create_search("lamia textor", filters=BRITISH)
    s = svc.get_search(sid)
    names = {n["name"]: n for n in s["names"]}
    assert s["target_tvk"] == TEXTOR and s["status"] == "draft"
    assert names["Cerambyx textor"]["included"] and not names["Lamia sutor"]["included"]
    assert names["Lamia sutor"]["held_back_reason"].startswith("Held back: epithet")
    assert s["filters"]["include_words"] == ["Entomologist", "Annals", "Transactions"]
    with pytest.raises(LectorError):
        svc.create_search("Lamia nonexistens")


def test_run_fetch_then_filters_changed_rerun_withdraws_without_asking_bhl(env):
    svc, fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    pre = svc.preview(sid)
    assert pre["not_counted"] == ["Lamia textor", "Cerambyx textor"]
    r1 = svc.run_search(sid)
    assert r1["status"] == "done" and r1["found"] == 10
    # 108 Canadian and 106/107 Le Naturaliste filtered; 102/103/104 kept (different volumes)
    assert r1["kept"] == 7 and r1["new"] == 7 and r1["text_fetched"] == 7
    assert r1["requests"] == 2 + 7
    svc.set_review(sid, 109, status="relevant", notes="Aviemore 1934")
    fake.asked.clear()
    svc.update_search(sid, filters={"to_year": 1930})
    r2 = svc.run_search(sid, refresh=False)
    assert fake.asked == []                       # cached BHL answers, no requests at all
    assert r2["withdrawn"] == 1 and r2["unchanged"] == 6 and r2["kept"] == 6
    held = {i["page_id"] for i in svc.review_queue(sid)}
    assert 109 not in held
    withdrawn = svc.review_queue(sid, include_withdrawn=True)
    p109 = next(i for i in withdrawn if i["page_id"] == 109)
    assert p109["state"] == "withdrawn" and p109["notes"] == "Aviemore 1934"
    diff = svc.run_diff(sid)
    assert [p["page_id"] for p in diff["withdrawn"]] == [109]
    assert diff["withdrawn"][0]["reason"].startswith("Now filtered out")
    assert any("to year" in x for x in diff["settings_changed"])
    svc.keep_page(sid, 109)
    assert 109 in {i["page_id"] for i in svc.review_queue(sid)}
    runs = svc.list_runs(sid)
    assert [r["mode"] for r in runs] == ["fetch", "fetch"] and runs[0]["withdrawn"] == 1
    assert json.dumps(runs[0]["settings"]["filters"]["to_year"]) == "1930"


def test_new_scan_found_on_refresh(env):
    svc, fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH, with_synonyms=False)
    svc.run_search(sid)
    fake.names["Lamia textor"][0]["Titles"][0]["Items"][1]["Pages"].append(_page(110, 302))
    fake.texts[110] = "More on Lamia textor at Rannoch."
    assert svc.run_search(sid, refresh=False)["new"] == 0
    r = svc.run_search(sid, refresh=True)
    assert r["new"] == 1 and svc.run_diff(sid)["new"][0]["page_id"] == 110


def test_progress_cancel_and_resume(env):
    svc, fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    stop, seen = threading.Event(), []

    def progress(stage, done, total, message):
        seen.append((stage, done, total))
        if stage == "text" and done == 2:
            stop.set()
    r = svc.run_search(sid, progress=progress, cancel=stop)
    assert r["status"] == "cancelled" and r["text_fetched"] == 2
    assert ("search", 0, 2) in seen and any(s[0] == "text" for s in seen)
    assert svc.list_runs(sid)[0]["status"] == "cancelled"
    r2 = svc.run_search(sid, refresh=False)
    assert r2["status"] == "done" and r2["text_fetched"] == 5 and r2["new"] == 0


def test_run_in_worker_thread_and_failed_pages_retry(env):
    svc, fake, _tmp = env
    fake.fail_pages = {102}
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    out = {}
    t = threading.Thread(target=lambda: out.update(r=svc.run_search(sid)))
    t.start()
    t.join(30)
    assert out["r"]["status"] == "done" and out["r"]["failed"] == 1
    fake.fail_pages = set()
    assert svc.run_search(sid, refresh=False)["text_fetched"] == 1


def test_review_queue_groups_duplicates_and_marks_index(env):
    svc, _fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    q = {i["page_id"]: i for i in svc.review_queue(sid)}
    group = [i for i in q.values() if i["copies"]]
    assert len(group) == 1 and {c["page_id"] for c in group[0]["copies"]} | {group[0]["page_id"]} \
        == {102, 103, 104}
    assert q[105]["likely_index"] and q[105]["index_reason"] == "BHL page type"
    rec = q[101]
    assert rec["authors"] == "Waterhouse, E. A." and "Waterhouse, E. A. (1870)" in rec["citation"]
    assert "E. A.. " not in rec["source"]
    assert any(bold for x in rec["excerpt_runs"] for _t, bold in x)
    assert q[109]["copyright_note"].startswith("May be in copyright")
    svc.set_review(sid, 101, status="record_extracted", tag="record", notes="Black Wood, Ripon?")
    assert svc.review_queue(sid, status="record_extracted")[0]["page_id"] == 101
    with pytest.raises(LectorError):
        svc.set_review(sid, 101, status="maybe")
    detail = svc.page_detail(sid, 101)
    assert detail["ocr_text"] == RECORD and detail["review_tag"] == "record"


def test_list_searches_dashboard_duplicate_and_presets(env):
    svc, _fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    svc.set_review(sid, 101, status="relevant")
    s = svc.list_searches()[0]
    assert (s["pages_held"], s["reviewed"]) == (7, 1)
    # sutor/sartor: epithet current in Monochamus; titillator: the suite's corrections list
    assert s["names_held_back"] == ["Lamia sutor", "Lamia sartor", "Lamia titillator"]
    d = svc.dashboard()
    assert d["pages_harvested"] == 7 and d["awaiting_review"] == 6 and d["activity"]
    copy = svc.duplicate_search(sid, "Lamia textor - Scottish")
    assert svc.get_search(copy)["status"] == "draft" and len(svc.list_searches()) == 2
    names = [p["name"] for p in svc.journal_presets()]
    assert names[:2] == ["British periodicals", "Scottish titles"]
    fam = svc.create_search("Cerambycidae", kind="family")
    assert {n["name"] for n in svc.get_search(fam)["names"]} == {
        "Lamia textor", "Monochamus sutor", "Monochamus sartor", "Rhagium mordax"}
    assert svc.resolve_taxon("Lamia sartor")["matched_as"] == "synonym"


def test_unrecorded_filters_block_a_fetch(env):
    svc, _fake, _tmp = env
    sid = svc.create_search("Lamia textor")
    with svc._store() as st:
        st.conn.execute("UPDATE search_filters SET recorded=0 WHERE search_id=?", (sid,))
    with pytest.raises(LectorError, match="Set its filters first"):
        svc.run_search(sid)
    svc.update_search(sid, filters={"max_pages": 50})
    assert svc.run_search(sid)["status"] == "done"


# --- LEC1 exports -----------------------------------------------------------------
def _docx_text(path):
    import docx
    d = docx.Document(str(path))
    body = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        body += "\n" + "\n".join(c.text for row in t.rows for c in row.cells)
    rels = [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype]
    return body, rels


def test_export_word_excel_text_with_links_authors_titles(env):
    pytest.importorskip("docx")
    pytest.importorskip("openpyxl")
    svc, _fake, tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    svc.set_review(sid, 109, status="not_relevant")
    svc.set_review(sid, 101, notes="Locality to confirm")
    out = svc.export(sid, formats=("txt", "docx", "xlsx"))
    assert set(out["files"]) == {"txt", "csv", "docx", "xlsx"}
    body, links = _docx_text(out["files"]["docx"])
    assert "Notes on Coleoptera taken during the past season" in body
    assert "Waterhouse" in body and "Same text also at" in body
    assert "Locality to confirm" in body and "Aviemore" not in body      # not_relevant left out
    assert "https://www.biodiversitylibrary.org/page/101" in links
    assert out["counts"]["duplicates_merged"] == 2 and out["counts"]["index_pages"] == 1
    import openpyxl
    ws = openpyxl.load_workbook(out["files"]["xlsx"]).active
    head = [c.value for c in ws[1]]
    assert "Authors" in head and "Citation" in head and ws.max_row == 1 + 6
    txt = Path(out["files"]["txt"]).read_text(encoding="utf-8")
    assert "Waterhouse, E. A. (1870)" in txt
    assert len(svc.list_exports(sid)) == 3
    out2 = svc.export(sid, formats=("docx",), order="journal",
                      include_statuses=("unreviewed", "relevant", "not_relevant"),
                      include_appendices=False, out_dir=tmp / "x")
    body2, _ = _docx_text(out2["files"]["docx"])
    assert "Aviemore" in body2 and "Appendix A" not in body2


def test_export_without_python_docx_says_so(env, monkeypatch):
    svc, _fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    import sys
    monkeypatch.setitem(sys.modules, "docx", None)
    out = svc.export(sid, formats=("docx", "xlsx", "pdf"))
    assert out["files"] == {} and "pip install python-docx" in out["messages"][0]


def test_pdf_reports_missing_converter(env, monkeypatch):
    pytest.importorskip("docx")
    from Lector import pdf
    svc, _fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    monkeypatch.setattr(pdf, "pdf_method", lambda: None)
    out = svc.export(sid, formats=("pdf",))
    assert "docx" in out["files"] and "pdf" not in out["files"]
    assert "LibreOffice" in out["messages"][-1]


def test_pdf_made_with_libreoffice_when_installed(env):
    pytest.importorskip("docx")
    from Lector import pdf
    if pdf.pdf_method() != "libreoffice":
        pytest.skip("LibreOffice not installed")
    svc, _fake, _tmp = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    out = svc.export(sid, formats=("pdf",))
    assert Path(out["files"]["pdf"]).read_bytes()[:4] == b"%PDF"


# --- CLI ------------------------------------------------------------------------------
@pytest.fixture
def cli(tmp_path, monkeypatch):
    make_uksi(tmp_path / "uksi.db")
    monkeypatch.setattr(config, "LECTOR_DB", tmp_path / "lector.db")
    monkeypatch.setattr(config, "UKSI_DB", tmp_path / "uksi.db")
    monkeypatch.setattr(config, "EXPORT_DIR", tmp_path / "exports")
    monkeypatch.setattr(config, "SYNONYM_FILE", tmp_path / "syn.txt")
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path / "_backups")
    fake = FakeBHL()
    import Lector.__main__ as m
    monkeypatch.setattr(m, "BHLClient", lambda key: fake)
    monkeypatch.setattr(config, "load_api_key", lambda key=None: "k")
    from Lector import species_source
    monkeypatch.setattr(species_source.UksiLookup.__init__, "__defaults__", (tmp_path / "uksi.db",))
    (tmp_path / "sp.txt").write_text("Lamia textor\n")
    return m, fake, tmp_path


def test_cli_fetch_status_export_forget_synonyms(cli, capsys):
    m, fake, tmp = cli
    m.main(["fetch", str(tmp / "sp.txt"), "--synonyms", "--include-title", "Entomologist",
            "Annals", "Transactions", "--exclude-title", "Canadian", "--max-pages", "0"])
    out = capsys.readouterr().out
    assert "Lamia sutor: Held back" in out and "✓ Cerambyx textor (synonym)" in out
    m.main(["fetch", str(tmp / "sp.txt"), "--synonyms", "--include-title", "Annals",
            "--max-pages", "0", "--replace-search"])   # the saved search takes these filters
    out = capsys.readouterr().out
    assert "--refresh asks BHL again" in out and "no longer within the filters" in out
    m.main(["status"])
    assert "Lamia textor" in capsys.readouterr().out
    m.main(["export", "--name", "Lamia textor"])
    assert (tmp / "exports" / "Lamia_textor_all_pages.txt").exists()   # not the window's name
    if pytest.importorskip("docx"):
        m.main(["export", "--name", "Lamia textor", "--docx", "--by-journal"])
        assert "duplicate scans merged" in capsys.readouterr().out
    m.main(["synonyms", "Lamia textor", "--exclude", "Lamia titillator", "--reason", "Nearctic"])
    out = capsys.readouterr().out
    assert "✗ Lamia titillator" in out and "Nearctic" in out
    st = LectorStore(tmp / "lector.db")
    st.conn.execute("UPDATE page_hits SET review_status='relevant' WHERE page_id=102")
    st.close()
    with pytest.raises(SystemExit, match="Nothing has been changed"):
        m.main(["forget", "--name", "Lamia textor"])
    m.main(["forget", "--name", "Lamia textor", "--yes"])
    st = LectorStore(tmp / "lector.db")
    assert st.conn.execute("SELECT COUNT(*) FROM page_hits").fetchone()[0] == 0
    assert st.conn.execute("SELECT COUNT(*) FROM pages WHERE ocr_text IS NOT NULL").fetchone()[0] > 0
    st.close()


# --- migration -------------------------------------------------------------------------
def _legacy_db(path):
    c = sqlite3.connect(path)
    c.executescript(schema.LEGACY)
    c.execute("INSERT INTO searches VALUES ('Lamia textor','Lamia textor',?,'accepted',"
              "'2026-10-07T22:20:21',1060,2)", (TEXTOR,))
    for pid in (1, 2):
        c.execute("INSERT INTO pages (page_id, title, year, page_url, ocr_text, fetched_at)"
                  " VALUES (?,?,?,?,?,?)", (pid, "Annals", 1841, f"u{pid}", "Lamia textor", "t"))
        c.execute("INSERT INTO page_hits (page_id, accepted_name, query_name, tvk, review_status)"
                  " VALUES (?,?,?,?,?)", (pid, "Lamia textor", "Lamia textor", TEXTOR,
                                          "relevant" if pid == 1 else "unreviewed"))
    c.commit()
    c.close()


def test_old_database_is_not_changed_without_asking(tmp_path):
    _legacy_db(tmp_path / "lector.db")
    svc = LectorService(db_path=tmp_path / "lector.db", uksi_db=tmp_path / "none.db",
                        client_factory=FakeBHL)
    with pytest.raises(NeedsUpdate):
        svc.list_searches()
    report = svc.needs_update()
    assert report["after"]["saved_searches"] == 1 and report["before"]["saved_searches"] is None
    assert upgrade.status(tmp_path / "lector.db")["version"] == 0      # dry run changed nothing
    done = svc.update_database()
    assert Path(done["backup"]).exists() and done["after"]["pages"] == done["before"]["pages"] == 2
    s = svc.list_searches()[0]
    assert (s["name"], s["pages_held"], s["reviewed"], s["filters_recorded"]) == \
        ("Lamia textor", 2, 1, False)
    assert svc.list_runs(s["search_id"])[0]["status"] == "migrated"
    with sqlite3.connect(tmp_path / "lector.db") as c:
        cols = {r[1] for r in c.execute("PRAGMA table_info(pages)")}
    assert {"page_type", "authors", "article_title"} <= cols


def test_migration_of_a_copy_of_the_real_lector_db(tmp_path):
    real = ROOT / "data" / "lector.db"
    if not real.exists():
        pytest.skip("no data/lector.db copy")
    src = sqlite3.connect(f"file:{real.as_posix()}?mode=ro", uri=True)
    dst = sqlite3.connect(tmp_path / "lector.db")
    src.backup(dst)
    src.close()
    before = schema.counts(dst)
    dst.close()
    if not upgrade.status(tmp_path / "lector.db")["needs_update"]:
        pytest.skip("data/lector.db already updated")
    report = upgrade.apply(tmp_path / "lector.db", tmp_path / "_backups")
    for k in ("searches", "pages", "page_hits", "pages with text", "reviewed hits"):
        assert report["after"][k] == before[k]
    if before["saved_searches"] is not None:
        # Already has saved searches (e.g. Wil's PC after the 12 Oct update, with later
        # runs): a later update only adds; every existing row is kept as it was.
        for k in ("saved_searches", "search_names", "runs", "run_pages", "search_pages"):
            assert report["after"][k] == before[k]
        return
    hits = sqlite3.connect(tmp_path / "lector.db").execute(
        "SELECT COUNT(DISTINCT accepted_name) FROM page_hits").fetchone()[0]
    assert report["after"]["saved_searches"] == hits
    assert report["after"]["search_pages"] == report["after"]["run_pages"]
    svc = LectorService(db_path=tmp_path / "lector.db", client_factory=FakeBHL)
    assert sum(s["pages_held"] for s in svc.list_searches()) == before["page_hits"]
    q = svc.review_queue(svc.list_searches()[0]["search_id"])
    assert q and sum(1 + len(i["copies"]) for i in q) == before["pages"]


def test_engine_rejects_unknown_mode_and_empty_names(tmp_path):
    st = LectorStore(tmp_path / "l.db")
    sid = saved.create_search(st.conn, "x", "Lamia textor", TEXTOR, names=[
        {"name": "Lamia textor", "origin": "accepted", "included": False}])
    with pytest.raises(ValueError):
        engine.run_search(st, sid, FakeBHL(), mode="fetch")
    with pytest.raises(ValueError):
        engine.run_search(st, sid, FakeBHL(), mode="all")
    st.close()
