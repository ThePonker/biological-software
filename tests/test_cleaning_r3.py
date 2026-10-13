"""Round 3 wave 2 (11 Oct), Wil's decisions 4-8 on the cleaning report and data clean-up:
D4 records on TVKs not in UKSI (report + one-off remap), D5 hybrid-by-rank listing,
D6 lat/long text in the grid ref, D7 dataset picker parents and sub-surveys, D8 kingdom /
taxon group backfill options.

Offscreen. The real data copies (data/*.db) are only read; anything that writes works on a
copy made with the backup API inside pytest's tmp_path. Tests on the real copies are skipped
when the copies are absent.
"""
import importlib.util
import os
import sqlite3

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import paths  # noqa: E402

HAVE_DATA = all(p.exists() for p in (paths.OBSERVATUM_DB, paths.UKSI_DB))
needs_data = pytest.mark.skipif(not HAVE_DATA, reason="real data copies not present")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _load(rel):
    p = os.path.join(ROOT, *rel.split("/"))
    spec = importlib.util.spec_from_file_location(os.path.basename(p)[:-3], p)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _copy_db(src, dest):
    s, d = sqlite3.connect(f"file:{src}?mode=ro", uri=True), sqlite3.connect(str(dest))
    try:
        s.backup(d)
    finally:
        d.close()
        s.close()
    return str(dest)


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


# ------------------------------------------------------------------ D4 TVKs not in UKSI

@needs_data
def test_report_tvks_not_in_uksi_real_copy():
    from scripts import report_common as rc
    from scripts import report_tvks_not_in_uksi as rep
    conn = rc.connect_ro(paths.OBSERVATUM_DB)
    try:
        found = rep.find(conn, paths.UKSI_DB)
    finally:
        conn.close()
    u = rc.connect_ro(paths.UKSI_DB)
    try:
        for g in found["groups"]:                        # every listed TVK really is missing
            assert u.execute("SELECT 1 FROM taxa WHERE tvk = ?", (g["tvk"],)).fetchone() is None
            m = g["match"]
            if m["match"] in rep.REMAPPABLE:             # and every target is a current taxon
                assert u.execute("SELECT scientific_name FROM taxa WHERE tvk = ?",
                                 (m["tvk"],)).fetchone()[0] == m["name"]
    finally:
        u.close()
    by_name = {g["name"]: g for g in found["groups"]}
    if "Forficula auricularia" in by_name:               # the case Wil named
        assert by_name["Forficula auricularia"]["match"]["match"] == "exact"
        assert by_name["Forficula auricularia"]["match"]["tvk"] == "NHMSYS0001387317"
    agg = by_name.get("Forficula auricularia/dentata agg.")
    if agg:                                              # no UKSI aggregate: never remapped
        assert agg["match"]["match"].startswith("no match")
    iso = by_name.get("Isoaeschna isoceles")
    if iso:
        assert iso["match"]["match"] == "synonym" and iso["match"]["name"] == "Aeshna isoceles"
    assert sum(g["records"] for g in found["groups"]) == sum(
        sum(c.values()) for c in found["by_tvk"].values())


@needs_data
def test_fix_tvks_not_in_uksi_apply_on_copy(tmp_path):
    fix = _load("scripts/_oneoff/fix_tvks_not_in_uksi_20261011.py")
    from Observatum.src.services.cleaning_report import build_report
    from scripts import report_common as rc
    from scripts import report_tvks_not_in_uksi as rep
    db = _copy_db(paths.OBSERVATUM_DB, tmp_path / "obs.db")
    before_no_tvk = len(build_report(db, paths.UKSI_DB, "observations").rows["no_tvk"])
    assert fix.main(["--db", db, "--uksi", str(paths.UKSI_DB)]) == 0      # dry run
    ro = rc.connect_ro(db)
    try:
        p = fix.plan(ro, paths.UKSI_DB)
        n_obs = ro.execute("SELECT COUNT(*) FROM observations").fetchone()[0]
    finally:
        ro.close()
    if not p["items"]:
        pytest.skip("nothing to remap on this copy")
    moved = {(i["table"], i["id"]): i for i in p["items"]}
    flagged_moved = sum(1 for i in p["items"] if i["match"]["match"] == "exact")
    assert fix.main(["--db", db, "--uksi", str(paths.UKSI_DB), "--apply",
                     "--backup-dir", str(tmp_path / "bk")]) == 0
    assert len(os.listdir(tmp_path / "bk")) == 1                         # backup first
    c = sqlite3.connect(db)
    try:
        assert c.execute("SELECT COUNT(*) FROM observations").fetchone()[0] == n_obs
        for (table, rid), i in moved.items():
            tvk, name, notes = c.execute(f"SELECT species_tvk, species_name, import_notes FROM "
                                         f"{table} WHERE id = ?", (rid,)).fetchone()
            assert tvk == i["match"]["tvk"]
            assert name == i["row"]["species_name"]                     # name untouched
            assert "decision 4" in notes
        found = rep.find(c, paths.UKSI_DB)
    finally:
        c.close()
    assert all(g["match"]["match"] not in rep.REMAPPABLE for g in found["groups"])
    after_no_tvk = len(build_report(db, paths.UKSI_DB, "observations").rows["no_tvk"])
    # the exact matches were on TVKs the cleaning report flags; Isoaeschna (synonym) was
    # already placed through UKSI's name_map, so it never counted there
    assert after_no_tvk <= before_no_tvk - flagged_moved


# ------------------------------------------------------------------ D5 hybrids by rank only

def test_hybrids_by_rank_only_temp(tmp_path):
    from scripts import report_tvks_not_in_uksi as rep
    from shared.record_kinds import classify
    main, uksi = tmp_path / "o.db", tmp_path / "u.db"
    c = sqlite3.connect(main)
    c.execute("CREATE TABLE observations (id INTEGER PRIMARY KEY, species_name TEXT, "
              "species_tvk TEXT, taxon_rank TEXT, source TEXT, date TEXT, site_name TEXT, "
              "irecord_id INTEGER, import_notes TEXT)")
    c.executemany("INSERT INTO observations VALUES (?,?,?,?,?,?,?,?,?)", [
        (1, "Rubus", "G1", "Species hybrid", "iRecord | x", "2014-10-22", "Hedge", 155384, None),
        (2, "Salix x multinervis", "H1", "Species hybrid", None, "2020-01-01", "Fen", None, None),
        (3, "Tilia cordata × platyphyllos", "H2", "Species hybrid", None, "", "", None, None),
        (4, "Rubus", "G1", "Genus", None, "", "", None, None),
        (5, "Zygiella x-notata", "S1", "Species", None, "", "", None, None),
    ])
    c.commit()
    c.close()
    u = sqlite3.connect(uksi)
    u.execute("CREATE TABLE taxa (tvk TEXT, scientific_name TEXT, rank TEXT, family TEXT)")
    u.execute("INSERT INTO taxa VALUES ('G1', 'Rubus', 'Genus', 'Rosaceae')")
    u.commit()
    u.close()
    # the rule itself: the stored rank wins, a hyphenated epithet is not a hybrid
    assert "hybrid" in classify("Rubus", "Species hybrid")
    assert "hybrid" not in classify("Rubus")
    assert classify("Zygiella x-notata", "Species") == frozenset()
    conn = sqlite3.connect(f"file:{main}?mode=ro", uri=True)
    try:
        items = rep.hybrids_by_rank_only(conn, str(uksi))
    finally:
        conn.close()
    assert [d["id"] for d in items] == [1]                 # real hybrid names are not listed
    d = items[0]
    assert (d["table"], d["uksi_name"], d["uksi_rank"], d["source"]) == (
        "observations", "Rubus", "Genus", "iRecord | x")
    row = next(rep.hybrid_rows(items))
    assert row[-1].endswith("occurrence_id=155384") and row[8] is None   # no dataset_name column


@needs_data
def test_hybrids_by_rank_only_real_copy():
    from scripts import report_common as rc
    from scripts import report_tvks_not_in_uksi as rep
    conn = rc.connect_ro(paths.OBSERVATUM_DB)
    try:
        items = rep.hybrids_by_rank_only(conn, paths.UKSI_DB)
    finally:
        conn.close()
    # Wil's four names, each on its genus TVK (not cut short: name == UKSI name of the TVK)
    assert {d["species_name"] for d in items} <= {"Armeria", "Dactylorhiza", "Helianthemum", "Rubus"}
    assert all(d["uksi_rank"] == "Genus" and d["uksi_name"] == d["species_name"] for d in items)


# ------------------------------------------------------------------ D6 lat/long in the grid ref

def _cleaning_dbs(tmp_path):
    """A scheme table (with source) and an observations table (without), plus a tiny UKSI."""
    main, uksi = tmp_path / "obs.db", tmp_path / "uksi.db"
    cols = ("id INTEGER PRIMARY KEY, irecord_id INTEGER, species_name TEXT, species_tvk TEXT, "
            "taxon_rank TEXT, date TEXT, date_type TEXT, grid_ref TEXT, site_name TEXT, "
            "recorder TEXT, latitude REAL, longitude REAL")
    rows = [
        (1, 11, "Clytus arietis", "T1", "Species", "2025-06-01", "D", "SP5820", "A", "W", None, None),
        (2, 12, "Clytus arietis", "T1", "Species", "2025-06-01", "D", "49.4557N2.5591W", "A", "W", None, None),
        (3, None, "Clytus arietis", "T1", "Species", "2025-06-01", "D", "~42.146N73.8892W", "A", "W", None, None),
        (4, 14, "Clytus arietis", "T1", "Species", "2025-06-01", "D", "SP58", "A", "W", None, None),
    ]
    c = sqlite3.connect(main)
    c.execute(f"CREATE TABLE recording_scheme ({cols}, source TEXT)")
    c.executemany("INSERT INTO recording_scheme VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  [r + (f"iRecord | survey {r[0]}",) for r in rows])
    c.execute(f"CREATE TABLE observations ({cols})")
    c.executemany("INSERT INTO observations VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    c.commit()
    c.close()
    u = sqlite3.connect(uksi)
    u.execute('CREATE TABLE taxa (tvk TEXT PRIMARY KEY, scientific_name TEXT, rank TEXT, '
              'parent_tvk TEXT, kingdom TEXT, "order" TEXT, family TEXT, genus TEXT)')
    u.execute("INSERT INTO taxa (tvk, scientific_name, rank) VALUES ('T1', 'Clytus arietis', 'Species')")
    u.execute("CREATE TABLE common_names (common_name TEXT, tvk TEXT, preferred INTEGER)")
    u.execute("CREATE TABLE synonyms (synonym TEXT, tvk TEXT)")
    u.execute("CREATE TABLE name_map (tvk TEXT PRIMARY KEY, recommended_tvk TEXT NOT NULL, name TEXT, "
              "recommended_name TEXT, name_status TEXT, deprecated TEXT, source TEXT)")
    u.execute("CREATE TABLE tvk_remap (old_tvk TEXT PRIMARY KEY, new_tvk TEXT, method TEXT, "
              "old_name TEXT, new_name TEXT)")
    u.commit()
    u.close()
    return str(main), str(uksi)


def test_latlon_problem_rule():
    from Observatum.src.services.cleaning_report import grid_problem, latlon_problem
    assert latlon_problem("49.4557N2.5591W") == "Latitude/longitude text, not a grid ref"
    assert latlon_problem("~49.4247N2.6394W") and latlon_problem("50.5385N2.541E")
    for ref in ("SP5820", "SP58", "SP58A", "", None, "SP583", "H1234"):
        assert latlon_problem(ref) is None
    assert grid_problem("49.4557N2.5591W") is None            # not also "coarse"/"unreadable"


def test_cleaning_report_latlon_list_export_and_screen(tmp_path, qapp):
    from Observatum.src.services.cleaning_export import HEADERS, SHEET_NAMES, export_xlsx
    from Observatum.src.services.cleaning_report import PROBLEM_LABELS, PROBLEMS, build_report
    main, uksi = _cleaning_dbs(tmp_path)
    assert PROBLEM_LABELS["latlon"] == "Not a grid ref (looks like latitude/longitude)"
    rep = build_report(main, uksi, "recording_scheme")
    assert [r.record_id for r in rep.rows["latlon"]] == [2, 3]
    assert [r.dataset for r in rep.rows["latlon"]] == ["iRecord | survey 2", "iRecord | survey 3"]
    assert [r.grid_ref for r in rep.rows["latlon"]] == ["49.4557N2.5591W", "~42.146N73.8892W"]
    assert [r.record_id for r in rep.rows["grid"]] == [4]
    assert rep.records_with_problem == 3
    obs = build_report(main, uksi, "observations")             # no source column: blank
    assert [r.dataset for r in obs.rows["latlon"]] == ["", ""]

    from openpyxl import load_workbook
    path = export_xlsx(rep, str(tmp_path / "c.xlsx"))
    wb = load_workbook(path, read_only=True)
    try:
        # the scheme's export has no "Came back from iRecord" sheet (observations only, F72)
        assert wb.sheetnames == ["Summary"] + [SHEET_NAMES[k] for k, _ in PROBLEMS if k != "returned"]
        sheet = list(wb[SHEET_NAMES["latlon"]].values)
        assert list(sheet[0]) == HEADERS and HEADERS[-1] == "Source"
        assert [(s[3], s[9], s[10]) for s in sheet[1:]] == [
            ("49.4557N2.5591W", 2, "iRecord | survey 2"), ("~42.146N73.8892W", 3, "iRecord | survey 3")]
        summary = {r[0]: r[1] for r in wb["Summary"].values if r and r[0]}
        assert summary["Not a grid ref (looks like latitude/longitude)"] == 2
    finally:
        wb.close()
    assert all(len(n) <= 31 for n in SHEET_NAMES.values())

    from PySide6.QtCore import Qt
    from src.views.stats.cleaning_dashboard import CleaningDashboard
    d = CleaningDashboard(main_db=main, uksi_db=uksi)
    d.cmb_source.setCurrentIndex(1)                             # Recording Scheme
    d.refresh()
    assert d._tiles["latlon"].text() == "Not a grid ref (looks like latitude/longitude)\n2"
    d._tiles["latlon"].click()
    assert d.current_key() == "latlon" and d._tiles["latlon"].isChecked()
    model = d.current_view().model()
    assert model.headerData(8, Qt.Orientation.Horizontal) == "Source"
    assert model.index(0, 8).data() == "iRecord | survey 2"
    d.deleteLater()


@needs_data
def test_latlon_real_copy():
    from Observatum.src.services.cleaning_report import build_report
    rep = build_report(paths.OBSERVATUM_DB, paths.UKSI_DB, "recording_scheme")
    rows = rep.rows["latlon"]
    assert rows and all(r.dataset for r in rows)
    assert all(r.grid_ref[:1] in "~0123456789" for r in rows)


# ------------------------------------------------------------------ D7 dataset parents / sub-surveys

DATASETS = [("NBN Atlas", 5), ("iRecord | iRecord App", 40), ("iRecord | iRecord App | Bayfordbury", 1),
            ("iRecord | iRecord App | Green Redgrave", 2), ("iRecord | iRecord App General Survey", 7),
            ("iRecord | iRecord App General Survey | Sands Meadow", 3),
            ("Suffolk BRC | Dunwich | Dunwich", 1)]


def test_dataset_tree_order_and_parents():
    from src.views.components.filter_wizard.dataset_picker import dataset_tree
    order, parent = dataset_tree(n for n, _ in DATASETS)
    assert [n for n, _ in order] == [
        "iRecord | iRecord App", "iRecord | iRecord App | Bayfordbury",
        "iRecord | iRecord App | Green Redgrave", "iRecord | iRecord App General Survey",
        "iRecord | iRecord App General Survey | Sands Meadow", "NBN Atlas",
        "Suffolk BRC | Dunwich | Dunwich"]
    assert parent["iRecord | iRecord App | Bayfordbury"] == "iRecord | iRecord App"
    assert parent["iRecord | iRecord App General Survey"] is None        # not " | "
    assert parent["Suffolk BRC | Dunwich | Dunwich"] is None             # its parent isn't listed
    assert dict(order)["iRecord | iRecord App | Green Redgrave"] == 1


def test_dataset_picker_parent_ticks_children(qapp):
    from PySide6.QtCore import Qt
    from src.views.components.filter_wizard.dataset_picker import DatasetPicker
    p = DatasetPicker(DATASETS, "#4a7c59")
    C, U, P = Qt.CheckState.Checked, Qt.CheckState.Unchecked, Qt.CheckState.PartiallyChecked
    app_ = p._item["iRecord | iRecord App"]
    bay = p._item["iRecord | iRecord App | Bayfordbury"]
    red = p._item["iRecord | iRecord App | Green Redgrave"]
    assert bay.text().strip() == "Bayfordbury (1)" and app_.text() == "iRecord | iRecord App (40)"

    app_.setCheckState(C)                                   # tick the parent: children follow
    assert p.selected() == ["iRecord | iRecord App", "iRecord | iRecord App | Bayfordbury",
                            "iRecord | iRecord App | Green Redgrave"]
    bay.setCheckState(U)                                    # untick one child: parent part-filled
    assert app_.checkState() == P
    assert p.selected() == ["iRecord | iRecord App", "iRecord | iRecord App | Green Redgrave"]
    bay.setCheckState(C)                                    # all ticked again: parent full
    assert app_.checkState() == C
    red.setCheckState(U)
    app_.setCheckState(C)                                   # clicking a part-filled parent re-ticks all
    assert red.checkState() == C
    app_.setCheckState(U)                                   # unticking it unticks all
    assert p.selected() == []
    red.setCheckState(C)                                    # a child alone; parent stays unticked
    assert app_.checkState() == U and p.selected() == ["iRecord | iRecord App | Green Redgrave"]

    # a saved filter: parent + one child loads part-filled, and gives back the same list
    p.set_selected(["irecord | irecord app", "iRecord | iRecord App | Bayfordbury"])
    assert app_.checkState() == P and bay.checkState() == C and red.checkState() == U
    assert p.selected() == ["iRecord | iRecord App", "iRecord | iRecord App | Bayfordbury"]
    p.set_selected(["iRecord | iRecord App"])               # a Round 2 filter: parent alone
    assert app_.checkState() == P and p.selected() == ["iRecord | iRecord App"]

    p.clear()
    p.search.setText("redgrave")                            # a match keeps its parent visible
    shown = [n for n, it in p._item.items() if not it.isHidden()]
    assert shown == ["iRecord | iRecord App", "iRecord | iRecord App | Green Redgrave"]
    p.deleteLater()


@needs_data
def test_dataset_filter_counts_real_copy(qapp):
    from src.services.filter_builder import matching_ids
    from src.views.components.filter_wizard.filter_wizard import FilterWizard
    from src.views.components.filter_wizard.status_filter_dialog import StatusFilterDialog
    w = FilterWizard(tab_name="recording_scheme")
    w.set_db_path(paths.OBSERVATUM_DB)
    w._ensure_tab_values()
    counts = dict(w._dataset_list)
    d = StatusFilterDialog(status_list=w._status_list, type_list=w._record_type_list,
                           dataset_list=w._dataset_list)
    pk = d.dataset_picker
    parent = "iRecord | iRecord Longhorn Trial"
    if parent not in pk._item:
        pytest.skip("no Longhorn Trial dataset on this copy")
    from PySide6.QtCore import Qt
    pk._item[parent].setCheckState(Qt.CheckState.Checked)
    chosen = d.get_values()["dataset"]
    kids = pk._descendants(parent)
    assert chosen == sorted([parent] + kids, key=lambda v: (v.casefold(), v))
    ids = matching_ids(paths.OBSERVATUM_DB, "recording_scheme", {"dataset": chosen})
    assert len(ids) == sum(counts[n] for n in chosen)
    for k in kids:
        pk._item[k].setCheckState(Qt.CheckState.Unchecked)
    only = d.get_values()["dataset"]
    assert only == [parent]
    assert len(matching_ids(paths.OBSERVATUM_DB, "recording_scheme", {"dataset": only})) == counts[parent]


# ------------------------------------------------------------------ D8 kingdom / taxon group backfill

def _backfill_dbs(tmp_path):
    u = sqlite3.connect(tmp_path / "uksi.db")
    u.execute('CREATE TABLE taxa (tvk TEXT PRIMARY KEY, scientific_name TEXT, rank TEXT, '
              'parent_tvk TEXT, kingdom TEXT, phylum TEXT, class TEXT, "order" TEXT, family TEXT, '
              'sort_code INTEGER, superfamily TEXT)')
    u.executemany("INSERT INTO taxa VALUES (?,?,?,?,?,?,?,?,?,?,?)", [
        ("B1", "Carabus nemoralis", "Species", None, "Animalia", "Arthropoda", "Insecta",
         "Coleoptera", "Carabidae", 18000, "Caraboidea"),
        ("F1", "Lepista nuda", "Species", None, "Fungi", "Basidiomycota", "Agaricomycetes",
         "Agaricales", "", 90000, None),
        ("X1", "Oddity", "Species", None, "", "", "", "", "", 5, None),
    ])
    u.commit()
    u.close()
    o = sqlite3.connect(tmp_path / "observatum.db")
    o.execute("CREATE TABLE observations (id INTEGER PRIMARY KEY, species_tvk TEXT, source TEXT, "
              "order_name TEXT, taxon_group TEXT, kingdom TEXT, taxon_rank TEXT, superfamily TEXT, "
              "taxonomic_sort_key, species_name TEXT)")
    o.execute("CREATE TABLE recording_scheme (id INTEGER PRIMARY KEY, species_tvk TEXT, source TEXT, "
              "order_name TEXT, taxon_group TEXT, kingdom TEXT)")
    o.executemany("INSERT INTO observations (id, species_tvk, taxon_group, kingdom, species_name) "
                  "VALUES (?,?,?,?,?)", [
                      (1, "B1", None, None, "Carabus nemoralis"),       # both filled
                      (2, "F1", "", "Fungi", "Lepista nuda"),           # group only
                      (3, "ZZ", None, None, "Unknown"),                 # TVK not in UKSI
                      (4, None, None, None, "No key"),                  # no TVK
                      (5, "X1", None, None, "Oddity"),                  # UKSI has nothing
                      (6, "B1", "insect - beetle (Coleoptera)", "Animalia", "Carabus nemoralis"),
                  ])
    o.execute("INSERT INTO recording_scheme VALUES (1, 'B1', 'x', 'Coleoptera', NULL, 'Animalia')")
    o.commit()
    o.close()
    return str(tmp_path / "observatum.db"), str(tmp_path / "uksi.db")


def test_backfill_kingdom_group_only_observations(tmp_path, capsys):
    m = _load("scripts/backfill_taxon_groups.py")
    db, uksi = _backfill_dbs(tmp_path)
    args = ["--db", db, "--uksi", uksi, "--tables", "observations", "--columns", "kingdom,taxon_group"]
    assert m.main(args) == 0                                                 # dry run
    out = capsys.readouterr().out
    assert "Nothing has been changed" in out
    assert "records with no kingdom or no taxon group: 5" in out
    assert "1  get  Animalia / insect - beetle (Coleoptera)" in out
    assert "1  get  Fungi / fungus" in out
    assert "cannot be resolved: TVK not in UKSI" in out and "cannot be resolved: no TVK" in out
    assert "1  cannot be resolved: UKSI gives no kingdom or taxon group for the TVK" in out
    assert "recording_scheme" not in out.split("only columns")[1]
    ro = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    found = m.plan(ro, uksi, ["observations"], ["kingdom", "taxon_group"])
    ro.close()
    assert set(found) == {"observations"}
    assert set(found["observations"]["fill"]) == {"taxon_group", "kingdom"}

    assert m.main(args + ["--apply", "--backup-dir", str(tmp_path / "bk")]) == 0
    assert len(os.listdir(tmp_path / "bk")) == 1
    c = sqlite3.connect(db)
    got = {r[0]: r[1:] for r in c.execute(
        "SELECT id, kingdom, taxon_group, taxon_rank, superfamily, taxonomic_sort_key FROM observations")}
    assert got[1] == ("Animalia", "insect - beetle (Coleoptera)", None, None, None)  # only the two
    assert got[2] == ("Fungi", "fungus", None, None, None)
    assert got[3][:2] == (None, None) and got[4][:2] == (None, None)
    assert got[6][:2] == ("Animalia", "insect - beetle (Coleoptera)")
    assert c.execute("SELECT taxon_group FROM recording_scheme").fetchone()[0] is None  # not chosen
    c.close()
    with pytest.raises(SystemExit):
        m.main(["--db", db, "--uksi", uksi, "--columns", "kingdom", "--relabel", "--apply"])
    with pytest.raises(SystemExit):
        m.main(["--db", db, "--uksi", uksi, "--tables", "nonsense"])
