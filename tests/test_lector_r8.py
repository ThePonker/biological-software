"""Lector Round 8 (12 Oct 2026): fixes from Wil's tests, L1-L16 (offscreen, temp DBs).

No network (test_lector_a's FakeBHL); no assertion depends on font metrics except the
1280x720 row check, which allows for wide text by scrolling the page, not the row.
"""
import csv
import json
import os
import sqlite3
import threading
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import QPushButton, QToolButton  # noqa: E402
from test_lector_a import BRITISH, TEXTOR, FakeBHL, _page, make_uksi  # noqa: E402
from test_lector_ui import SlowBHL, make_env, make_window, pump, run_and_wait  # noqa: E402

from Lector import engine, schema, upgrade  # noqa: E402
from Lector.service import LectorService  # noqa: E402
from Lector.store import LectorStore  # noqa: E402

ANNALS_ONLY = {"include_words": ["Annals"], "max_pages": 0}


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.fixture
def env(tmp_path):
    return make_env(tmp_path)


@pytest.fixture
def window(qapp, env):
    win = make_window(qapp, env[0])
    yield win
    win.tasks.cancel()
    win.tasks.wait(5000)
    win.close()
    win.deleteLater()
    pump(qapp)


def buttons(widget):
    return [b.text() for b in widget.findChildren(QPushButton)]


# --- L1 navigation ---------------------------------------------------------------------
def test_back_buttons_go_where_you_came_from(qapp, window, env):
    svc, _fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    for key in ("builder", "history", "review", "export"):
        assert not window.screens[key].back_btn.isHidden(), key
        assert window.screens[key].back_btn.text() == "← Back"
    assert window.screens["searches"].back_btn.isHidden()
    window.go("searches")
    window.go("history", search_id=sid)
    window.go("builder", search_id=sid)                         # Edit search
    window.screens["builder"].back_btn.click()
    assert window.current() is window.screens["history"]
    window.go("review", search_id=sid)
    window.screens["review"].back_btn.click()
    assert window.current() is window.screens["history"]
    window.screens["history"].back_btn.click()
    assert window.current() is window.screens["searches"]
    window.screens["export"].back_btn.click()                   # nothing before: home
    assert window.current() is window.screens["searches"]
    # the editor is skipped on the way back after "Save and run"
    window.go("history", search_id=sid)
    b = window.go("builder", search_id=sid)
    assert b.save(run=True) == sid
    assert pump(qapp, lambda: not window.tasks.busy)
    runs = len(svc.list_runs(sid))
    window.screens["history"].back_btn.click()                  # not the editor, not twice
    assert window.current() is window.screens["searches"]
    # Back to a search's page never runs it again
    window.go("history", search_id=sid, run="fetch")
    assert pump(qapp, lambda: not window.tasks.busy)
    window.go("review", search_id=sid)
    window.screens["review"].back_btn.click()
    assert window.current() is window.screens["history"] and not window.tasks.busy
    assert len(svc.list_runs(sid)) == runs + 1


def test_left_menu_always_navigates(qapp, window, env):
    from PySide6.QtTest import QTest
    svc, _fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    for key in ("history", "builder"):
        window.go(key, search_id=sid)
        assert window.nav_buttons["searches"].isChecked()       # highlighted already
        QTest.mouseClick(window.nav_buttons["searches"], Qt.MouseButton.LeftButton)
        pump(qapp)
        assert window.current() is window.screens["searches"], key
    window.go("review", search_id=sid)
    QTest.mouseClick(window.nav_buttons["review"], Qt.MouseButton.LeftButton)
    assert window.current() is window.screens["review"]
    QTest.mouseClick(window.nav_buttons["export"], Qt.MouseButton.LeftButton)
    assert window.current() is window.screens["export"]


# --- L2 / L4 dashboard rows and filter box ---------------------------------------------------
def test_every_row_has_edit_and_the_filter_box(qapp, window, env):
    svc, _fake = env
    ran = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(ran)
    old = svc.create_search("Lamia textor", name="Old fetch", filters=BRITISH)
    svc.run_search(old, mode="count")
    svc.update_search(old, filters={"recorded": False})
    draft = svc.create_search("Monochamus sutor")
    dash = window.go("searches")
    rows = {s["search_id"]: r for r, s in enumerate(dash.rows)}
    assert {"Run again", "Edit", "Review"} <= set(buttons(dash.table.cellWidget(rows[ran], 5)))
    assert {"Set filters", "Edit"} <= set(buttons(dash.table.cellWidget(rows[old], 5)))
    assert {"Edit", "Run"} <= set(buttons(dash.table.cellWidget(rows[draft], 5)))
    edit = next(b for b in dash.table.cellWidget(rows[ran], 5).findChildren(QPushButton)
                if b.text() == "Edit")
    edit.click()
    assert window.current() is window.screens["builder"]
    assert window.screens["builder"].search_id == ran
    dash = window.go("searches")
    box = dash.filter_box
    assert box.placeholderText() == "Filter your searches…" and box.isClearButtonEnabled()
    box.setText("sutor")                                        # live, no Enter
    shown = [s["search_id"] for r, s in enumerate(dash.rows) if not dash.table.isRowHidden(r)]
    assert shown == [draft]
    box.clear()
    assert not any(dash.table.isRowHidden(r) for r in range(dash.table.rowCount()))


def test_dashboard_rows_fit_1280x720(qapp, tmp_path):
    svc, _fake = make_env(tmp_path)
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    win = make_window(qapp, svc)
    try:
        win.resize(1280, 720)
        dash = win.go("searches")
        pump(qapp, timeout=0.3)
        t = dash.table
        actions = t.cellWidget(0, 5)
        # every button of the row is inside its column (none cut off)
        for b in actions.findChildren(QPushButton) + actions.findChildren(QToolButton):
            assert b.geometry().right() <= t.columnWidth(5), b.text()
        assert t.columnWidth(0) >= 200
        assert t.horizontalScrollBar().maximum() == 0               # the row fits: no scrolling
    finally:
        win.close()


# --- L3 fetch missing text -------------------------------------------------------------------
def _with_missing_text(svc, fake, recorded=True):
    fake.fail_pages = {101, 102}
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    fake.fail_pages = set()
    if not recorded:
        svc.update_search(sid, filters={"recorded": False})
    return sid


def test_fetch_missing_text_downloads_only_text(env):
    svc, fake = env
    sid = _with_missing_text(svc, fake, recorded=False)        # works without filters
    st = LectorStore(svc.db_path)
    before = dict(st.conn.execute("SELECT page_id, state FROM search_pages WHERE search_id=?",
                                  (sid,)).fetchall())
    runs = st.conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0]
    st.close()
    s = next(x for x in svc.list_searches() if x["search_id"] == sid)
    assert s["pages_without_text"] == 2
    fake.asked.clear()
    out = svc.fetch_missing_text(sid)
    assert out["status"] == "done" and out["text_fetched"] == 2 and out["still_missing"] == 0
    assert sorted(fake.asked) == [("page", 101), ("page", 102)]   # no name searches
    st = LectorStore(svc.db_path)
    after = dict(st.conn.execute("SELECT page_id, state FROM search_pages WHERE search_id=?",
                                 (sid,)).fetchall())
    assert after == before                                       # nothing withdrawn or added
    assert st.conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == runs
    st.close()
    dash = svc.dashboard()
    assert any(a["kind"] == "fetch_text" and "2 pages" in a["text"] for a in dash["activity"])
    assert not any(i["action"] == "fetch_text" for i in dash["attention_items"])
    again = svc.fetch_missing_text(sid)
    assert again["text_fetched"] == 0 and "already has its text" in again["message"]


def test_fetch_missing_text_cancel(env):
    svc, fake = env
    sid = _with_missing_text(svc, fake)
    stop = threading.Event()

    def progress(stage, done, total, msg):
        if stage == "text" and done == 1:
            stop.set()
    out = svc.fetch_missing_text(sid, progress=progress, cancel=stop)
    assert out["status"] == "cancelled" and out["text_fetched"] == 1
    assert out["still_missing"] == 1 and "Cancelled" in out["message"]


def test_fetch_missing_text_in_the_window(qapp, window, env):
    svc, fake = env
    sid = _with_missing_text(svc, fake)
    dash = window.go("searches")
    more = dash.table.cellWidget(0, 5).findChildren(QToolButton)[0]
    acts = [a.text() for a in more.menu().actions()]
    assert "Fetch missing text (2 pages)" in acts
    fetch_now = [b for b in dash.attention.findChildren(QPushButton) if b.text() == "Fetch now"]
    assert len(fetch_now) == 1
    run_and_wait(qapp, window, fetch_now[0].click)
    assert next(s for s in svc.list_searches() if s["search_id"] == sid)["pages_without_text"] == 0
    dash = window.screens["searches"]
    more = dash.table.cellWidget(0, 5).findChildren(QToolButton)[0]
    assert not any(a.text().startswith("Fetch missing") for a in more.menu().actions())
    assert not [b for b in dash.attention.findChildren(QPushButton) if b.text() == "Fetch now"]


# --- L5 held-back names in migrated searches -----------------------------------------------
def _legacy_with_synonyms(path):
    """A 7 Oct lector.db: fetched with --synonyms, so held-back names were searched."""
    c = sqlite3.connect(path)
    c.executescript(schema.LEGACY)
    names = [("Lamia textor", "accepted"), ("Cerambyx textor", "synonym"),
             ("Lamia sartor", "synonym"), ("Lamia sutor", "synonym")]
    for q, role in names:
        c.execute("INSERT INTO searches VALUES ('Lamia textor',?,?,?,'2026-10-07T22:20:21',5,2)",
                  (q, TEXTOR, role))
    for pid, q in ((1, "Lamia textor"), (2, "Lamia sartor"), (3, "Cerambyx textor")):
        c.execute("INSERT INTO pages (page_id, title, year, page_url, ocr_text, fetched_at)"
                  " VALUES (?,?,?,?,?,?)", (pid, "Annals", 1841, f"u{pid}", q, "t"))
        c.execute("INSERT INTO page_hits (page_id, accepted_name, query_name, tvk)"
                  " VALUES (?,?,?,?)", (pid, "Lamia textor", q, TEXTOR))
    c.commit()
    c.close()


def _uksi_env(tmp_path):
    uksi = tmp_path / "uksi.db"
    make_uksi(uksi)
    return uksi


def test_update_checks_old_names_instead_of_ticking_them(tmp_path):
    uksi = _uksi_env(tmp_path)
    _legacy_with_synonyms(tmp_path / "lector.db")
    svc = LectorService(db_path=tmp_path / "lector.db", uksi_db=uksi, client_factory=FakeBHL,
                        synonym_file=tmp_path / "syn.txt", backup_dir=tmp_path / "bk")
    report = svc.needs_update()
    held = dict(report["searches_made"][0]["held_back"])
    # the two the old fetch searched, and titillator (never searched, held back) listed too
    assert set(held) == {"Lamia sartor", "Lamia sutor", "Lamia titillator"}
    assert "Monochamus sutor" in held["Lamia sutor"]
    assert "Lamia sutor: not ticked" in upgrade.format_report(report)
    svc.update_database()
    s = svc.get_search(svc.list_searches()[0]["search_id"])
    names = {n["name"]: n for n in s["names"]}
    assert names["Lamia textor"]["included"] and names["Cerambyx textor"]["included"]
    assert not names["Lamia sartor"]["included"] and not names["Lamia sutor"]["included"]
    assert "Monochamus sartor" in names["Lamia sartor"]["held_back_reason"]
    assert svc.list_searches()[0]["pages_held"] == 3            # every page is still held
    # no UKSI at hand: the names are stored as fetched (and nothing is lost)
    _legacy_with_synonyms(tmp_path / "other.db")
    report = upgrade.dry_run(tmp_path / "other.db", upgrade.name_checker(tmp_path / "none.db"))
    assert report["searches_made"][0]["held_back"] == []


def _migrated_the_old_way(tmp_path):
    """Wil's lector.db: updated on 12 Oct before this fix (every old name ticked)."""
    _legacy_with_synonyms(tmp_path / "lector.db")
    conn = sqlite3.connect(tmp_path / "lector.db")
    schema.migrate(conn, check=None)
    conn.close()
    return tmp_path / "lector.db"


def test_fix_script_unticks_what_the_update_ticked(tmp_path, capsys):
    from scripts import fix_lector_name_choices as fix
    uksi = _uksi_env(tmp_path)
    db = _migrated_the_old_way(tmp_path)
    syn = tmp_path / "syn.txt"
    kw = dict(db=db, uksi_db=uksi, synonym_file=syn, backup_dir=tmp_path / "bk")

    def ticked():
        with sqlite3.connect(db) as c:
            return [r[0] for r in c.execute("SELECT name FROM search_names WHERE included=1"
                                            " ORDER BY position")]
    assert ticked() == ["Lamia textor", "Cerambyx textor", "Lamia sartor", "Lamia sutor"]
    assert fix.main([], **kw) == 0
    out = capsys.readouterr().out
    assert "untick 2 names" in out and "Lamia sartor  -  Held back: epithet is current" in out
    assert "found only by these names: 1" in out
    assert "DRY RUN - Nothing has been changed" in out
    assert len(ticked()) == 4 and not (tmp_path / "bk").exists()
    assert fix.main(["--apply"], **kw) == 0
    out = capsys.readouterr().out
    assert "Backup:" in out and len(list((tmp_path / "bk").glob("lector_pre_fix_names_*.db"))) == 1
    assert ticked() == ["Lamia textor", "Cerambyx textor"]
    with sqlite3.connect(db) as c:
        reason = c.execute("SELECT held_back_reason FROM search_names WHERE name='Lamia sutor'"
                           ).fetchone()[0]
    assert "Monochamus sutor" in reason
    assert fix.main([], **kw) == 0
    assert "No ticked names to change" in capsys.readouterr().out


def test_fix_script_leaves_your_own_searches_alone(tmp_path, capsys):
    from scripts import fix_lector_name_choices as fix
    uksi = _uksi_env(tmp_path)
    svc = LectorService(db_path=tmp_path / "lector.db", uksi_db=uksi, client_factory=FakeBHL,
                        synonym_file=tmp_path / "syn.txt")
    sid = svc.create_search("Lamia textor")
    names = svc.get_search(sid)["names"]
    for n in names:
        n["included"] = True                                    # you ticked them yourself
    svc.update_search(sid, names=names)
    fix.main([], db=tmp_path / "lector.db", uksi_db=uksi, synonym_file=tmp_path / "syn.txt")
    assert "No ticked names to change" in capsys.readouterr().out


def test_edit_lists_held_back_names_of_a_migrated_search(qapp, tmp_path):
    uksi = _uksi_env(tmp_path)
    db = _migrated_the_old_way(tmp_path)
    svc = LectorService(db_path=db, uksi_db=uksi, client_factory=FakeBHL,
                        synonym_file=tmp_path / "syn.txt")
    sid = svc.list_searches()[0]["search_id"]
    stored = {n["name"] for n in svc.get_search(sid)["names"]}
    shown = {n["name"]: n for n in svc.edit_names(sid)}
    assert "Lamia titillator" not in stored and "Lamia titillator" in shown
    assert not shown["Lamia titillator"]["included"]
    assert shown["Lamia titillator"]["held_back_reason"]
    assert "Monochamus sutor" in shown["Lamia sutor"]["held_back_reason"]
    win = make_window(qapp, svc)
    try:
        b = win.go("builder", search_id=sid)
        assert "Lamia titillator" in {n["name"] for n in b.names.names()}
        assert b.names._checks["Lamia titillator"][0].isChecked() is False
        # before the fix script: sutor is ticked, and it says who ticked it
        from PySide6.QtWidgets import QLabel
        texts = [lab.text() for lab in b.names.box.findChildren(QLabel)]
        assert any("ticked when lector.db was updated, not by you" in t for t in texts)
        assert not any(t.endswith("ticked by you") for t in texts)
    finally:
        win.close()


# --- L6 tick boxes ---------------------------------------------------------------------------
def test_ticked_boxes_have_a_white_tick(qapp):
    from PySide6.QtGui import QColor, QImage

    from Lector.ui import style as S
    qss = S.app_qss()
    assert "QCheckBox::indicator:checked" in qss and "image: url(" in qss
    path = qss.split('image: url("')[1].split('")')[0]
    img = QImage(path)
    assert not img.isNull() and img.width() == 16 and Path(path.replace(".png", "@2x.png")).exists()
    mark = QColor(S.c("button_primary_text"))
    fill = QColor(S.ACCENT)
    centre = img.pixelColor(7, 9)            # on the tick's stroke near its foot
    corner = img.pixelColor(13, 13)          # the box, away from the tick
    assert abs(centre.lightness() - mark.lightness()) < abs(centre.lightness() - fill.lightness())
    assert corner.name() == fill.name()
    # one tick for the suite: Observatum's tick column draws it with the same function
    from src.views.components import tick_column
    assert tick_column.paint_tick.__module__ == "shared.tick_mark"


# --- L7 / L13 / L14 reading list ---------------------------------------------------------------
def test_notes_box_is_outlined_and_labelled(qapp, window, env):
    svc, _fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    r = window.go("review", search_id=sid)
    notes = r.panel.notes
    assert "border: 1px solid" in notes.styleSheet() and notes.objectName() == "reviewNotes"
    from PySide6.QtWidgets import QLabel
    assert any(lab.text() == "Notes" and lab.buddy() is notes
               for lab in r.panel.findChildren(QLabel))


def test_reading_list_header_counts_entries_and_pages(qapp, window, env):
    from Lector.ui.review_screen import entries_text
    assert entries_text([{"copies": []}] * 3) == "3 entries"
    assert entries_text([{"copies": [1, 2]}, {"copies": []}]) == \
        "2 entries (4 pages, duplicates merged)"
    svc, _fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    r = window.go("review", search_id=sid, status="all")
    entries = len([i for i in r.items if i["state"] != "withdrawn"])
    pages = sum(1 + len(i["copies"]) for i in r.items if i["state"] != "withdrawn")
    assert pages > entries
    assert f"{entries} entries ({pages} pages, duplicates merged)" in r.subtitle.text()


def test_keys_1_to_4_mark_and_move_on_but_not_while_typing(qapp, window, env):
    svc, _fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    r = window.go("review", search_id=sid, status="all")
    keys = {sc.key().toString(): sc for sc in r.shortcuts}
    first = r.current_item()["page_id"]
    keys["1"].activated.emit()
    second = r.current_item()["page_id"]
    assert second != first
    assert {i["page_id"]: i for i in svc.review_queue(sid)}[first]["review_status"] == "relevant"
    window.activateWindow()
    r.panel.notes.setFocus()
    pump(qapp)
    from PySide6.QtWidgets import QApplication
    if QApplication.focusWidget() is r.panel.notes:             # offscreen focus is reliable
        keys["2"].activated.emit()
        assert r.current_item()["page_id"] == second
        assert {i["page_id"]: i for i in svc.review_queue(sid)}[second]["review_status"] == \
            "unreviewed"
    from PySide6.QtTest import QTest
    r.panel.notes.setFocus()
    QTest.keyClicks(r.panel.notes, "1834")
    assert r.current_item()["page_id"] == second and "1834" in r.panel.notes.toPlainText()


# --- L8 / L9 / L10 runs: numbers, cancel message, requests -------------------------------------
def test_runs_are_numbered_within_their_search(qapp, window, env):
    svc, fake = env
    a = svc.create_search("Lamia textor", filters=BRITISH)
    b = svc.create_search("Lamia textor", name="Second", filters=ANNALS_ONLY)
    svc.run_search(a)                     # run_id 1
    svc.run_search(b)                     # 2
    svc.run_search(a, mode="count")       # 3 (count only)
    r = svc.run_search(a, refresh=False)  # 4 -> Run 2 of a
    assert (r["run_id"], r["number"], r["label"]) == (4, 2, "Run 2")
    runs = svc.list_runs(a)
    assert [(x["label"], x["number"]) for x in runs] == [("Run 2", 2), ("Count", None),
                                                         ("Run 1", 1)]
    assert svc.list_runs(b)[0]["label"] == "Run 1"
    d = svc.run_diff(a)
    assert (d["run_b"], d["number_b"], d["number_a"]) == (4, 2, 1)
    h = window.go("history", search_id=a)
    assert h.run_list.item(0).text().startswith("Run 2 ·")
    assert h.diff_title.text() == "What changed in run 2"
    assert h.diff_against.text() == "Compared with run 1"
    assert any(x["text"].startswith("Run 2 of “Lamia textor”") for x in svc.dashboard()["activity"])
    run_and_wait(qapp, window, lambda: h.run("fetch"))
    assert h.status_line.text().startswith("Run 3 finished")
    if pytest.importorskip("docx"):
        out = svc.export(a, formats=("docx",))
        from docx import Document
        text = "\n".join(p.text for p in Document(out["files"]["docx"]).paragraphs)
        assert "Run: 3, " in text


def test_requests_count_names_and_page_texts(env):
    svc, fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    r = svc.run_search(sid)
    assert (r["requests"], r["name_requests"], r["text_requests"]) == (9, 2, 7)
    assert r["requests_text"] == "9 BHL requests (2 name searches, 7 page texts)"
    run = svc.list_runs(sid)[0]
    assert (run["name_requests"], run["text_requests"]) == (2, 7)
    assert "9 BHL requests (2 name searches, 7 page texts)" in svc.dashboard()["activity"][0]["text"]
    assert svc.list_searches()[0]["last_run"]["requests_text"] == run["requests_text"]
    r2 = svc.run_search(sid, refresh=False)
    assert r2["requests_text"] == "0 BHL requests (0 name searches, 0 page texts)"


def test_cancel_message_says_it_once(qapp, tmp_path):
    svc, fake = make_env(tmp_path, SlowBHL())
    win = make_window(qapp, svc)
    try:
        sid = svc.create_search("Lamia textor", filters=BRITISH)
        h = win.go("history", search_id=sid, run="fetch")
        assert pump(qapp, lambda: fake.request_count >= 3)
        win.cancel_task()
        assert pump(qapp, lambda: not win.tasks.busy)
        text = h.status_line.text()
        assert text == ("Run 1 cancelled. Whatever was downloaded is kept; running again "
                        "carries on from there.")
        assert svc.list_runs(sid)[0]["message"] == engine.CANCELLED_MESSAGE
    finally:
        win.close()
        pump(qapp, lambda: not win.tasks.busy)


# --- L11 a cancelled run and withdrawals --------------------------------------------------------
def test_cancel_during_name_searches_withdraws_nothing(env):
    svc, fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    held = {i["page_id"] for i in svc.review_queue(sid)}
    stop = threading.Event()

    def progress(stage, done, total, msg):
        if stage == "search" and done == 1:
            stop.set()                    # after the first name, before the second
    fake.names["Lamia textor"] = []       # BHL "lost" every Lamia textor page
    r = svc.run_search(sid, progress=progress, cancel=stop)
    assert r["status"] == "cancelled" and r["withdrawn"] == 0
    assert {i["page_id"] for i in svc.review_queue(sid)} == held


def test_cancel_during_text_keeps_withdrawals_of_complete_lookups(env):
    svc, fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    svc.update_search(sid, filters={"to_year": 1930})
    fake.names["Lamia textor"][0]["Titles"][0]["Items"][1]["Pages"].append(_page(110, 302))
    stop = threading.Event()

    def progress(stage, done, total, msg):
        if stage == "text":
            stop.set()
    r = svc.run_search(sid, progress=progress, cancel=stop)
    # both names were looked up in full before the cancel: 109 (1935) is rightly withdrawn
    assert r["status"] == "cancelled" and r["withdrawn"] == 1
    assert [p["reason"] for p in svc.run_diff(sid)["withdrawn"]] == [
        "Now filtered out (years, journals or index pages)"]


def test_a_failed_name_withdraws_none_of_its_pages(env):
    svc, fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)

    class Failing(FakeBHL):
        def name_metadata(self, name):
            if name == "Cerambyx textor":
                self.request_count += 1
                from Lector.bhl_client import BHLError
                raise BHLError("GetNameMetadata failed after 4 attempts (HTTP 500)")
            return super().name_metadata(name)
    svc._client = Failing()
    r = svc.run_search(sid, refresh=True)
    assert r["withdrawn"] == 0
    assert 201 in {i["page_id"] for i in svc.review_queue(sid)}


def test_withdrawn_reasons_name_duplicate_scans(env):
    """L11 investigation: with no filters, a run still drops duplicate scans (on by
    default); pages held from before (no de-duplication then) are withdrawn as such."""
    svc, fake = env
    # another library's scan of Annals v.7 p.145 (page 102)
    fake.names["Lamia textor"][0]["Titles"][1]["Items"].append(
        {"ItemID": 23, "Volume": "7", "Pages": [_page(111, 145)]})
    fake.texts[111] = fake.texts[102]
    sid = svc.create_search("Lamia textor", filters={"max_pages": 0, "skip_duplicates": False})
    svc.run_search(sid)
    svc.update_search(sid, filters={"skip_duplicates": True})
    r = svc.run_search(sid, refresh=False)
    d = svc.run_diff(sid)
    assert r["withdrawn"] == len(d["withdrawn"]) >= 1
    assert {p["reason"] for p in d["withdrawn"]} == {
        "Duplicate scan: the same printed page is kept from another copy"}
    # and a page BHL no longer lists
    fake.names["Cerambyx textor"] = []
    svc.run_search(sid, refresh=True)
    assert [p["reason"] for p in svc.run_diff(sid)["withdrawn"]] == [
        "No longer found on BHL for the names searched"]


# --- L12 one description of the filters ---------------------------------------------------------
@pytest.mark.parametrize("filters, scope", [
    ({}, "Any journal"),
    ({"include_words": ["Annals"]}, "Your title words (1)"),
    ({"exclude_words": ["Canadian"]}, "Any journal, 1 word left out"),
    ({"journal_preset": "British periodicals", "include_words": ["Entomologist"]},
     "British periodicals"),
    ({"recorded": False}, "Filters not recorded"),
])
def test_scope_is_described_identically(qapp, window, env, filters, scope):
    svc, _fake = env
    sid = svc.create_search("Lamia textor", filters={"max_pages": 0})
    svc.update_search(sid, filters=filters)
    listed = next(s for s in svc.list_searches() if s["search_id"] == sid)
    assert listed["scope"] == svc.get_search(sid)["scope"] == scope
    dash = window.go("searches")
    assert dash.table.cellWidget(0, 1).text().startswith(scope.replace(", ", ", "))
    h = window.go("history", search_id=sid)
    texts = [w.text() for w in h.summary.findChildren(type(h.title))]
    assert scope in texts and listed["years"] in texts


def test_builder_any_journal_clears_words(qapp, window, env):
    from Lector.ui.builder_filters import ANY_JOURNAL, OWN_WORDS
    svc, _fake = env
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    b = window.go("builder", search_id=sid)
    f = b.filters
    assert f.preset.currentText() == OWN_WORDS and f.include.words()
    f.preset.setCurrentIndex(f.preset.findText(ANY_JOURNAL))
    f._preset_chosen(0)
    assert f.filters()["include_words"] == [] and f.filters()["journal_preset"] is None
    f.preset.setCurrentIndex(f.preset.findText("British periodicals"))
    f._preset_chosen(0)
    assert f.filters()["journal_preset"] == "British periodicals" and f.include.words()
    b.save()
    assert svc.get_search(sid)["scope"] == "British periodicals"
    b = window.go("builder", search_id=svc.create_search("Monochamus sutor"))
    assert b.filters.preset.currentText() == ANY_JOURNAL


# --- L15 exports without authors / article titles -----------------------------------------------
def test_export_leaves_out_empty_author_columns(env):
    openpyxl = pytest.importorskip("openpyxl")
    pytest.importorskip("docx")
    svc, _fake = env
    plain = svc.create_search("Lamia textor", filters=ANNALS_ONLY, with_synonyms=False)
    svc.run_search(plain)
    out = svc.export(plain, formats=("docx", "xlsx", "txt"))
    wb = openpyxl.load_workbook(out["files"]["xlsx"])
    head = [c.value for c in wb["Index"][1]]
    assert "Authors" not in head and "Article title" not in head and "BHL link" in head
    about = [tuple(r) for r in wb["About"].iter_rows(values_only=True)]
    assert ("Note", "Authors and article titles: none given by BHL for these pages") in about
    link = wb["Index"].cell(row=2, column=head.index("BHL link") + 1)
    assert link.hyperlink is not None
    with open(out["files"]["csv"], encoding="utf-8-sig", newline="") as fh:
        assert not {"authors", "article_title"} & set(next(csv.reader(fh)))
    # a search where BHL gives an article (page 101, Waterhouse): the columns stay
    withp = svc.create_search("Lamia textor", name="With 101", filters=BRITISH)
    svc.run_search(withp)
    out = svc.export(withp, formats=("docx", "xlsx", "txt"))
    head = [c.value for c in openpyxl.load_workbook(out["files"]["xlsx"])["Index"][1]]
    assert "Authors" in head and "Article title" in head
    with open(out["files"]["csv"], encoding="utf-8-sig", newline="") as fh:
        assert {"authors", "article_title"} <= set(next(csv.reader(fh)))


# --- L16 closing during a run --------------------------------------------------------------------
def test_close_during_a_run_asks_and_can_keep_running(qapp, tmp_path):
    from Lector.ui.app import CLOSE_QUESTION
    svc, fake = make_env(tmp_path, SlowBHL())
    win = make_window(qapp, svc)
    try:
        sid = svc.create_search("Lamia textor", filters=BRITISH)
        win.go("history", search_id=sid, run="fetch")
        assert pump(qapp, lambda: fake.request_count >= 2)
        win.answers = [False]                                   # "Keep running"
        assert win.close() is False
        assert win.messages[-1] == CLOSE_QUESTION
        n = fake.request_count
        assert pump(qapp, lambda: fake.request_count > n)       # still going
        assert win.isVisible() and win.tasks.busy
    finally:
        win.answers = [True]
        win.close()
        pump(qapp, lambda: not win.tasks.busy and not win.isVisible())
    assert svc.list_runs(sid)[0]["status"] == "cancelled"


def test_schema_v3_adds_only(tmp_path):
    """Version 2 -> 3 adds the request columns and the activity table; no stored row
    changes and no saved search is made again."""
    svc, fake = make_env(tmp_path)
    sid = svc.create_search("Lamia textor", filters=BRITISH)
    svc.run_search(sid)
    db = tmp_path / "lector.db"
    with sqlite3.connect(db) as c:                              # back to a version 2 file
        c.execute("PRAGMA user_version = 2")
        before = schema.counts(c)
    assert upgrade.status(db)["needs_update"]
    done = svc.update_database()
    assert done["applied"] and Path(done["backup"]).name.startswith("lector_pre_v3_")
    assert done["searches_made"] == []
    for k, v in before.items():
        assert done["after"][k] == v, k
    assert json.loads(json.dumps(svc.list_runs(sid)[0]["settings"]))["filters"]
