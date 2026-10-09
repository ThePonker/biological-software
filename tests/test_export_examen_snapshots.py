"""Export of examen.db's frozen assessments to workbooks (backlog E13, 9 Oct 2026)."""
import csv
import hashlib
import sqlite3

from scripts import export_examen_snapshots as ex

# The schema examen.db carries (from the retired Examen/snapshot_manager.py).
SCHEMA = """
CREATE TABLE site_snapshots (id INTEGER PRIMARY KEY AUTOINCREMENT, site_name TEXT NOT NULL,
  project_name TEXT DEFAULT '', client TEXT DEFAULT '', survey_year INTEGER NOT NULL,
  first_visit TEXT, last_visit TEXT, visit_count INTEGER DEFAULT 0,
  species_count INTEGER DEFAULT 0, key_species_count INTEGER DEFAULT 0,
  key_species_pct REAL DEFAULT 0.0, rare_count INTEGER DEFAULT 0, scarce_count INTEGER DEFAULT 0,
  priority_count INTEGER DEFAULT 0, sqi REAL DEFAULT 0.0, sqi_reliable INTEGER DEFAULT 1,
  scoring_species INTEGER DEFAULT 0, analysis_mode TEXT NOT NULL DEFAULT 'codex_full',
  codex_version TEXT DEFAULT '', pantheon_version TEXT DEFAULT '', frozen_date TEXT NOT NULL,
  notes TEXT DEFAULT '', UNIQUE(site_name, project_name, survey_year, analysis_mode));
CREATE TABLE snapshot_species (id INTEGER PRIMARY KEY AUTOINCREMENT, snapshot_id INTEGER NOT NULL,
  tvk TEXT NOT NULL, species_name TEXT NOT NULL, short_status TEXT DEFAULT '',
  tier TEXT DEFAULT '', sqs INTEGER DEFAULT 0);
CREATE TABLE snapshot_log (id INTEGER PRIMARY KEY AUTOINCREMENT, snapshot_id INTEGER NOT NULL,
  action TEXT NOT NULL, timestamp TEXT NOT NULL, notes TEXT);
"""


def _db(path, rows=True):
    c = sqlite3.connect(path)
    c.executescript(SCHEMA)
    if rows:
        c.execute("INSERT INTO site_snapshots (site_name, project_name, client, survey_year, "
                  "sqi, analysis_mode, frozen_date) VALUES "
                  "('Glory Park', 'GP Survey', 'NE', 2024, 117, 'codex_full', '2025-01-02T10:00')")
        c.execute("INSERT INTO site_snapshots (site_name, project_name, survey_year, "
                  "analysis_mode, frozen_date) VALUES ('Glory Park', 'GP Survey', 2024, "
                  "'pantheon_only', '2025-01-02T10:05')")
        c.executemany("INSERT INTO snapshot_species (snapshot_id, tvk, species_name, sqs) "
                      "VALUES (?,?,?,?)",
                      [(1, "T1", "Lamia textor", 16), (1, "T2", "Carabus nemoralis", 1),
                       (2, "T1", "Lamia textor", 16), (9, "T3", "Orphan species", 0)])
        c.execute("INSERT INTO snapshot_log (snapshot_id, action, timestamp) "
                  "VALUES (1, 'frozen', '2025-01-02T10:00')")
    c.commit()
    c.close()
    return path


def _md5(p):
    return hashlib.md5(p.read_bytes()).hexdigest()


def test_dry_run_writes_nothing(tmp_path):
    db = _db(tmp_path / "examen.db")
    out = tmp_path / "out"
    planned = ex.export(db, str(out), write=False)
    assert [k for _, k, _ in planned] == ["assessment", "assessment", "other", "index"]
    assert not out.exists()


def test_write_one_workbook_per_assessment(tmp_path):
    from openpyxl import load_workbook
    db = _db(tmp_path / "examen.db")
    before = _md5(db)
    out = tmp_path / "out"
    ex.export(db, str(out), write=True)
    assert _md5(db) == before                       # examen.db untouched
    files = sorted(p.name for p in out.iterdir())
    assert files == ["001_Glory_Park_GP_Survey_2024_codex_full.xlsx",
                     "002_Glory_Park_GP_Survey_2024_pantheon_only.xlsx",
                     "examen_other_tables.xlsx", "index.csv"]
    wb = load_workbook(out / files[0])
    assert wb.sheetnames == ["About", "Assessment", "Species", "Log"]
    about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True)}
    assert about["Site"] == "Glory Park" and about["Date frozen"] == "2025-01-02T10:00"
    assert "archived frozen assessment" in about["Note"]
    species = list(wb["Species"].iter_rows(values_only=True))
    assert len(species) == 3 and species[1][3] == "Lamia textor"
    other = load_workbook(out / "examen_other_tables.xlsx")
    assert len(list(other.worksheets[1].iter_rows(values_only=True))) == 2   # header + orphan
    with open(out / "index.csv", encoding="utf-8-sig") as f:
        idx = list(csv.DictReader(f))
    assert [r["species_rows"] for r in idx[:2]] == ["2", "1"]


def test_never_overwrites(tmp_path):
    db = _db(tmp_path / "examen.db")
    out = tmp_path / "out"
    ex.export(db, str(out), write=True)
    first = _md5(out / "index.csv")
    ex.export(db, str(out), write=True)
    assert _md5(out / "index.csv") == first
    assert (out / "index_2.csv").exists()
    assert (out / "001_Glory_Park_GP_Survey_2024_codex_full_2.xlsx").exists()


def test_empty_database_plans_nothing(tmp_path):
    db = _db(tmp_path / "examen.db", rows=False)
    assert ex.export(db, str(tmp_path / "out"), write=True) == []
    assert not (tmp_path / "out").exists()
