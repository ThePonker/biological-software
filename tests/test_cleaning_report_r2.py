"""Cleaning report (backlog B, 10 Oct 2026): services/cleaning_report.py, cleaning_export.py
and the Stats/Reports screen views/stats/cleaning_dashboard.py.

Small temp databases for the rules; the data copies (read-only) for a smoke test and timing.
"""
import os
import sqlite3
import time

import pytest

from src.services.cleaning_report import (
    IRECORD_URL, PROBLEMS, base_name, build_report, date_problem, display_date, grid_problem,
    grid_size_m, irecord_url, synonym_groups,
)
from src.services.cleaning_export import SHEET_NAMES, export_csv_folder, export_xlsx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_MAIN = os.path.join(ROOT, "data", "observatum.db")
DATA_UKSI = os.path.join(ROOT, "data", "uksi.db")
HAVE_DATA = os.path.exists(DATA_MAIN) and os.path.exists(DATA_UKSI)


# ------------------------------------------------------------------ pure rules

@pytest.mark.parametrize("ref,size", [
    ("SP58", 10000), ("SP 58", 10000), ("SP", 100000), ("~SX", 100000), ("SP5820", 1000),
    ("SP580207", 100), ("TQ 53083 77347", 1), ("SP58A", 2000), ("SP58Z", 2000), ("SP58NW", 5000),
    ("H45", 10000), ("H4567", 1000), ("49.4557N2.5591W", 0), ("~49.442N2.641W", 0),
    ("SP583", None), ("rubbish", None), ("", None), (None, None),
])
def test_grid_size(ref, size):
    assert grid_size_m(ref) == size


def test_grid_problem():
    assert grid_problem("SP58") == "Coarse grid ref (10 km)"
    assert grid_problem("SK") == "Coarse grid ref (100 km)"
    assert grid_problem("SP58A") is None                      # tetrad: not coarse
    assert grid_problem("SP5820") is None
    assert grid_problem("49.4557N2.5591W") is None            # lat/long text: not flagged
    assert grid_problem("") == "No grid ref"
    assert grid_problem(None, has_latlon=True) == "No grid ref (lat/long held)"
    assert grid_problem("SP583") == "Grid ref not readable"


def test_dates():
    assert date_problem("2025-06-01", "D") is None
    assert date_problem("2025-06-01", None) is None
    assert date_problem("", "D") == "No date"
    assert date_problem(None, None) == "No date"
    assert date_problem("1956-06-01", "O") == "Vague date: month only (O)"
    assert date_problem("2024-05-30", "DD") == "Vague date: day range (DD)"
    assert date_problem("2014-01-01", "Y") == "Vague date: year only (Y)"
    assert date_problem("2014-01-01", "ZZ") == "Vague date: not a single day (ZZ)"
    assert display_date("1956-06-01", "O") == "1956-06"
    assert display_date("2014-01-01", "YY") == "2014"
    assert display_date("2024-05-30", "DD") == "2024-05-30"


def test_irecord_url_one_pattern():
    assert irecord_url(53796418) == IRECORD_URL.format(id=53796418)
    assert irecord_url("127178").endswith("occurrence_id=127178")
    assert irecord_url(None) is None and irecord_url("") is None and irecord_url(0) is None


def test_base_name():
    assert base_name("Leiopus linnei/nebulosus agg.") == base_name("Leiopus linnei/nebulosus")
    assert base_name("Xanthogramma pedissequum sensu lato") == "xanthogramma pedissequum"
    assert base_name("Enoplognatha ovata s.l.") == "enoplognatha ovata"
    assert base_name("Rhopalus (Rhopalus) subrufus") == "rhopalus subrufus"
    assert base_name("Leiopus sp.") == "leiopus"


def test_synonym_groups_only_true_doubles():
    current = {"OLD1": "CUR1", "CUR1": "CUR1", "CUR2": "CUR2", "AGG": "AGG", "LOST": None}
    pairs = {("OLD1", "Strangalia old"), ("CUR1", "Rutpela maculata"),        # two TVKs, one sp.
             ("CUR2", "Clytus arietis"),                                       # alone
             ("AGG", "Leiopus linnei/nebulosus"), ("AGG", "Leiopus linnei/nebulosus agg."),
             ("LOST", "Nowhere"), ("", "No tvk")}
    out = synonym_groups(pairs, current)
    assert set(out) == {("OLD1", "Strangalia old"), ("CUR1", "Rutpela maculata")}
    assert out[("OLD1", "Strangalia old")] == ("CUR1", ["Rutpela maculata"])
    assert out[("CUR1", "Rutpela maculata")] == ("CUR1", ["Strangalia old"])
    # same TVK, two genuinely different names (scheme counts by name): flagged
    out2 = synonym_groups({("CUR1", "Rutpela maculata"), ("CUR1", "Strangalia maculata")}, current)
    assert len(out2) == 2


# ------------------------------------------------------------------ temp databases

COLS = ("id INTEGER PRIMARY KEY, irecord_id INTEGER, species_name TEXT, species_tvk TEXT, "
        "taxon_rank TEXT, date TEXT, date_type TEXT, grid_ref TEXT, site_name TEXT, "
        "recorder TEXT, latitude REAL, longitude REAL")
ROWS = [  # id, irecord_id, name, tvk, rank, date, date_type, grid, site, recorder, lat, lon
    (1, 1001, "Rutpela maculata", "CRT0000000000001", "Species", "2025-06-01", "D", "SP5820", "A", "W", None, None),
    (2, None, "Leiopus sp.", "CRT0000000000002", None, "2025-06-02", "D", "SP5820", "A", "W", None, None),
    (3, 1003, "Cerambycidae", "CRT0000000000003", "Family", "2025-06-03", "D", "SP58A", "A", "W", None, None),
    (4, 1004, "Leiopus linnei/nebulosus", "CRT0000000000004", "Species aggregate", "2025-06-04", "D", "SP5820", "A", "W", None, None),
    (5, 1005, "Hepialus humuli thulensis", "CRT0000000000005", None, "2025-06-05", "D", "SP5820", "A", "W", None, None),
    (6, 1006, "Strangalia maculata", "CRT0000000000009", "Species", "2025-06-06", "D", "SP5820", "A", "W", None, None),
    (7, 1007, "Clytus arietis", "CRT0000000000006", "Species", "", "D", "SP58", "A", "W", None, None),
    (8, 1008, "Clytus arietis", "CRT0000000000006", "Species", "1956-06-01", "O", "", "A", "W", 52.0, -1.0),
    (9, 1009, "Mystery beetle", "", None, "2025-06-09", "D", "SP5820", "A", "W", None, None),
    (10, 1010, "Unplaced beetle", "CRT0000000000099", "Species", "2025-06-10", "D", "SP5820", "A", "W", None, None),
    (11, 1011, "Salix x multinervis", "CRT0000000000007", "Species hybrid", "2025-06-11", "D", "SP5820", "A", "W", None, None),
    (12, 1012, "Bombus terrestris s.s.", "CRT0000000000008", "Species", "2025-06-12", "D", "SP5820", "A", "W", None, None),
    (13, 1013, "Xanthogramma pedissequum", "CRT0000000000010", "Species sensu lato", "2025-06-13", "D", "SP5820", "A", "W", None, None),
]
TAXA = [  # tvk, name, rank, parent
    ("CRT0000000000001", "Rutpela maculata", "Species", None),
    ("CRT0000000000002", "Leiopus", "Genus", None),
    ("CRT0000000000003", "Cerambycidae", "Family", None),
    ("CRT0000000000004", "Leiopus linnei/nebulosus", "Species aggregate", None),
    ("CRT0000000000005", "Hepialus humuli thulensis", "Subspecies", None),
    ("CRT0000000000006", "Clytus arietis", "Species", None),
    ("CRT0000000000007", "Salix x multinervis", "Species hybrid", None),
    ("CRT0000000000008", "Bombus terrestris", "Species", None),
    ("CRT0000000000010", "Xanthogramma pedissequum", "Species sensu lato", None),
]


@pytest.fixture()
def dbs(tmp_path):
    main, uksi = tmp_path / "obs.db", tmp_path / "uksi.db"
    c = sqlite3.connect(main)
    for table in ("observations", "recording_scheme"):
        c.execute(f"CREATE TABLE {table} ({COLS})")
        c.executemany(f"INSERT INTO {table} VALUES ({','.join('?' * 12)})", ROWS)
    c.commit()
    c.close()
    u = sqlite3.connect(uksi)
    u.execute('CREATE TABLE taxa (tvk TEXT PRIMARY KEY, scientific_name TEXT, rank TEXT, '
              'parent_tvk TEXT, kingdom TEXT, "order" TEXT, family TEXT, genus TEXT)')
    u.executemany("INSERT INTO taxa (tvk, scientific_name, rank, parent_tvk) VALUES (?,?,?,?)", TAXA)
    u.execute("CREATE TABLE common_names (common_name TEXT, tvk TEXT, preferred INTEGER)")
    u.execute("CREATE TABLE synonyms (synonym TEXT, tvk TEXT)")
    u.execute("CREATE TABLE name_map (tvk TEXT PRIMARY KEY, recommended_tvk TEXT NOT NULL, name TEXT, "
              "recommended_name TEXT, name_status TEXT, deprecated TEXT, source TEXT)")
    u.execute("INSERT INTO name_map (tvk, recommended_tvk) VALUES ('CRT0000000000009', 'CRT0000000000001')")
    u.execute("CREATE TABLE tvk_remap (old_tvk TEXT PRIMARY KEY, new_tvk TEXT, method TEXT, "
              "old_name TEXT, new_name TEXT)")
    u.commit()
    u.close()
    return str(main), str(uksi)


def _ids(rep, key):
    return sorted(r.record_id for r in rep.rows[key])


def test_build_report_problems(dbs):
    main, uksi = dbs
    for source in ("observations", "recording_scheme"):
        rep = build_report(main, uksi, source)
        assert rep.total_records == 13
        assert _ids(rep, "genus") == [2]                     # rank from UKSI (none stored)
        assert _ids(rep, "family_plus") == [3]
        assert _ids(rep, "agg") == [4]
        assert _ids(rep, "subsp") == [5]                     # UKSI rank Subspecies
        assert _ids(rep, "hybrid") == [11]
        assert _ids(rep, "sl") == [13]
        assert _ids(rep, "no_tvk") == [9, 10]
        assert _ids(rep, "synonym") == [1, 6]                # old TVK + current TVK, one species
        assert _ids(rep, "date") == [7, 8]
        assert _ids(rep, "grid") == [7, 8]                   # hectad, none; tetrad (3) is fine
        assert 12 not in {r.record_id for rows in rep.rows.values() for r in rows}   # s.s. fine
        syn = {r.record_id: r.detail for r in rep.rows["synonym"]}
        assert syn[1] == "Same species as: Strangalia maculata"
        assert syn[6] == "Same species as: Rutpela maculata"
        assert rep.records_with_problem == 12
        assert sum(rep.counts.values()) == 14
        row8 = rep.rows["date"][1]
        assert (row8.date, row8.detail, row8.irecord_id) == ("1956-06", "Vague date: month only (O)", 1008)


def test_unknown_source_refused(dbs):
    with pytest.raises(ValueError):
        build_report(dbs[0], dbs[1], "specimens; DROP TABLE x")


def test_export_xlsx_and_csv(dbs, tmp_path):
    rep = build_report(dbs[0], dbs[1], "observations")
    path = export_xlsx(rep, str(tmp_path / "c.xlsx"))
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        assert wb.sheetnames == ["Summary"] + [SHEET_NAMES[k] for k, _ in PROBLEMS]
        summary = {r[0]: r[1] for r in wb["Summary"].values if r and r[0]}
        assert summary["Records checked"] == 13 and summary["Records with at least one problem"] == 12
        grid = list(wb["Grid ref"].values)
        assert grid[0][7] == "iRecord ID" and {g[7] for g in grid[1:]} == {1007, 1008}
        assert grid[1][8] == IRECORD_URL.format(id=grid[1][7])
        genus = list(wb["Genus only"].values)
        assert genus[1][7] is None and genus[1][9] == 2          # no iRecord ID: blank, Observatum ID kept
    finally:
        wb.close()
    folder = export_csv_folder(rep, str(tmp_path / "csv"))
    assert sorted(os.listdir(folder))[0] == "aggregates.csv" and "summary.csv" in os.listdir(folder)


# ------------------------------------------------------------------ the screen (offscreen)

@pytest.fixture(scope="module")
def qapp():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_dashboard_lists_and_opens(qapp, dbs, monkeypatch):
    from src.views.stats.cleaning_dashboard import CleaningDashboard
    d = CleaningDashboard(main_db=dbs[0], uksi_db=dbs[1])
    d.refresh()
    assert d.tabs.tabText(0) == "Genus only (1)"
    assert d._tiles["grid"].text().endswith("\n2")
    opened, shown = [], []
    d.open_url = lambda u: opened.append(u.toString())
    monkeypatch.setattr(d, "open_in_observatum", lambda row: shown.append(row.record_id))
    d.tabs.setCurrentIndex([k for k, _ in PROBLEMS].index("grid"))
    view = d.tabs.currentWidget()
    d._on_double_click(view.model().index(0, 0))
    assert opened == [IRECORD_URL.format(id=1007)]
    d.tabs.setCurrentIndex(0)                                 # genus: record 2 has no iRecord ID
    view = d.tabs.currentWidget()
    d._on_double_click(view.model().index(0, 0))
    assert shown == [2]
    d.cmb_source.setCurrentIndex(1)
    assert d.report().source == "recording_scheme"


@pytest.mark.skipif(not HAVE_DATA, reason="data copies absent")
def test_detail_dialog_view_only_on_data_copy(qapp):
    from PySide6.QtWidgets import QPushButton
    from src.services.cleaning_report import CleaningRow
    from src.views.stats.cleaning_dashboard import CleaningDashboard
    conn = sqlite3.connect(f"file:{DATA_MAIN}?mode=ro", uri=True)
    obs_id = conn.execute("SELECT MIN(id) FROM observations").fetchone()[0]
    rs_id = conn.execute("SELECT MIN(id) FROM recording_scheme").fetchone()[0]
    conn.close()
    d = CleaningDashboard(main_db=DATA_MAIN, uksi_db=DATA_UKSI)
    for idx, rid in ((0, obs_id), (1, rs_id)):
        d.cmb_source.blockSignals(True)
        d.cmb_source.setCurrentIndex(idx)
        d.cmb_source.blockSignals(False)
        row = CleaningRow(rid, "", "", "", "", "", "", None, "grid", "")
        dlg = d.detail_dialog(row)
        texts = [b.text() for b in dlg.findChildren(QPushButton) if b.isVisibleTo(dlg)]
        assert "Edit" not in texts and "Delete" not in texts
        dlg.deleteLater()


@pytest.mark.skipif(not HAVE_DATA, reason="data copies absent")
def test_data_copy_figures_and_speed():
    conn = sqlite3.connect(f"file:{DATA_MAIN}?mode=ro", uri=True)
    no_date = conn.execute("SELECT COUNT(*) FROM recording_scheme WHERE COALESCE(date,'')=''").fetchone()[0]
    vague = conn.execute("SELECT COUNT(*) FROM recording_scheme WHERE COALESCE(date,'')<>'' "
                         "AND COALESCE(date_type,'D') NOT IN ('D','')").fetchone()[0]
    conn.close()
    from shared.species_search import get_index
    get_index(DATA_UKSI)                                    # warmed at start-up in the app
    t0 = time.perf_counter()
    rep = build_report(DATA_MAIN, DATA_UKSI, "recording_scheme")
    assert time.perf_counter() - t0 < 10
    assert rep.counts["date"] == no_date + vague
    assert rep.total_records > 100000
